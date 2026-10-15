"""V2-a: is the VTM decoder output bit-identical to the encoder reconstruction on 59 pairs? Run: python studies/P001_vvc-dental/run_v2a.py

Study P001, cycle 3: check V2-a of STUDY.md section 8 for the VTM wrapper (../REFERENCE.md#b05),
plus determinism of the bitstream. Paths come from lab_paths.json (images_root, cache_root,
coders.vtm). Writes <cache_root>/cycles/3/v2a_vtm.json (refuses by name if it exists) and the
book v2a_vtm.xlsx beside it; `--rebuild-book` rebuilds the book from the run file.
"""

import argparse
import datetime
import hashlib
import json
import math
import platform
import sys
import time
from pathlib import Path

import numpy as np
from PIL import Image

import imstudies
from imstudies.core.metrics import Refusal, psnr
from imstudies.lab.coders import vtm
from imstudies.lab.images import load_test_images
from imstudies.lab.runs import RunRefusal, git_state, json_number, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]

COMMAND = "python studies/P001_vvc-dental/run_v2a.py"
QUESTION = ("V2-a: совпадает ли побитово выход DecoderApp с восстановлением EncoderApp (VTM-23.0) на 59 парах "
            "«изображение × QP», и детерминирован ли поток?")
CYCLE, RUN_FILE, BOOK_FILE = 3, "v2a_vtm.json", "v2a_vtm.xlsx"
QPS = (-12, 0, 12, 22, 32, 42, 51, 63)
ODD_SHAPES = ((8, 8), (21, 13), (517, 389))  # (H, W), cropped from the E0 recipe on 517 x 512
# The six test images of cycle 1, read as in run_v1.py (imstudies.lab.images, parity tested).
TEST_IMAGES = ["BABOON.BMP", "BARBARA.BMP", "goldhill.bmp", "lenna.bmp", "frisco.bmp", "DIEGO.BMP"]
TIMEOUT_S = {"value": 600, "basis": "per EncoderApp or DecoderApp run: over 10 times the longest call of "
                                    "cycle 3 (43 s, 512 x 512 at QP 0); an operational limit, no effect on numbers"}
LEVEL = {"value": "5", "status": "provisional",
         "basis": "smallest VVC level covering all P001 sources by header sizes (19 files: max H*W 4170592, "
                  "max H 1504, max W 2789): Level 4/4.1 MaxLumaPs 2228224 < 4170592 <= 8912896 of Level 5; "
                  "side <= sqrt(8*MaxLumaPs) = 8444; table mainTierLevelInfo, VTM-23.0 "
                  "source/Lib/CommonLib/ProfileTierLevel.cpp lines 114-133"}
FRAME_RATE = {"value": 1, "status": "provisional",
              "basis": "EncoderApp refuses FrameRate 0 ('Frame rate cannot be 0 or less'); B05 names no value; "
                       "owner decision 2026-10-14, 17:04 UTC: 1, provisional. FrameRate 1, 25, 50, 60 gave "
                       "bit-identical bitstream and reconstruction on frame E0 at QP 32 (diagnostic, cycle 3)"}
DECLARED = {
    "data": ("Все изображения отладочные, в приёмке P001 не участвуют: синтетический кадр E0; шесть тестовых BMP "
             "цикла 1 (Images/BMP, чтение как в run_v1.py); синтетика 8 × 8, 21 × 13, 517 × 389 (H × W; рецепт E0 "
             "на 517 × 512, seed 1, вырезка от левого верхнего угла). Стоматологические пиксели не используются."),
    "conditions": ("QP ∈ {−12, 0, 12, 22, 32, 42, 51, 63} для E0 и шести BMP (56 пар); QP 32 для трёх "
                   "нестандартных размеров (3 пары). Настройки — B05 «Внешний кодер»."),
    "criterion": ("V2-a: на всех 59 парах encoder_recon == recon побитово и recon.shape == image.shape; "
                  "детерминизм: кадр E0 на всех 8 QP кодируется второй раз — sha256 потока совпадает."),
    "refutation": "Хотя бы одно несовпадение — «не пройдено», первая строка отчёта, разбор без подгонки.",
    "expected": ("Без вердикта, печатается рядом: число бит не возрастает с ростом QP на каждом изображении; "
                 "|PSNR-Y лога − наш PSNR (B02) по recon| — сотые доли дБ (E0: 31.4004 против 31.3932 дБ при QP 32)."),
    "path_control": ("Первым: кадр E0 при QP 32 — поток, recon, совпадение V2-a, наш PSNR против PSNR-Y лога; "
                     "расхождение — стоп."),
    "pages": "Страницы для глаз не применяются — побитовая сверка; визуальная проверка — V2-d, позже.",
}
MAX_PIXELS_RULE = ("max_pixels = H*W of the image itself: the limit is not under test in V2-a, and "
                   "config.json core_params.io.max_image_pixels is null (NOT ASSIGNED)")


def e0_recipe(h, w):
    y, x = np.mgrid[0:h, 0:w]
    return (x / 2 + y / 4 + np.random.default_rng(1).normal(0, 8, (h, w))).clip(0, 255).astype(np.uint8)


def images_of_the_cycle(images_root):
    images = [{"code": "e0", "class": "control", "source": "E0 recipe 512 x 512, seed 1", "array": e0_recipe(512, 512)}]
    for img in load_test_images(images_root, TEST_IMAGES, (512, 512)):
        images.append({"code": img["code"], "class": "test", "source": img["source"],
                       "source_sha256": img["source_sha256"], "array": img["array"]})
    base = e0_recipe(517, 512)
    for h, w in ODD_SHAPES:
        images.append({"code": f"syn{h}x{w}", "class": "synthetic", "source": "E0 recipe 517 x 512, seed 1, top-left crop",
                       "array": np.ascontiguousarray(base[:h, :w])})
    return images


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {k: (ROOT / raw[k]).resolve() for k in ("images_root", "cache_root") if raw.get(k)}
    for key, value in (raw.get("coders", {}).get("vtm") or {}).items():
        paths[key] = (ROOT / value).resolve()
    for key in ("images_root", "cache_root", "encoder", "decoder", "cfg"):
        if key not in paths:
            raise RunRefusal(f"lab_paths.json has no key '{key}' (coders.vtm.{key} for the codec)")
    for key in ("encoder", "decoder", "cfg"):
        if not paths[key].is_file():
            raise RunRefusal(f"VTM {key} not found: {paths[key]} (build: tools/build_coders.ps1 vtm)")
    manifest = paths["encoder"].parent / "MANIFEST.json"
    if not manifest.is_file():
        raise RunRefusal(f"{manifest} not found: the executables are not a build of tools/build_coders.ps1")
    paths["manifest"] = json.loads(manifest.read_text(encoding="utf-8"))
    for key in ("encoder", "decoder", "cfg"):
        if sha256_file(paths[key]) != paths["manifest"]["files_sha256"].get(paths[key].name):
            raise RunRefusal(f"{paths[key].name} differs from its MANIFEST.json sha256")
    return paths


NO_RESULT = {"bits": None, "bpp": None, "cr": None, "bitstream_sha256": None, "recon_sha256": None,
             "encoder_recon_sha256": None, "shape_ok": False, "v2a_match": False, "psnr_b02": None,
             "log_psnr_y": None, "log_minus_ours": None, "seconds": None}


def encode(paths, img, qp, work):
    """One wrapper call mapped to SPEC-RESULT: (row fields, result or None)."""
    image = img["array"]
    started = time.perf_counter()
    try:
        result = vtm.encode_decode(image, qp, max_pixels=image.size, level=LEVEL["value"],
                                   frame_rate=FRAME_RATE["value"], timeout_s=TIMEOUT_S["value"],
                                   encoder=paths["encoder"],
                                   decoder=paths["decoder"], cfg=paths["cfg"], work_root=work)
    except Refusal as refusal:
        return NO_RESULT | {"status": "no_result", "reason": refusal.reason, "detail": str(refusal)}, None
    except Exception as exc:  # noqa: BLE001 - status `error` by SPEC-RESULT; the detail is kept
        detail = f"{type(exc).__name__}: {exc}"[:600]
        return NO_RESULT | {"status": "error", "reason": "internal_error", "detail": detail}, None
    h, w = image.shape
    bits = 8 * len(result.bitstream)
    ours = psnr(image, result.recon, max_pixels=image.size).value
    row = {
        "status": "valid", "reason": None, "detail": None, "bits": bits, "bpp": bits / (h * w),
        "cr": (8 * h * w) / bits, "bitstream_sha256": hashlib.sha256(result.bitstream).hexdigest(),
        "recon_sha256": hashlib.sha256(result.recon.tobytes()).hexdigest(),
        "encoder_recon_sha256": hashlib.sha256(result.encoder_recon.tobytes()).hexdigest(),
        "shape_ok": result.recon.shape == image.shape,
        "v2a_match": result.recon.shape == image.shape and np.array_equal(result.recon, result.encoder_recon),
        "psnr_b02": json_number(ours), "log_psnr_y": json_number(result.log_psnr_y),
        "log_minus_ours": (result.log_psnr_y - ours if result.log_psnr_y is not None
                           and math.isfinite(ours) and math.isfinite(result.log_psnr_y) else None),
        "seconds": round(time.perf_counter() - started, 2),
    }
    return row, result


def base_row(img, qp):
    h, w = img["array"].shape
    return {"image_code": img["code"], "image_class": img["class"], "qp": qp, "h": h, "w": w,
            "level": LEVEL["value"], "level_status": LEVEL["status"],
            "frame_rate": FRAME_RATE["value"], "frame_rate_status": FRAME_RATE["status"]}


def summarize(rows, determinism):
    valid = [r for r in rows if r["status"] == "valid"]
    monotonic = {}
    for code in dict.fromkeys(r["image_code"] for r in rows):
        series = sorted((r["qp"], r["bits"]) for r in valid if r["image_code"] == code)
        if len(series) > 1:
            monotonic[code] = all(b2 <= b1 for (_, b1), (_, b2) in zip(series, series[1:]))
    gaps = [abs(r["log_minus_ours"]) for r in valid if r["log_minus_ours"] is not None]
    v2a = len(rows) > 0 and all(r["status"] == "valid" and r["v2a_match"] for r in rows)
    det = len(determinism) > 0 and all(d["equal"] for d in determinism)
    return {
        "n_pairs": len(rows), "n_images": len({r["image_code"] for r in rows}), "n_valid": len(valid),
        "n_v2a_match": sum(r["status"] == "valid" and r["v2a_match"] for r in rows),
        "n_determinism": len(determinism), "n_determinism_equal": sum(d["equal"] for d in determinism),
        "bits_non_increasing_by_image": monotonic,
        "n_images_bits_non_increasing": sum(monotonic.values()), "n_images_with_series": len(monotonic),
        "n_log_psnr_compared": len(gaps), "max_abs_log_minus_ours_db": max(gaps) if gaps else None,
        "median_abs_log_minus_ours_db": float(np.median(gaps)) if gaps else None,
        "n_infinite_psnr_b02": sum(r["psnr_b02"] == "+inf" for r in valid),
        "verdict_v2a": "pass" if v2a else "fail", "verdict_determinism": "pass" if det else "fail",
        "refutation_fired": not (v2a and det),
    }


def run(paths, images, plan, determinism_plan, out_dir, state):
    """Path control, the pairs of `plan` ((code, qp)), re-encodes of `determinism_plan`, run file, book."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path = out_dir / RUN_FILE, out_dir / BOOK_FILE
    if run_path.exists():
        raise RunRefusal(f"{run_path} already exists: a run file of cycle {CYCLE} is never overwritten")
    started = datetime.datetime.now(datetime.timezone.utc)
    by_code = {img["code"]: img for img in images}
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")

    control_row, control = encode(paths, by_code["e0"], 32, out_dir)
    control_row = base_row(by_code["e0"], 32) | control_row
    print(f"  path control E0 QP 32: status {control_row['status']}, bits {control_row.get('bits')}, "
          f"V2-a {control_row.get('v2a_match')}, PSNR B02 {control_row.get('psnr_b02')} dB, "
          f"log PSNR-Y {control_row.get('log_psnr_y')} dB (E0 journal: 2216 bits, 31.4004 / 31.3932 dB)")
    if control is None or not control_row["v2a_match"]:
        raise RunRefusal(f"path control failed: {control_row.get('detail') or 'V2-a mismatch'}; stop before the pairs")

    rows = []
    for k, (code, qp) in enumerate(plan, start=1):
        row, _ = encode(paths, by_code[code], qp, out_dir)
        rows.append({"pair_index": k} | base_row(by_code[code], qp) | row)
        print(f"  {k:>2}/{len(plan)} {code:<10} QP {qp:>3}: {row['status']} bits {row.get('bits')} "
              f"V2-a {row.get('v2a_match')} ({row.get('seconds')} s)")
    determinism = []
    for code, qp in determinism_plan:
        row, _ = encode(paths, by_code[code], qp, out_dir)
        first = next(r for r in rows if (r["image_code"], r["qp"]) == (code, qp))
        determinism.append({"image_code": code, "qp": qp, "first_sha256": first.get("bitstream_sha256"),
                            "second_sha256": row.get("bitstream_sha256"),
                            "equal": row["status"] == "valid" and row["bitstream_sha256"] == first.get("bitstream_sha256")
                            and row["recon_sha256"] == first.get("recon_sha256")})
    summary = summarize(rows, determinism)
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        **state, "declared_before_numbers": DECLARED, "level": LEVEL, "frame_rate": FRAME_RATE, "timeout_s": TIMEOUT_S,
        "max_pixels_rule": MAX_PIXELS_RULE, "build_manifest": paths["manifest"],
        "wrapper_provenance": {k: v for k, v in control.provenance.items() if not k.endswith("_command")}
        | {"encoder_command_template": [a if not a.endswith((".yuv", ".bin")) else Path(a).name
                                        for a in control.provenance["encoder_command"]],
           "decoder_command_template": [a if not a.endswith((".yuv", ".bin")) else Path(a).name
                                        for a in control.provenance["decoder_command"]]},
        "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__,
                        "numpy": np.__version__, "pillow": Image.__version__, "platform": platform.platform()},
        "images": [{k: v for k, v in img.items() if k != "array"} | {
            "shape": list(img["array"].shape), "pixels_sha256": hashlib.sha256(img["array"].tobytes()).hexdigest()}
            for img in images],
        "unit_of_analysis": "pair = image x QP; n is counted by images",
        "path_control": control_row, "rows": rows, "determinism": determinism, "summary": summary,
    }
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        json.dump(record, f, ensure_ascii=False, indent=1, allow_nan=False)
    build_book(run_path, book_path)
    s = summary
    print(f"\nV2-a: {s['n_v2a_match']} of {s['n_pairs']} pairs bit-identical ({s['n_images']} images) -> "
          f"{s['verdict_v2a']}; determinism {s['n_determinism_equal']} of {s['n_determinism']} -> "
          f"{s['verdict_determinism']}; refutation fired: {s['refutation_fired']}")
    print(f"Expected, no verdict: bits non-increasing with QP on {s['n_images_bits_non_increasing']} of "
          f"{s['n_images_with_series']} images; |log PSNR-Y - ours| max {s['max_abs_log_minus_ours_db']!r} dB, "
          f"median {s['median_abs_log_minus_ours_db']!r} dB over {s['n_log_psnr_compared']} pairs; "
          f"PSNR B02 = +inf on {s['n_infinite_psnr_b02']} pairs")
    print(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}")
    return run_path, book_path


TABLE = [  # (header, run-file key)
    ("№ пары", "pair_index"), ("Изображение (код)", "image_code"), ("Класс", "image_class"), ("QP", "qp"),
    ("H, пикс", "h"), ("W, пикс", "w"), ("Статус", "status"), ("Причина", "reason"), ("Поток, бит", "bits"),
    ("bpp, бит/пикс", "bpp"), ("CR, безразм.", "cr"), ("V2-a: recon = encoder_recon", "v2a_match"),
    ("Форма recon = форма входа", "shape_ok"), ("PSNR B02 по recon, дБ", "psnr_b02"),
    ("PSNR-Y лога, дБ", "log_psnr_y"), ("Лог − наш, дБ", "log_minus_ours"), ("sha256 потока", "bitstream_sha256"),
    ("sha256 recon", "recon_sha256"), ("Level", "level"), ("Статус Level", "level_status"),
    ("FrameRate", "frame_rate"), ("Статус FrameRate", "frame_rate_status"), ("Время, с", "seconds"),
]
YES_NO = {True: "да", False: "нет"}


def _cell(value):
    return YES_NO[value] if isinstance(value, bool) else value


def book_cells(record, run_sha256):
    cells = {}
    s = record["summary"]
    d = "Описание"
    lines = [
        ("Команда", record["command"]), ("Вопрос", record["question"]),
        ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
        ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
        ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
        *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
        ("Level", f"{record['level']['value']} ({record['level']['status']}): {record['level']['basis']}"),
        ("FrameRate", f"{record['frame_rate']['value']} ({record['frame_rate']['status']}): "
                      f"{record['frame_rate']['basis']}"),
        ("Кодер", record["build_manifest"]["encoder_version"]), ("Компилятор", record["build_manifest"]["compiler"]),
        ("max_pixels", record["max_pixels_rule"]),
        ("Пар, шт", s["n_pairs"]), ("Изображений, шт", s["n_images"]), ("Пар со значением, шт", s["n_valid"]),
        ("Пар с совпадением V2-a, шт", s["n_v2a_match"]), ("Вердикт V2-a", s["verdict_v2a"]),
        ("Повторных кодирований, шт", s["n_determinism"]), ("Из них поток совпал, шт", s["n_determinism_equal"]),
        ("Вердикт детерминизма", s["verdict_determinism"]), ("Опровержение сработало", YES_NO[s["refutation_fired"]]),
        ("Ожидаемо: изображений с невозрастающими битами, шт", s["n_images_bits_non_increasing"]),
        ("из изображений с рядом QP, шт", s["n_images_with_series"]),
        ("Ожидаемо: наибольшее |лог − наш|, дБ", s["max_abs_log_minus_ours_db"]),
        ("Ожидаемо: медиана |лог − наш|, дБ", s["median_abs_log_minus_ours_db"]),
        ("Пар с PSNR B02 = +inf, шт", s["n_infinite_psnr_b02"]),
        ("Точность чисел книги", "16 значащих цифр (openpyxl пишет %.16g); каждая клетка сверена с числом файла "
                                 "прогона, округлённым до 16 цифр; полная точность — в файле прогона"),
        ("Листы «Рисунки» и «Страницы»", "не строятся: страницы не применяются (объявлено до чисел)"),
        ("На пальцах", "Каждое изображение сжато кодером VVC и разжато его же декодером. Мы проверили, что картинка "
                       "из декодера в точности та же, что кодер считал у себя, и что повторное сжатие даёт тот же файл."),
    ]
    for i, (k, v) in enumerate(lines, start=1):
        cells[d, i, 1], cells[d, i, 2] = k, v
    t = "Таблица"
    for j, (head, key) in enumerate(TABLE, start=1):
        cells[t, 1, j], cells[t, 2, j] = head, f"rows[].{key}"
    for i, row in enumerate(record["rows"]):
        for j, (_, key) in enumerate(TABLE, start=1):
            cells[t, 3 + i, j] = _cell(row[key])
    g = "График"
    codes = list(dict.fromkeys(r["image_code"] for r in record["rows"] if r["image_class"] != "synthetic"))
    qps = sorted({r["qp"] for r in record["rows"] if r["image_code"] in codes})
    cells[g, 1, 1] = "Поток, бит, по QP для каждого изображения; логарифмическая ось. Ключ: rows[].bits"
    cells[g, 2, 1] = "QP"
    for j, code in enumerate(codes, start=2):
        cells[g, 2, j] = code
    for i, qp in enumerate(qps):
        cells[g, 3 + i, 1] = qp
        for j, code in enumerate(codes, start=2):
            cells[g, 3 + i, j] = next((r["bits"] for r in record["rows"]
                                       if (r["image_code"], r["qp"]) == (code, qp) and r["status"] == "valid"), None)
    return cells


def build_book(run_path, book_path):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    record = json.loads(run_bytes.decode("utf-8"))
    cells = book_cells(record, hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Таблица", "График"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    wb["Описание"].column_dimensions["A"].width = 44
    wb["Описание"].column_dimensions["B"].width = 120
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for c in wb["Таблица"][1]:
        c.font = Font(bold=True)
    g = wb["График"]
    n_rows = max(r for s, r, _ in cells if s == "График")
    n_cols = max(c for s, _, c in cells if s == "График")
    if n_rows >= 3 and n_cols >= 2:
        chart = LineChart()
        chart.title = "Поток, бит, по QP"
        chart.y_axis.title, chart.x_axis.title = "бит", "QP"
        chart.y_axis.scaling.logBase = 10
        chart.x_axis.delete = chart.y_axis.delete = False
        chart.add_data(Reference(g, min_col=2, max_col=n_cols, min_row=2, max_row=n_rows), titles_from_data=True)
        chart.set_categories(Reference(g, min_col=1, min_row=3, max_row=n_rows))
        chart.width, chart.height = 28, 13
        g.add_chart(chart, "K2")
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g, as openpyxl writes them); no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        images = images_of_the_cycle(paths["images_root"])
        plan = [(img["code"], qp) for img in images if img["class"] != "synthetic" for qp in QPS]
        plan += [(f"syn{h}x{w}", 32) for h, w in ODD_SHAPES]
        run(paths, images, plan, [("e0", qp) for qp in QPS], out_dir, state)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
