"""V4 AVIF on B12: does AVIF through libheif (pillow-heif 0.22.0) keep the OOP of noisy Frisco inside the QF range and within 6 of RS24? Run: python studies/P001_vvc-dental/run_v4_avif.py

Study P001, cycle 9: the repetition of check V4 for AVIF of STUDY.md section 8 with the wrapper B12
(../REFERENCE.md#b12) that replaces B08 for P001. Order, each step a stop if it fails: (1) threads - Frisco,
sigma^2 = 100, seed 1, QF 35 and 80 coded with 1, 2, 4, 8 aom threads: 2, 4 and 8 must give one file;
(2) regression - every Frisco row (1515) equals variant (d) of cycles/8/avif_diag.json by file_sha256 and
psnr_tc_db; (3) the criterion of section 8 on Frisco; Diego is printed without a criterion. Conditions as in
cycle 7: Frisco and Diego, sigma^2 = 25, 100, 196, seeds 1..5 (noise sha256 = cycle 7), QF 0..100, OOP rule
of section 8. Writes <cache_root>/cycles/9/v4_avif.json (refuses by name if it exists) and the book
v4_avif.xlsx beside it; `--rebuild-book` rebuilds the book from the run file.
"""

import argparse
import datetime
import hashlib
import json
import math
import platform
import statistics
import sys
import tempfile
import time
from pathlib import Path

import numpy as np

import imstudies
from imstudies.core.metrics import Refusal, psnr
from imstudies.lab.coders import avif_libheif as b12
from imstudies.lab.coders import isobmff
from imstudies.lab.coders.process import run_tool
from imstudies.lab.coders.still import check_versions
from imstudies.lab.images import load_test_images
from imstudies.lab.runs import RunRefusal, git_state, json_number, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
COMMAND = "python studies/P001_vvc-dental/run_v4_avif.py"
QUESTION = ("V4 AVIF на B12: лежит ли наилучшая точка AVIF через libheif на зашумлённом Frisco внутри диапазона QF "
            "и не дальше 6 от RS24 при σ² = 25, 100, 196?")
CYCLE, RUN_FILE, BOOK_FILE = 9, "v4_avif.json", "v4_avif.xlsx"
TEST_IMAGES = ["frisco.bmp", "DIEGO.BMP"]
FULL_PLAN = {"sigma2": (25, 100, 196), "seeds": (1, 2, 3, 4, 5), "qualities": tuple(range(0, 101))}
THREADS_CHECK = {"image": "frisco", "sigma2": 100, "seed": 1, "qualities": (35, 80), "threads": (1, 2, 4, 8)}
TOL_QF = 6
TIMEOUT_S = 3600
REFERENCE_SHA256 = "88f7661b38309874accdf8bbf9d5f720596d341e9594c070076eadeb198a85e0"  # v3b_v4_reference.json
CYCLE7_SHA256 = "5718f5e8615c77ab64db41a3bf0a7416bc4da21339e3abfd9c359aa3f76d6b14"  # cycles/7/v3b_v4.json
CYCLE8_SHA256 = "cbb24be4260c0efe72b77d0f8d82757e240e3e6232c1c541681e3bc3cfdb76b3"  # cycles/8/avif_diag.json
DECLARED = {
    "data": "Frisco и Diego из Images/BMP (чтение цикла 1), 512 × 512; тестовые изображения, в P001 не входят.",
    "noise": "Шум цикла 7 (STUDY.md §8): σ² = 25, 100, 196, зёрна 1…5, float64, rint, отсечение; sha256 зашумлённых "
             "сверяется с файлом цикла 7.",
    "codec": f"B12: pillow-heif 0.22.0 → libheif 1.19.7, aom 3.12.0, enc_params {b12.ENC_PARAMS}, QF 0…100, шаг 1; "
             "версии окружения сверяются в каждом вызове.",
    "threads": "Шаг 1 — потоки: Frisco, σ² = 100, зерно 1, QF 35 и 80, threads 1, 2, 4, 8 (остальные параметры B12); "
               "печатаются sha256 и размер; вердикт — 2, 4, 8 дают один файл (1 — назвать); иначе стоп.",
    "regression": "Шаг 2 — регрессия: каждая строка Frisco (1515) = строке варианта (г) cycles/8/avif_diag.json по "
                  "file_sha256 и psnr_tc_db (sha256 файла цикла 8 сверяется до сравнения); не совпало — стоп, V4 не "
                  "считается.",
    "criterion": f"Шаг 3 — V4 AVIF (STUDY.md §8, объявлено 2026-10-20, 05:32 UTC): на Frisco при σ² = 25, 100, 196 "
                 f"медиана QF_OOP по зёрнам строго внутри 0…100 и |медиана QF_OOP − RS24| ≤ {TOL_QF}; PSNR_max и CR в OOP "
                 f"— без критерия, рядом с RS24; Diego — без критерия.",
    "oop": "OOP — QF с наибольшим PSNR_tc среди строк valid, при равенстве — меньший QF (§8); медиана, наименьшее и "
           "наибольшее по пяти зёрнам.",
    "cr": "CR по файлу = W·H / байт файла; CR по mdat = W·H / байт полезной нагрузки mdat — диагностика.",
    "pages": "Страницы для глаз не применяются — численная сверка.",
}


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {}
    for key in ("images_root", "cache_root", "v3b_v4_reference", "avif_libheif_python"):
        if not raw.get(key) or raw[key].startswith("<"):
            raise RunRefusal(f"lab_paths.json has no key '{key}' or it is the example placeholder")
        paths[key] = (ROOT / raw[key]).resolve()
    if not paths["avif_libheif_python"].is_file():
        raise RunRefusal(f"the python of the B12 environment is not found: {paths['avif_libheif_python']}")
    return paths


def load_pinned(path, sha, what):
    found = sha256_file(path)
    if found != sha:
        raise RunRefusal(f"{what} {Path(path).name}: sha256 {found} differs from the pinned {sha}")
    return json.loads(Path(path).read_text(encoding="utf-8"))


def noisy_image(true, sigma2, seed):
    """The noise of cycle 7: float64 sum, rint, clip 0..255, uint8."""
    noise = np.random.default_rng(seed).normal(0, math.sqrt(sigma2), true.shape)
    return np.clip(np.rint(true.astype(np.float64) + noise), 0, 255).astype(np.uint8)


def oop(points):
    """(qf, psnr_tc) pairs -> the qf of the largest psnr_tc, the smaller qf on a tie; None if empty."""
    return min(points, key=lambda p: (-p[1], p[0]))[0] if points else None


def code_with_threads(image, qualities, threads, python, work_root):
    """Diagnosis of step 1: the B12 executor with enc_params of B12 except threads (B12 itself fixes 4)."""
    with tempfile.TemporaryDirectory(dir=work_root, prefix="threads_") as work:
        work = Path(work)
        np.save(work / "image.npy", image)
        request = {"qualities": list(qualities), "enc_params": b12.ENC_PARAMS | {"threads": str(threads)},
                   "options": b12.OPTIONS}
        (work / "request.json").write_text(json.dumps(request), encoding="utf-8")
        run_tool([str(python), "-I", str(b12.EXEC), str(work)], TIMEOUT_S)
        result = json.loads((work / "result.json").read_text(encoding="utf-8"))
        check_versions("AVIF (libheif)", b12.EXPECTED_VERSIONS, result["versions"])
        return {q: (work / f"q{q:03d}.avif").read_bytes() for q in qualities}


def threads_step(image, python, work_root):
    t = THREADS_CHECK
    files = {n: code_with_threads(image, t["qualities"], n, python, work_root) for n in t["threads"]}
    rows = [{"threads": n, "qf": q, "file_bytes": len(files[n][q]), "file_sha256": hashlib.sha256(files[n][q]).hexdigest()}
            for n in t["threads"] for q in t["qualities"]]
    same_2_4_8 = all(files[2][q] == files[4][q] == files[8][q] for q in t["qualities"])
    one_differs = {q: files[1][q] != files[4][q] for q in t["qualities"]}
    return {"rows": rows, "same_2_4_8": same_2_4_8, "threads_1_differs_from_4": one_differs,
            "verdict": "pass" if same_2_4_8 else "stop"}


def make_row(image_code, s2, seed, true, noisy, coded):
    data, recon = coded.file, coded.recon
    h, w = true.shape
    mdat = sum(len(p) for path, _, p, _ in isobmff.boxes(data) if path == "mdat")
    return {"image": image_code, "sigma2": s2, "seed": seed, "qf": coded.provenance["encode"]["quality"],
            "status": "valid", "reason": None, "file_bytes": len(data), "mdat_payload_bytes": mdat,
            "cr_file": h * w / len(data), "cr_mdat": h * w / mdat,
            "psnr_tc_db": psnr(true, recon, max_pixels=true.size).value,
            "psnr_nc_db": psnr(noisy, recon, max_pixels=true.size).value,
            "file_sha256": hashlib.sha256(data).hexdigest(), "av1c_monochrome": coded.provenance["av1C_monochrome"]}


def code_image(img, plan, noise7, python, work_root):
    rows, provenance = [], None
    true = img["array"]
    for s2 in plan["sigma2"]:
        for seed in plan["seeds"]:
            noisy = noisy_image(true, s2, seed)
            if hashlib.sha256(noisy.tobytes()).hexdigest() != noise7.get((img["code"], s2, seed)):
                raise RunRefusal(f"noisy {img['code']} s2 {s2} seed {seed} differs from cycle 7 (sha256)")
            try:
                coded = b12.encode_decode(noisy, list(plan["qualities"]), max_pixels=true.size, timeout_s=TIMEOUT_S,
                                          python=python, work_root=work_root)
            except Refusal as refusal:
                rows += [{"image": img["code"], "sigma2": s2, "seed": seed, "qf": q, "status": "no_result",
                          "reason": refusal.reason} for q in plan["qualities"]]
                continue
            except (RuntimeError, OSError, TimeoutError) as exc:
                if "versions differ" in str(exc):
                    raise RunRefusal(str(exc)) from exc
                rows += [{"image": img["code"], "sigma2": s2, "seed": seed, "qf": q, "status": "error",
                          "reason": "internal_error", "detail": str(exc)[:600]} for q in plan["qualities"]]
                continue
            provenance = provenance or coded[0].provenance
            rows += [make_row(img["code"], s2, seed, true, noisy, c) for c in coded]
    return rows, provenance


def regression(rows, cycle8):
    ref = {(r["sigma2"], r["seed"], r["qf"]): r for r in cycle8["rows"] if r["variant"] == "d"}
    differ = []
    for r in rows:
        c = ref.get((r["sigma2"], r["seed"], r["qf"]))
        if c is None or r.get("file_sha256") != c["file_sha256"] or json_number(r.get("psnr_tc_db")) != c["psnr_tc_db"]:
            differ.append(f"s2 {r['sigma2']} seed {r['seed']} QF {r['qf']}")
    return {"rows_compared": len(rows), "rows_equal": len(rows) - len(differ), "differ_count": len(differ),
            "differ_first": differ[:20], "verdict": "pass" if rows and not differ else "stop"}


def curves_of(rows):
    groups = {}
    for r in rows:
        groups.setdefault((r["image"], r["sigma2"], r["seed"]), []).append(r)
    out = []
    for (image, s2, seed), rs in sorted(groups.items()):
        valid = [r for r in rs if r["status"] == "valid"]
        q = oop([(r["qf"], r["psnr_tc_db"]) for r in valid])
        at = next((r for r in valid if r["qf"] == q), {})
        out.append({"image": image, "sigma2": s2, "seed": seed, "n_valid": len(valid), "n_not_valid": len(rs) - len(valid),
                    "qf_oop": q, "psnr_max_db": at.get("psnr_tc_db"), "cr_file_oop": at.get("cr_file"),
                    "cr_mdat_oop": at.get("cr_mdat"), "oop_inside_range": None if q is None else 0 < q < 100})
    return out


def summary_of(curves, rs24):
    out = []
    for image, s2 in sorted({(c["image"], c["sigma2"]) for c in curves}):
        cs = [c for c in curves if (c["image"], c["sigma2"]) == (image, s2)]
        entry = {"image": image, "sigma2": s2, "n_seeds": len(cs),
                 "n_oop_inside_range": sum(bool(c["oop_inside_range"]) for c in cs)}
        for name in ("qf_oop", "psnr_max_db", "cr_file_oop", "cr_mdat_oop"):
            values = [c[name] for c in cs if c[name] is not None]
            entry |= {f"{name}_median": statistics.median(values) if values else None,
                      f"{name}_min": min(values, default=None), f"{name}_max": max(values, default=None)}
        ref = rs24.get((image, s2))
        med = entry["qf_oop_median"]
        entry |= {"rs24_qf_oop": ref and ref["qf_oop"], "rs24_psnr_max_db": ref and ref["psnr_max_db"],
                  "rs24_cr": ref and ref["cr"], "rs24_note": ref and ref.get("note"),
                  "d_qf_oop": None if not ref or med is None else med - ref["qf_oop"],
                  "d_psnr_max_db": None if not ref or entry["psnr_max_db_median"] is None
                  else entry["psnr_max_db_median"] - ref["psnr_max_db"],
                  "cr_rel": None if not ref or entry["cr_file_oop_median"] is None
                  else entry["cr_file_oop_median"] / ref["cr"] - 1}
        criterion = image == "frisco"
        entry["criterion"] = "V4 AVIF" if criterion else "—"
        if criterion and med is not None and ref:
            entry["inside_range"] = 0 < med < 100
            entry["qf_within_tolerance"] = abs(med - ref["qf_oop"]) <= TOL_QF
            entry["in_tolerance"] = entry["inside_range"] and entry["qf_within_tolerance"]
        else:
            entry |= {"inside_range": None, "qf_within_tolerance": None, "in_tolerance": None}
        out.append(entry)
    return out


def verdict_of(summary, sigma2s):
    rows = [s for s in summary if s["criterion"] == "V4 AVIF"]
    have = {s["sigma2"] for s in rows if s["in_tolerance"] is not None}
    failed = [f"frisco s2 {s['sigma2']}" for s in rows if s["in_tolerance"] is False]
    missing = [s2 for s2 in (25, 100, 196) if s2 not in have]
    verdict = "fail" if failed else ("incomplete" if missing else "pass")
    return {"rows_in_tolerance": sum(s["in_tolerance"] is True for s in rows), "rows": 3, "failed": failed,
            "missing_sigma2": missing, "verdict": verdict}


def median_curves(rows):
    groups = {}
    for r in rows:
        if r["status"] == "valid":
            groups.setdefault((r["image"], r["sigma2"], r["qf"]), []).append(r["psnr_tc_db"])
    return [{"image": i, "sigma2": s, "qf": q, "n_seeds": len(p), "psnr_tc_median_db": statistics.median(p)}
            for (i, s, q), p in sorted(groups.items())]


def _jsonable(obj):
    if isinstance(obj, float):
        return json_number(obj)
    if isinstance(obj, dict):
        return {k: _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    return obj


def run(paths, images, plan, out_dir, state):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path = out_dir / RUN_FILE, out_dir / BOOK_FILE
    if run_path.exists():
        raise RunRefusal(f"{run_path} already exists: a run file of cycle {CYCLE} is never overwritten")
    reference = load_pinned(paths["v3b_v4_reference"], REFERENCE_SHA256, "V3-b/V4 reference")
    rs24 = {(r["image"], r["sigma2"]): r for r in reference["rs24_oop_by_psnr"] if r["codec"] == "avif"}
    cycle7 = load_pinned(paths["cache_root"] / "cycles" / "7" / "v3b_v4.json", CYCLE7_SHA256, "cycle 7 run file")
    noise7 = {(n["image"], n["sigma2"], n["seed"]): n["noisy_sha256"] for n in cycle7["noise"]}
    cycle8 = load_pinned(paths["cache_root"] / "cycles" / "8" / "avif_diag.json", CYCLE8_SHA256, "cycle 8 run file")
    python, work_root = paths["avif_libheif_python"], paths["cache_root"] / "work_b12"
    work_root.mkdir(parents=True, exist_ok=True)
    started, t0 = datetime.datetime.now(datetime.timezone.utc), time.perf_counter()
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")
    by_code = {img["code"]: img for img in images}
    steps, rows, provenance, stopped = {}, [], None, None
    t = time.perf_counter()
    frisco = by_code["frisco"]
    tc = THREADS_CHECK
    steps["threads"] = threads_step(noisy_image(frisco["array"], tc["sigma2"], tc["seed"]), python, work_root)
    steps["threads"]["seconds"] = round(time.perf_counter() - t, 1)
    print(f"  step 1, threads: 2 = 4 = 8 {steps['threads']['same_2_4_8']}; 1 differs from 4 "
          f"{steps['threads']['threads_1_differs_from_4']} -> {steps['threads']['verdict']}")
    if steps["threads"]["verdict"] != "pass":
        stopped = "step 1 (threads): 2, 4 and 8 threads do not give one file"
    if not stopped:
        t = time.perf_counter()
        frisco_rows, provenance = code_image(frisco, plan, noise7, python, work_root)
        steps["regression"] = regression(frisco_rows, cycle8) | {"seconds": round(time.perf_counter() - t, 1)}
        rows += frisco_rows
        print(f"  step 2, regression: {steps['regression']['rows_equal']} of {steps['regression']['rows_compared']} "
              f"Frisco rows equal cycle 8 variant (d) -> {steps['regression']['verdict']}")
        if steps["regression"]["verdict"] != "pass":
            stopped = (f"step 2 (regression): {steps['regression']['differ_count']} of "
                       f"{steps['regression']['rows_compared']} Frisco rows differ from cycle 8")
    if not stopped:
        for code, img in by_code.items():
            if code == "frisco":
                continue
            t = time.perf_counter()
            more, _ = code_image(img, plan, noise7, python, work_root)
            rows += more
            print(f"  {code}: {sum(r['status'] == 'valid' for r in more)}/{len(more)} valid, "
                  f"{time.perf_counter() - t:.1f} s")
    for k, r in enumerate(rows, start=1):
        r["row_index"] = k
    curves = curves_of(rows)
    summary = summary_of(curves, rs24)
    verdict = {"verdict": "not computed", "stopped_at": stopped} if stopped else verdict_of(summary, plan["sigma2"])
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "wall_seconds": round(time.perf_counter() - t0, 1), **state, "declared_before_numbers": DECLARED,
        "plan": {"sigma2": list(plan["sigma2"]), "seeds": list(plan["seeds"]),
                 "qualities": [min(plan["qualities"]), max(plan["qualities"]), len(plan["qualities"])],
                 "images": list(by_code)},
        "b12": {"enc_params": b12.ENC_PARAMS, "options": b12.OPTIONS, "expected_versions": b12.EXPECTED_VERSIONS,
                "internal": b12.INTERNAL, "provenance_of_first_call": provenance},
        "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__,
                        "numpy": np.__version__, "platform": platform.platform()},
        "reference": {"file": paths["v3b_v4_reference"].name, "sha256": REFERENCE_SHA256,
                      "rs24_avif": list(rs24.values())},
        "pinned": {"cycles/7/v3b_v4.json": CYCLE7_SHA256, "cycles/8/avif_diag.json": CYCLE8_SHA256},
        "images": [{k: v for k, v in img.items() if k != "array"}
                   | {"pixels_sha256": hashlib.sha256(img["array"].tobytes()).hexdigest()} for img in images],
        "unit_of_analysis": "row = image x sigma2 x seed x QF; summary per (image, sigma2) over seeds",
        "steps": steps, "stopped": stopped, "rows": rows, "curves": curves, "median_curves": median_curves(rows),
        "summary": summary, "verdict_v4_avif": verdict,
    }
    text = json.dumps(_jsonable(record), ensure_ascii=False, indent=1, allow_nan=False)
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        f.write(text)
    build_book(run_path, book_path)
    for s in summary:
        print(f"  {s['image']:<6} s2 {s['sigma2']:>3}: QF_OOP {s['qf_oop_median']} [{s['qf_oop_min']}; {s['qf_oop_max']}] "
              f"(RS24 {s['rs24_qf_oop']}), PSNR_max {s['psnr_max_db_median']} (RS24 {s['rs24_psnr_max_db']}), "
              f"CR {s['cr_file_oop_median']} (RS24 {s['rs24_cr']}) {s['criterion']} in tolerance {s['in_tolerance']}")
    print(f"V4 AVIF on B12: {verdict}\nRun file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}")
    return run_path, book_path


SUMMARY = [("Изображение", "image"), ("σ²", "sigma2"), ("Зёрен, шт", "n_seeds"), ("Критерий", "criterion")] + [
    (f"{label}: {stat}", f"{key}_{k}") for key, label in (("qf_oop", "QF_OOP, QF"), ("psnr_max_db", "PSNR_max, дБ"),
                                                          ("cr_file_oop", "CR по файлу в OOP, безразм."),
                                                          ("cr_mdat_oop", "CR по mdat в OOP, безразм."))
    for k, stat in (("median", "медиана"), ("min", "наименьшее"), ("max", "наибольшее"))] + [
    ("RS24 QF_OOP, QF", "rs24_qf_oop"), ("RS24 PSNR_max, дБ", "rs24_psnr_max_db"), ("RS24 CR, безразм.", "rs24_cr"),
    ("RS24 примечание", "rs24_note"), ("ΔQF_OOP (медиана − RS24), QF", "d_qf_oop"),
    ("ΔPSNR_max (медиана − RS24), дБ", "d_psnr_max_db"), ("CR: медиана / RS24 − 1, %", "cr_rel"),
    ("OOP внутри 0…100, зёрен", "n_oop_inside_range"), ("Медиана QF_OOP внутри 0…100", "inside_range"),
    ("|ΔQF_OOP| ≤ 6", "qf_within_tolerance"), ("В допуске", "in_tolerance")]
CURVES = [("Изображение", "image"), ("σ²", "sigma2"), ("Зерно", "seed"), ("Valid, шт", "n_valid"),
          ("Не valid, шт", "n_not_valid"), ("QF_OOP, QF", "qf_oop"), ("PSNR_max, дБ", "psnr_max_db"),
          ("CR по файлу в OOP, безразм.", "cr_file_oop"), ("CR по mdat в OOP, безразм.", "cr_mdat_oop"),
          ("OOP внутри 0…100", "oop_inside_range")]
ROWS = [("№", "row_index"), ("Изображение", "image"), ("σ²", "sigma2"), ("Зерно", "seed"), ("QF", "qf"),
        ("Статус", "status"), ("Причина", "reason"), ("Файл, байт", "file_bytes"),
        ("mdat (нагрузка), байт", "mdat_payload_bytes"), ("CR по файлу, безразм.", "cr_file"),
        ("CR по mdat, безразм.", "cr_mdat"), ("PSNR_tc, дБ", "psnr_tc_db"), ("PSNR_nc, дБ", "psnr_nc_db"),
        ("av1C monochrome", "av1c_monochrome"), ("sha256 файла", "file_sha256")]
THREADS = [("Потоков aom", "threads"), ("QF", "qf"), ("Файл, байт", "file_bytes"), ("sha256 файла", "file_sha256")]
VERDICT_WORDS = {True: "да", False: "нет", None: "без критерия"}


def _table(cells, formats, sheet, columns, records, percent=(), first_row=1):
    for j, (head, key) in enumerate(columns, start=1):
        cells[sheet, first_row, j], cells[sheet, first_row + 1, j] = head, key
    for i, rec in enumerate(records):
        for j, (_, key) in enumerate(columns, start=1):
            v = rec.get(key)
            cells[sheet, first_row + 2 + i, j] = VERDICT_WORDS[v] if key == "in_tolerance" else (
                ("да" if v else "нет") if isinstance(v, bool) else v)
            if key in percent:
                formats[sheet, first_row + 2 + i, j] = "0.0%"


def book_cells(record, run_sha256):
    cells, formats = {}, {}
    steps, verdict = record["steps"], record["verdict_v4_avif"]
    reg = steps.get("regression")
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
             ("Закончен, UTC", record["finished_utc"]), ("Время прогона, с", record["wall_seconds"]),
             ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
             *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
             ("Шаг 1, потоки", f"{steps['threads']['verdict']}: 2 = 4 = 8 — "
                               f"{'да' if steps['threads']['same_2_4_8'] else 'нет'}; 1 ≠ 4 по QF — "
                               f"{json.dumps(steps['threads']['threads_1_differs_from_4'])}"),
             ("Шаг 2, регрессия", "не выполнялся" if not reg else
              f"{reg['verdict']}: равны {reg['rows_equal']} из {reg['rows_compared']} строк Frisco циклу 8 (г)"),
             ("Остановлен", record["stopped"] or "нет"),
             ("Вердикт V4 AVIF на B12", json.dumps(verdict, ensure_ascii=False)),
             ("Версии B12", json.dumps(record["b12"]["expected_versions"], ensure_ascii=False)),
             ("Строк, шт", len(record["rows"])),
             ("Ключи листов", "под шапкой — имя поля: «Сводка» — summary[], «Потоки» — steps.threads.rows[], "
                              "«Зёрна» — curves[], «Строки» — rows[], «График» — median_curves[]"),
             ("Листы «Рисунки» и «Страницы»", "не строятся: страницы не применяются (объявлено до чисел)"),
             ("На пальцах", "Два тестовых аэроснимка с шумом сжаты в AVIF тем путём, которым шла группа, при каждом "
                            "значении качества; сначала проверено, что результат не зависит от числа потоков и "
                            "совпадает с прошлым циклом, затем найдена наилучшая точка и сравнена со статьёй.")]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v
    _table(cells, formats, "Сводка", SUMMARY, record["summary"], percent=("cr_rel",))
    _table(cells, formats, "Потоки", THREADS, steps["threads"]["rows"])
    _table(cells, formats, "Зёрна", CURVES, record["curves"])
    _table(cells, formats, "Строки", ROWS, record["rows"])
    series = sorted({(m["image"], m["sigma2"]) for m in record["median_curves"]})
    qs = sorted({m["qf"] for m in record["median_curves"]})
    med = {(m["image"], m["sigma2"], m["qf"]): m["psnr_tc_median_db"] for m in record["median_curves"]}
    cells["График", 1, 1], cells["График", 2, 1] = "median_curves[].qf", "QF"
    for j, (i, s) in enumerate(series, start=2):
        cells["График", 1, j], cells["График", 2, j] = "median_curves[].psnr_tc_median_db", f"{i} σ²={s}"
    for r, q in enumerate(qs, start=3):
        cells["График", r, 1] = q
        for j, (i, s) in enumerate(series, start=2):
            if (i, s, q) in med:
                cells["График", r, j] = med[i, s, q]
    return cells, formats, len(series)


def build_book(run_path, book_path):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    cells, formats, n_series = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Сводка", "Потоки", "Зёрна", "Строки", "График"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    for (sheet, r, c), fmt in formats.items():
        wb[sheet].cell(row=r, column=c).number_format = fmt
    wb["Описание"].column_dimensions["A"].width = 34
    wb["Описание"].column_dimensions["B"].width = 140
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for name in ("Сводка", "Потоки", "Зёрна", "Строки", "График"):
        for c in wb[name][2 if name == "График" else 1]:
            c.font = Font(bold=True)
        wb[name].freeze_panes = "A3"
    g = wb["График"]
    n_rows = max(r for s, r, _ in cells if s == "График")
    if n_rows >= 3 and n_series:
        chart = LineChart()
        chart.title = "AVIF через libheif (B12): PSNR_tc (медиана по зёрнам) по QF"
        chart.y_axis.title, chart.x_axis.title = "PSNR_tc, дБ", "QF"
        chart.x_axis.delete = chart.y_axis.delete = False
        chart.add_data(Reference(g, min_col=2, max_col=1 + n_series, min_row=2, max_row=n_rows), titles_from_data=True)
        chart.set_categories(Reference(g, min_col=1, min_row=3, max_row=n_rows))
        chart.width, chart.height = 28, 14
        g.add_chart(chart, "B" + str(n_rows + 3))
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g, as openpyxl writes them); no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())[0]
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        images = load_test_images(paths["images_root"], TEST_IMAGES, (512, 512))
        run(paths, images, FULL_PLAN, out_dir, state)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
