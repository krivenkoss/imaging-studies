"""AVIF gap (Q-AVIF-GAP): is PSNR_max of our AVIF on noisy Frisco below RS24 because of the encoder speed or of the libavif path? Run: python studies/P001_vvc-dental/run_avif_diag.py

Study P001, cycle 8: a diagnosis, not a check - no criterion, B08 is not changed. Frisco with AWGN
(sigma^2 = 25, 100, 196; seeds 1..5; the noise of cycle 7) is coded at AVIF quality 0..100 by four
variants: (a) B08 as is (regression: every file equals the row of cycles/7/v3b_v4.json by sha256),
(b) B08 with speed 4, (c) B08 with speed 2, (d) the group's path - AVIF through libheif by the last
pillow-heif that encodes AVIF, in a venv outside the project (key avif_libheif_python of
lab_paths.json), defaults. OOP by the rule of STUDY.md section 8; median and range over seeds of
QF_OOP, PSNR_max, CR at the OOP and the difference with RS24; encoding time. Writes
<cache_root>/cycles/8/avif_diag.json (refuses by name if it exists) and the book avif_diag.xlsx beside
it; `--rebuild-book` rebuilds the book from the run file.
"""

import argparse
import concurrent.futures
import datetime
import hashlib
import io
import json
import math
import os
import platform
import statistics
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import numpy as np
from PIL import Image

import imstudies
from imstudies.core.metrics import psnr
from imstudies.lab.coders import avif, isobmff
from imstudies.lab.coders.process import run_tool
from imstudies.lab.images import load_test_images
from imstudies.lab.runs import RunRefusal, git_state, json_number, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
COMMAND = "python studies/P001_vvc-dental/run_avif_diag.py"
QUESTION = ("Q-AVIF-GAP: почему PSNR_max нашего AVIF (B08) на зашумлённом Frisco на 0,85–1,15 дБ ниже RS24 — "
            "скорость кодера или путь libavif против libheif?")
CYCLE, RUN_FILE, BOOK_FILE = 8, "avif_diag.json", "avif_diag.xlsx"
IMAGE = "frisco.bmp"
FULL_PLAN = {"sigma2": (25, 100, 196), "seeds": (1, 2, 3, 4, 5), "qualities": tuple(range(0, 101)),
             "variants": ("a", "b", "c", "d")}
REFERENCE_SHA256 = "88f7661b38309874accdf8bbf9d5f720596d341e9594c070076eadeb198a85e0"  # v3b_v4_reference.json
CYCLE7_SHA256 = "5718f5e8615c77ab64db41a3bf0a7416bc4da21339e3abfd9c359aa3f76d6b14"  # cycles/7/v3b_v4.json
PILLOW_VARIANTS = {"a": dict(avif.ENCODE_ARGS), "b": avif.ENCODE_ARGS | {"speed": 4}, "c": avif.ENCODE_ARGS | {"speed": 2}}
VARIANT_TEXT = {"a": "B08 как есть (Pillow → libavif, aom, speed 6, max_threads 1)",
                "b": "B08 со speed 4", "c": "B08 со speed 2",
                "d": "путь группы: AVIF через libheif (pillow-heif 0.22.0, последняя версия с AVIF), умолчания"}
LIBHEIF_DEFAULTS = ("libheif 1.19.7 plugins/encoder_aom.cc (тег v1.19.7, прочитан 2026-10-20): speed 6; threads = число "
                    "ядер машины; usage all-intra; tune ssim; chroma 420 по умолчанию, но изображение L пишется "
                    "монохромным (av1C monochrome проверяется в строке); quality q → cq_level = ((100 − q)·63 + 50)/100, "
                    "AOM_Q")
DECLARED = {
    "data": "Frisco из Images/BMP (чтение цикла 1), 512 × 512; тестовое изображение, в P001 не входит.",
    "noise": "Шум цикла 7 (STUDY.md §8): σ² = 25, 100, 196, зёрна 1…5, float64, rint, отсечение; sha256 зашумлённых "
             "сверяется с файлом цикла 7.",
    "variants": "; ".join(f"({k}) {v}" for k, v in VARIANT_TEXT.items()) + ". B08 не меняется.",
    "oop": "OOP — QF с наибольшим PSNR_tc среди строк valid, при равенстве — меньший QF (STUDY.md §8); медиана, "
           "наименьшее и наибольшее по пяти зёрнам; разница с RS24 (табл. 1) — медиана минус эталон, CR — медиана / "
           "эталон − 1.",
    "regression": "Вариант (а): sha256 каждого файла, PSNR_tc и PSNR_nc = строке cycles/7/v3b_v4.json (sha256 файла "
                  "цикла 7 сверяется до сравнения); любое расхождение — регрессия не пройдена.",
    "timing": "Время кодирования — секунды одного вызова кодера (без декодирования), сумма и медиана на вызов; (а)–(в) "
              "кодируются в пуле процессов, каждый вызов с max_threads = 1, поэтому файл от пула не зависит (это и "
              "проверяет регрессия); (г) — один процесс, потоки libheif по умолчанию; время (г) с (а)–(в) не сравнимо.",
    "criterion": "Критерия нет — диагностика; вывод о причине разрыва делает владелец.",
    "pages": "Страницы для глаз не применяются — численный разбор.",
}
HELPER = r'''
import io, json, sys, time
from pathlib import Path
import numpy as np, PIL, pillow_heif
from PIL import Image
work, qualities = Path(sys.argv[1]), [int(q) for q in sys.argv[2].split(",")]
out = {"pillow_heif": pillow_heif.__version__, "pillow": PIL.__version__, "numpy": np.__version__,
       "libheif_info": pillow_heif.libheif_info(), "python": sys.version.split()[0],
       "options": {k: repr(getattr(pillow_heif.options, k)) for k in dir(pillow_heif.options) if k.isupper()}, "items": []}
for npy in sorted(work.glob("noisy_*.npy")):
    img = np.load(npy)
    for q in qualities:
        t = time.perf_counter()
        buf = io.BytesIO()
        pillow_heif.from_pillow(Image.fromarray(img, "L")).save(buf, format="AVIF", quality=q)
        seconds = time.perf_counter() - t
        data = buf.getvalue()
        heif = pillow_heif.open_heif(io.BytesIO(data), convert_hdr_to_8bit=False)
        stem = f"{npy.stem}_q{q:03d}"
        (work / f"{stem}.avif").write_bytes(data)
        np.save(work / f"{stem}_recon.npy", np.asarray(heif[0]))
        out["items"].append({"stem": stem, "noisy": npy.stem, "q": q, "seconds": seconds, "mode": heif[0].mode})
(work / "result.json").write_text(json.dumps(out), encoding="utf-8")
'''


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {}
    for key in ("images_root", "cache_root", "v3b_v4_reference"):
        if not raw.get(key):
            raise RunRefusal(f"lab_paths.json has no key '{key}'")
        paths[key] = (ROOT / raw[key]).resolve()
    value = raw.get("avif_libheif_python")
    paths["avif_libheif_python"] = None if not value or value.startswith("<") else (ROOT / value).resolve()
    return paths


def load_pinned(path, sha, what):
    found = sha256_file(path)
    if found != sha:
        raise RunRefusal(f"{what} {Path(path).name}: sha256 {found} differs from the pinned {sha}")
    return json.loads(Path(path).read_text(encoding="utf-8"))


def noisy_image(true, sigma2, seed):
    """The noise of cycle 7: float64 sum, rint, clip 0..255, uint8."""
    noise = np.random.default_rng(seed).normal(0, math.sqrt(sigma2), true.shape)
    return np.clip(np.rint(true.astype(np.float64) + noise), 0, 255).astype(np.uint8)


def encode_pillow(task):
    """One AVIF encode and decode through Pillow (libavif) - a top-level function for the process pool."""
    noisy_bytes, shape, q, args = task
    image = np.frombuffer(noisy_bytes, np.uint8).reshape(shape)
    started = time.perf_counter()
    buf = io.BytesIO()
    Image.fromarray(image, "L").save(buf, "AVIF", quality=q, **args)
    seconds = time.perf_counter() - started
    data = buf.getvalue()
    with Image.open(io.BytesIO(data)) as im:
        mode = im.mode
        recon = np.asarray(im).tobytes() if mode == "L" else None
    return data, recon, mode, seconds


def oop(points):
    """(qf, psnr_tc) pairs -> the qf of the largest psnr_tc, the smaller qf on a tie; None if empty."""
    return min(points, key=lambda p: (-p[1], p[0]))[0] if points else None


def make_row(variant, s2, seed, q, true, noisy, data, recon, mode, seconds):
    row = {"variant": variant, "image": "frisco", "sigma2": s2, "seed": seed, "qf": q, "encode_seconds": seconds}
    mono = isobmff.av1_monochrome(data) if data else None
    if mode != "L" or recon is None or recon.shape != true.shape or mono != 1:
        return row | {"status": "error", "reason": "internal_error", "av1c_monochrome": mono,
                      "detail": f"decoded mode {mode}, av1C monochrome {mono}"}
    mdat = sum(len(p) for path, _, p, _ in isobmff.boxes(data) if path == "mdat")
    h, w = true.shape
    return row | {"status": "valid", "reason": None, "av1c_monochrome": mono, "file_bytes": len(data),
                  "mdat_payload_bytes": mdat, "cr_file": h * w / len(data), "cr_mdat": h * w / mdat,
                  "psnr_tc_db": psnr(true, recon, max_pixels=true.size).value,
                  "psnr_nc_db": psnr(noisy, recon, max_pixels=true.size).value,
                  "file_sha256": hashlib.sha256(data).hexdigest()}


def code_pillow(variant, true, noisies, qualities, workers):
    tasks = [(noisy.tobytes(), noisy.shape, q, PILLOW_VARIANTS[variant]) for (_, _, noisy) in noisies for q in qualities]
    keys = [(s2, seed, noisy) for (s2, seed, noisy) in noisies for _ in qualities]
    if workers > 1:
        with concurrent.futures.ProcessPoolExecutor(max_workers=workers) as pool:
            results = list(pool.map(encode_pillow, tasks, chunksize=4))
    else:
        results = [encode_pillow(t) for t in tasks]
    rows = []
    for (s2, seed, noisy), task, (data, recon, mode, seconds) in zip(keys, tasks, results):
        recon = None if recon is None else np.frombuffer(recon, np.uint8).reshape(true.shape)
        rows.append(make_row(variant, s2, seed, task[2], true, noisy, data, recon, mode, seconds))
    return rows


def code_libheif(python, true, noisies, qualities, work_root):
    """Variant (d) in the venv outside the project: every noisy image through libheif; files and decodings back."""
    with tempfile.TemporaryDirectory(dir=work_root) as work:
        work = Path(work)
        for s2, seed, noisy in noisies:
            np.save(work / f"noisy_s{s2:03d}_seed{seed}.npy", noisy)
        run_tool([str(python), "-W", "ignore", "-c", HELPER, str(work), ",".join(map(str, qualities))], 7200)
        result = json.loads((work / "result.json").read_text(encoding="utf-8"))
        by_stem = {(s2, seed): noisy for s2, seed, noisy in noisies}
        rows = []
        for item in result["items"]:
            s2, seed = int(item["noisy"].split("_")[1][1:]), int(item["noisy"].split("seed")[1])
            data = (work / f"{item['stem']}.avif").read_bytes()
            recon = np.load(work / f"{item['stem']}_recon.npy")
            recon = recon if recon.dtype == np.uint8 else None
            rows.append(make_row("d", s2, seed, item["q"], true, by_stem[s2, seed], data, recon, item["mode"],
                                 item["seconds"]))
        env = {k: v for k, v in result.items() if k != "items"}
    return rows, env


def regression(rows_a, cycle7):
    """Variant (a) against cycle 7: file sha256, PSNR_tc and PSNR_nc of every row (JSON values compared)."""
    ref = {(r["sigma2"], r["seed"], r["qf"]): r for r in cycle7["rows"]
           if r["image"] == "frisco" and r["codec"] == "avif"}
    differ = []
    for r in rows_a:
        c = ref.get((r["sigma2"], r["seed"], r["qf"]))
        same = c is not None and r.get("file_sha256") == c["file_sha256"] and all(
            json_number(r.get(k)) == c[k] for k in ("psnr_tc_db", "psnr_nc_db"))
        if not same:
            differ.append(f"s2 {r['sigma2']} seed {r['seed']} QF {r['qf']}")
    return {"rows_compared": len(rows_a), "rows_equal": len(rows_a) - len(differ), "differ_first": differ[:10],
            "verdict": "pass" if rows_a and not differ else "fail"}


def curves_of(rows):
    groups = {}
    for r in rows:
        groups.setdefault((r["variant"], r["sigma2"], r["seed"]), []).append(r)
    out = []
    for (variant, s2, seed), rs in sorted(groups.items()):
        valid = [r for r in rs if r["status"] == "valid"]
        q = oop([(r["qf"], r["psnr_tc_db"]) for r in valid])
        at = next((r for r in valid if r["qf"] == q), {})
        out.append({"variant": variant, "sigma2": s2, "seed": seed, "n_valid": len(valid), "n_not_valid": len(rs) - len(valid),
                    "qf_oop": q, "psnr_max_db": at.get("psnr_tc_db"), "cr_file_oop": at.get("cr_file"),
                    "oop_inside_range": None if q is None else 0 < q < 100})
    return out


def summary_of(curves, rows, rs24, walls):
    out = []
    for variant, s2 in sorted({(c["variant"], c["sigma2"]) for c in curves}):
        cs = [c for c in curves if (c["variant"], c["sigma2"]) == (variant, s2)]
        secs = [r["encode_seconds"] for r in rows if (r["variant"], r["sigma2"]) == (variant, s2)]
        ref = rs24[s2]
        entry = {"variant": variant, "sigma2": s2, "n_seeds": len(cs),
                 "rs24_qf_oop": ref["qf_oop"], "rs24_psnr_max_db": ref["psnr_max_db"], "rs24_cr": ref["cr"]}
        for name in ("qf_oop", "psnr_max_db", "cr_file_oop"):
            values = [c[name] for c in cs if c[name] is not None]
            entry |= {f"{name}_median": statistics.median(values) if values else None,
                      f"{name}_min": min(values, default=None), f"{name}_max": max(values, default=None)}
        med = entry
        entry |= {"d_qf_oop": None if med["qf_oop_median"] is None else med["qf_oop_median"] - ref["qf_oop"],
                  "d_psnr_max_db": None if med["psnr_max_db_median"] is None else med["psnr_max_db_median"] - ref["psnr_max_db"],
                  "cr_rel": None if med["cr_file_oop_median"] is None else med["cr_file_oop_median"] / ref["cr"] - 1,
                  "n_oop_inside_range": sum(bool(c["oop_inside_range"]) for c in cs),
                  "encode_seconds_total": sum(secs), "encode_seconds_median": statistics.median(secs) if secs else None,
                  "variant_wall_seconds": walls.get(variant)}
        out.append(entry)
    return out


def median_curves(rows):
    groups = {}
    for r in rows:
        if r["status"] == "valid":
            groups.setdefault((r["variant"], r["sigma2"], r["qf"]), []).append(r["psnr_tc_db"])
    return [{"variant": v, "sigma2": s, "qf": q, "n_seeds": len(p), "psnr_tc_median_db": statistics.median(p)}
            for (v, s, q), p in sorted(groups.items())]


def _jsonable(obj):
    if isinstance(obj, float):
        return json_number(obj)
    if isinstance(obj, dict):
        return {k: _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    return obj


def run(paths, image, plan, out_dir, state, workers):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path = out_dir / RUN_FILE, out_dir / BOOK_FILE
    if run_path.exists():
        raise RunRefusal(f"{run_path} already exists: a run file of cycle {CYCLE} is never overwritten")
    try:
        avif.check_versions("AVIF", avif.EXPECTED_VERSIONS, avif.versions())
    except RuntimeError as exc:
        raise RunRefusal(str(exc)) from exc
    reference = load_pinned(paths["v3b_v4_reference"], REFERENCE_SHA256, "V3-b/V4 reference")
    rs24 = {r["sigma2"]: r for r in reference["rs24_oop_by_psnr"] if r["image"] == "frisco" and r["codec"] == "avif"}
    cycle7 = load_pinned(paths["cache_root"] / "cycles" / "7" / "v3b_v4.json", CYCLE7_SHA256, "cycle 7 run file")
    noise7 = {(n["sigma2"], n["seed"]): n["noisy_sha256"] for n in cycle7["noise"] if n["image"] == "frisco"}
    true = image["array"]
    noisies = []
    for s2 in plan["sigma2"]:
        for seed in plan["seeds"]:
            noisy = noisy_image(true, s2, seed)
            if hashlib.sha256(noisy.tobytes()).hexdigest() != noise7.get((s2, seed)):
                raise RunRefusal(f"noisy Frisco s2 {s2} seed {seed} differs from cycle 7 (sha256)")
            noisies.append((s2, seed, noisy))
    started, t0 = datetime.datetime.now(datetime.timezone.utc), time.perf_counter()
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")
    rows, walls, unavailable, env_d = [], {}, {}, None
    for variant in plan["variants"]:
        t = time.perf_counter()
        if variant == "d":
            python = paths["avif_libheif_python"]
            if python is None or not python.is_file():
                unavailable["d"] = "lab_paths.json has no avif_libheif_python or the file does not exist"
                print(f"  variant (d) unavailable: {unavailable['d']}")
                continue
            try:
                got, env_d = code_libheif(python, true, noisies, plan["qualities"], out_dir)
            except (subprocess.CalledProcessError, TimeoutError, OSError, KeyError, ValueError) as exc:
                unavailable["d"] = f"the libheif venv failed: {type(exc).__name__}: {str(exc)[:400]}"
                print(f"  variant (d) unavailable: {unavailable['d']}")
                continue
        else:
            got = code_pillow(variant, true, noisies, plan["qualities"], workers)
        walls[variant] = round(time.perf_counter() - t, 1)
        rows += got
        valid = [r for r in got if r["status"] == "valid"]
        print(f"  variant ({variant}): {len(valid)}/{len(got)} valid, {walls[variant]} s")
    for k, r in enumerate(rows, start=1):
        r["row_index"] = k
    curves = curves_of(rows)
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "wall_seconds": round(time.perf_counter() - t0, 1), **state, "declared_before_numbers": DECLARED,
        "plan": {"sigma2": list(plan["sigma2"]), "seeds": list(plan["seeds"]),
                 "qualities": [min(plan["qualities"]), max(plan["qualities"]), len(plan["qualities"])],
                 "variants": list(plan["variants"]), "workers": workers, "cpu_count": os.cpu_count()},
        "variants": {k: {"what": VARIANT_TEXT[k], **({"encode_args": PILLOW_VARIANTS[k]} if k in PILLOW_VARIANTS else {})}
                     for k in plan["variants"]},
        "variant_d_environment": env_d, "variant_d_defaults": LIBHEIF_DEFAULTS, "unavailable": unavailable,
        "b08_versions": avif.versions(), "b08_codec_versions": avif.codec_versions(),
        "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__,
                        "numpy": np.__version__, "platform": platform.platform()},
        "reference": {"file": paths["v3b_v4_reference"].name, "sha256": REFERENCE_SHA256, "rs24_frisco_avif": rs24},
        "cycle7_run_file": {"file": "cycles/7/v3b_v4.json", "sha256": CYCLE7_SHA256},
        "image": {k: v for k, v in image.items() if k != "array"} | {"pixels_sha256": hashlib.sha256(true.tobytes()).hexdigest()},
        "noise": [{"sigma2": s2, "seed": seed, "noisy_sha256": hashlib.sha256(n.tobytes()).hexdigest()}
                  for s2, seed, n in noisies],
        "unit_of_analysis": "row = variant x sigma2 x seed x QF; summary per (variant, sigma2) over seeds",
        "rows": rows, "curves": curves, "median_curves": median_curves(rows),
        "summary": summary_of(curves, rows, rs24, walls),
        "regression_a": regression([r for r in rows if r["variant"] == "a"], cycle7) if "a" in plan["variants"] else None,
    }
    text = json.dumps(_jsonable(record), ensure_ascii=False, indent=1, allow_nan=False)
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        f.write(text)
    build_book(run_path, book_path)
    print_summary(record, run_path, book_path)
    return run_path, book_path


def print_summary(record, run_path, book_path):
    print(f"\nRows: {len(record['rows'])}; wall time {record['wall_seconds']} s")
    for s in record["summary"]:
        print(f"  ({s['variant']}) s2 {s['sigma2']:>3}: QF_OOP {s['qf_oop_median']} [{s['qf_oop_min']}; {s['qf_oop_max']}] "
              f"(RS24 {s['rs24_qf_oop']}), PSNR_max {s['psnr_max_db_median']:.3f} dB (RS24 {s['rs24_psnr_max_db']}, "
              f"diff {s['d_psnr_max_db']:+.3f}), CR {s['cr_file_oop_median']:.2f} (RS24 {s['rs24_cr']}), "
              f"encode {s['encode_seconds_median']:.3f} s/call")
    reg = record["regression_a"]
    if reg:
        print(f"Regression (a) against cycle 7: {reg['rows_equal']} of {reg['rows_compared']} rows equal -> {reg['verdict']}")
    for k, why in record["unavailable"].items():
        print(f"Variant ({k}) unavailable: {why}")
    print(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}")


SUMMARY = [("Вариант", "variant"), ("σ²", "sigma2"), ("Зёрен, шт", "n_seeds")] + [
    (f"{label}: {stat}", f"{key}_{k}") for key, label in (("qf_oop", "QF_OOP, QF"), ("psnr_max_db", "PSNR_max, дБ"),
                                                          ("cr_file_oop", "CR по файлу в OOP, безразм."))
    for k, stat in (("median", "медиана"), ("min", "наименьшее"), ("max", "наибольшее"))] + [
    ("RS24 QF_OOP, QF", "rs24_qf_oop"), ("RS24 PSNR_max, дБ", "rs24_psnr_max_db"), ("RS24 CR, безразм.", "rs24_cr"),
    ("ΔQF_OOP (медиана − RS24), QF", "d_qf_oop"), ("ΔPSNR_max (медиана − RS24), дБ", "d_psnr_max_db"),
    ("CR: медиана / RS24 − 1, %", "cr_rel"), ("OOP внутри 0…100, зёрен", "n_oop_inside_range"),
    ("Кодирование: сумма, с", "encode_seconds_total"), ("Кодирование: медиана на вызов, с", "encode_seconds_median"),
    ("Вариант целиком (стена), с", "variant_wall_seconds")]
CURVES = [("Вариант", "variant"), ("σ²", "sigma2"), ("Зерно", "seed"), ("Valid, шт", "n_valid"),
          ("Не valid, шт", "n_not_valid"), ("QF_OOP, QF", "qf_oop"), ("PSNR_max, дБ", "psnr_max_db"),
          ("CR по файлу в OOP, безразм.", "cr_file_oop"), ("OOP внутри 0…100", "oop_inside_range")]
ROWS = [("№", "row_index"), ("Вариант", "variant"), ("σ²", "sigma2"), ("Зерно", "seed"), ("QF", "qf"),
        ("Статус", "status"), ("Причина", "reason"), ("Файл, байт", "file_bytes"), ("mdat (нагрузка), байт", "mdat_payload_bytes"),
        ("CR по файлу, безразм.", "cr_file"), ("CR по mdat, безразм.", "cr_mdat"), ("PSNR_tc, дБ", "psnr_tc_db"),
        ("PSNR_nc, дБ", "psnr_nc_db"), ("av1C monochrome", "av1c_monochrome"), ("Кодирование, с", "encode_seconds"),
        ("sha256 файла", "file_sha256")]


def _shown(value):
    return ("да" if value else "нет") if isinstance(value, bool) else value


def _table(cells, formats, sheet, columns, records, percent=()):
    for j, (head, key) in enumerate(columns, start=1):
        cells[sheet, 1, j], cells[sheet, 2, j] = head, key
    for i, rec in enumerate(records):
        for j, (_, key) in enumerate(columns, start=1):
            cells[sheet, 3 + i, j] = _shown(rec.get(key))
            if key in percent:
                formats[sheet, 3 + i, j] = "0.0%"


def book_cells(record, run_sha256):
    cells, formats = {}, {}
    reg = record.get("regression_a")
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
             ("Закончен, UTC", record["finished_utc"]), ("Время прогона, с", record["wall_seconds"]),
             ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
             *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
             ("Версии B08", json.dumps(record["b08_versions"], ensure_ascii=False)),
             ("Окружение (г)", json.dumps(record["variant_d_environment"], ensure_ascii=False)),
             ("Умолчания (г)", record["variant_d_defaults"]),
             ("Недоступно", json.dumps(record["unavailable"], ensure_ascii=False)),
             ("Регрессия (а)", "не выполнялась" if not reg else
              f"{reg['verdict']}: равны {reg['rows_equal']} из {reg['rows_compared']} строк цикла 7"),
             ("Строк, шт", len(record["rows"])),
             ("Ключи листов", "под шапкой — имя поля файла прогона: «Сводка» — summary[], «Зёрна» — curves[], «Строки» — "
                              "rows[], «График» — median_curves[]"),
             ("Листы «Рисунки» и «Страницы»", "не строятся: страницы не применяются (объявлено до чисел)"),
             ("На пальцах", "Один тестовый аэроснимок с шумом сжат в AVIF четырьмя способами при каждом значении "
                            "качества; для каждого способа найдена степень сжатия, при которой сжатый снимок ближе "
                            "всего к чистому, и это сравнивается со статьёй группы.")]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v
    _table(cells, formats, "Сводка", SUMMARY, record["summary"], percent=("cr_rel",))
    _table(cells, formats, "Зёрна", CURVES, record["curves"])
    _table(cells, formats, "Строки", ROWS, record["rows"])
    series = sorted({(m["variant"], m["sigma2"]) for m in record["median_curves"]})
    qs = sorted({m["qf"] for m in record["median_curves"]})
    med = {(m["variant"], m["sigma2"], m["qf"]): m["psnr_tc_median_db"] for m in record["median_curves"]}
    cells["График", 1, 1], cells["График", 2, 1] = "median_curves[].qf", "QF"
    for j, (v, s) in enumerate(series, start=2):
        cells["График", 1, j], cells["График", 2, j] = "median_curves[].psnr_tc_median_db", f"({v}) σ²={s}"
    for r, q in enumerate(qs, start=3):
        cells["График", r, 1] = q
        for j, (v, s) in enumerate(series, start=2):
            if (v, s, q) in med:
                cells["График", r, j] = med[v, s, q]
    return cells, formats, len(series)


def build_book(run_path, book_path):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    cells, formats, n_series = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Сводка", "Зёрна", "Строки", "График"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    for (sheet, r, c), fmt in formats.items():
        wb[sheet].cell(row=r, column=c).number_format = fmt
    wb["Описание"].column_dimensions["A"].width = 34
    wb["Описание"].column_dimensions["B"].width = 140
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for name in ("Сводка", "Зёрна", "Строки", "График"):
        for c in wb[name][2 if name == "График" else 1]:
            c.font = Font(bold=True)
        wb[name].freeze_panes = "A3"
    g = wb["График"]
    n_rows = max(r for s, r, _ in cells if s == "График")
    if n_rows >= 3 and n_series:
        chart = LineChart()
        chart.title = "AVIF: PSNR_tc (медиана по зёрнам) по QF, варианты (а)–(г)"
        chart.y_axis.title, chart.x_axis.title = "PSNR_tc, дБ", "QF"
        chart.x_axis.delete = chart.y_axis.delete = False
        chart.add_data(Reference(g, min_col=2, max_col=1 + n_series, min_row=2, max_row=n_rows), titles_from_data=True)
        chart.set_categories(Reference(g, min_col=1, min_row=3, max_row=n_rows))
        chart.width, chart.height = 28, 14
        g.add_chart(chart, "B" + str(n_rows + 3))
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g, as openpyxl writes them); no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())[0]
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    parser.add_argument("--workers", type=int, default=os.cpu_count() or 1, help="processes for variants (a)-(c)")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        (image,) = load_test_images(paths["images_root"], [IMAGE], (512, 512))
        run(paths, image, FULL_PLAN, out_dir, state, args.workers)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
