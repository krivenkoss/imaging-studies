"""V2-e: does the VVC stream packed into HEIF by GPAC decode back to the same image on all 59 V2-a pairs? Run: python studies/P001_vvc-dental/run_v2e.py

Study P001, cycle 5: check V2-e of STUDY.md section 8 for the packing (../REFERENCE.md#b09). Every pair
of V2-a is encoded again by B05 with the settings of cycle 3 and must give the bitstream and recon
of cycles/3/v2a_vtm.json (otherwise the run stops); then B09 packs and checks the round trip. Paths
come from lab_paths.json (images_root, cache_root, coders.vtm, coders.gpac). Writes
<cache_root>/cycles/5/v2e_heif.json (refuses by name if it exists) and the book v2e_heif.xlsx beside
it; `--rebuild-book` rebuilds the book from the run file.
"""

import argparse
import datetime
import hashlib
import json
import platform
import sys
import time
from pathlib import Path

import numpy as np

import imstudies
from imstudies.core.metrics import Refusal
from imstudies.lab import synthetic
from imstudies.lab.coders import vtm, vvc_heif
from imstudies.lab.images import load_test_images
from imstudies.lab.runs import RunRefusal, git_state, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
COMMAND = "python studies/P001_vvc-dental/run_v2e.py"
QUESTION = ("V2-e: восстанавливается ли поток VVC, упакованный GPAC в HEIF, в то же изображение на всех 59 парах "
            "V2-a, и воспроизводится ли при этом поток цикла 3?")
CYCLE, RUN_FILE, BOOK_FILE = 5, "v2e_heif.json", "v2e_heif.xlsx"
TEST_IMAGES = ["BABOON.BMP", "BARBARA.BMP", "goldhill.bmp", "lenna.bmp", "frisco.bmp", "DIEGO.BMP"]
QPS = (-12, 0, 12, 22, 32, 42, 51, 63)
ODD_SHAPES = ((8, 8), (21, 13), (517, 389))
LEVEL = {"value": "5", "status": "approved", "basis": "B05, owner 2026-10-14 18:13 UTC"}
FRAME_RATE = {"value": 1, "status": "approved", "basis": "B05, owner 2026-10-14 18:13 UTC"}
TIMEOUT_S = {"value": 600, "basis": "B05 settings of cycle 3; per EncoderApp, DecoderApp, MP4Box or gpac run"}
REPACK = (("e0", 32), ("syn517x389", 32))
DECLARED = {
    "v2e": ("Все 59 пар V2-a: B05 повторно с параметрами цикла 3 (Level=5, FrameRate=1, timeout_s=600) — sha256 "
            "потока и recon = строке cycles/3/v2a_vtm.json (иначе стоп — первая строка отчёта); затем B09 — упаковка, "
            "извлечение, DecoderApp — побитово = recon. Повторная упаковка двух пар — тот же sha256 файла, или "
            "названо, что меняется."),
    "refutation": "Хоть одна пара без побитового совпадения — V2-e не пройдена.",
    "sizes": ("Числа размеров и CR в этом цикле — не результат исследования (тестовые и синтетические изображения): "
              "печатаются для контроля."),
    "pages": "Страницы для глаз не применяются — побитовая сверка.",
}


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {k: (ROOT / raw[k]).resolve() for k in ("images_root", "cache_root") if raw.get(k)}
    for codec, keys in (("vtm", ("encoder", "decoder", "cfg")), ("gpac", ("mp4box", "gpac"))):
        for key in keys:
            value = (raw.get("coders", {}).get(codec) or {}).get(key)
            if not value:
                raise RunRefusal(f"lab_paths.json has no key coders.{codec}.{key}")
            paths[key] = (ROOT / value).resolve()
            if not paths[key].is_file():
                raise RunRefusal(f"{codec} {key} not found: {paths[key]}")
    for key in ("images_root", "cache_root"):
        if key not in paths:
            raise RunRefusal(f"lab_paths.json has no key '{key}'")
    manifest = paths["encoder"].parent / "MANIFEST.json"
    if not manifest.is_file():
        raise RunRefusal(f"{manifest} not found: the executables are not a build of tools/build_coders.ps1")
    paths["manifest"] = json.loads(manifest.read_text(encoding="utf-8"))
    for key in ("encoder", "decoder", "cfg"):
        if sha256_file(paths[key]) != paths["manifest"]["files_sha256"].get(paths[key].name):
            raise RunRefusal(f"{paths[key].name} differs from its MANIFEST.json sha256")
    return paths


def images_of_the_cycle(images_root):
    images = [{"code": "e0", "class": "control", "array": synthetic.e0_frame()}]
    images += [{"code": i["code"], "class": "test", "array": i["array"]}
               for i in load_test_images(images_root, TEST_IMAGES, (512, 512))]
    images += [{"code": f"syn{h}x{w}", "class": "synthetic", "array": synthetic.e0_crop(h, w)} for h, w in ODD_SHAPES]
    return images


def check_against_cycle_3(images, run3):
    recorded = {i["code"]: i["pixels_sha256"] for i in run3["images"]}
    differ = [i["code"] for i in images if recorded.get(i["code"]) != hashlib.sha256(i["array"].tobytes()).hexdigest()]
    if differ:
        raise RunRefusal(f"images differ from cycle 3 (pixels sha256): {differ}")


def encode(paths, image, qp, work):
    return vtm.encode_decode(image, qp, max_pixels=image.size, level=LEVEL["value"], frame_rate=FRAME_RATE["value"],
                             timeout_s=TIMEOUT_S["value"], encoder=paths["encoder"], decoder=paths["decoder"],
                             cfg=paths["cfg"], work_root=work)


def package(paths, coded, work):
    return vvc_heif.pack(coded.bitstream, coded.recon, timeout_s=TIMEOUT_S["value"], mp4box=paths["mp4box"],
                         gpac=paths["gpac"], decoder=paths["decoder"], work_root=work)


def run(paths, images, plan, repack, out_dir, state):
    """Pairs of `plan` ((code, qp)): B05 against cycle 3, then B09; `repack` pairs packed a second time."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path = out_dir / RUN_FILE, out_dir / BOOK_FILE
    if run_path.exists():
        raise RunRefusal(f"{run_path} already exists: a run file of cycle {CYCLE} is never overwritten")
    run3 = json.loads((paths["cache_root"] / "cycles" / "3" / "v2a_vtm.json").read_text(encoding="utf-8"))
    check_against_cycle_3(images, run3)
    cycle3 = {(r["image_code"], r["qp"]): r for r in run3["rows"]}
    by_code = {img["code"]: img for img in images}
    started = datetime.datetime.now(datetime.timezone.utc)
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")
    rows, files, first_provenance = [], {}, None
    for k, (code, qp) in enumerate(plan, start=1):
        image, then = by_code[code]["array"], cycle3[code, qp]
        t0 = time.perf_counter()
        coded = encode(paths, image, qp, out_dir)
        t1 = time.perf_counter()
        stream_sha, recon_sha = hashlib.sha256(coded.bitstream).hexdigest(), hashlib.sha256(coded.recon.tobytes()).hexdigest()
        if (stream_sha, recon_sha) != (then["bitstream_sha256"], then["recon_sha256"]):
            raise RunRefusal(f"pair {code} QP {qp}: B05 does not reproduce cycle 3 (bitstream {stream_sha} vs "
                             f"{then['bitstream_sha256']}, recon {recon_sha} vs {then['recon_sha256']}) - stop")
        row = {"pair_index": k, "image_code": code, "image_class": by_code[code]["class"], "qp": qp,
               "h": image.shape[0], "w": image.shape[1], "bitstream_sha256": stream_sha, "recon_sha256": recon_sha,
               "reproduces_cycle_3": True, "stream_bytes": len(coded.bitstream),
               "level": LEVEL["value"], "level_status": LEVEL["status"],
               "frame_rate": FRAME_RATE["value"], "frame_rate_status": FRAME_RATE["status"],
               "vtm_seconds": round(t1 - t0, 2)}
        try:
            packed = package(paths, coded, out_dir)
        except Refusal as refusal:
            row |= {"status": "no_result", "reason": refusal.reason, "detail": str(refusal), "round_trip": False}
        except Exception as exc:  # noqa: BLE001 - status `error` by SPEC-RESULT; the detail is kept
            row |= {"status": "error", "reason": "internal_error", "detail": f"{type(exc).__name__}: {exc}"[:600],
                    "round_trip": False}
        else:
            p = packed.provenance
            first_provenance = first_provenance or p
            files[code, qp] = hashlib.sha256(packed.file).hexdigest()
            row |= {"status": "valid", "reason": None, "detail": None, "round_trip": True,
                    "heif_bytes": len(packed.file), "heif_sha256": files[code, qp], "mdat_bytes": p["mdat_bytes"],
                    "free_bytes": p["free_bytes"], "meta_bytes": p["box_bytes"].get("meta"),
                    "ftyp_bytes": p["box_bytes"].get("ftyp"), "overhead_bytes": len(packed.file) - len(coded.bitstream),
                    "extracted_stream_bytes": p["extracted_stream_bytes"], "ispe_ok": p["ispe"] == [image.shape[1], image.shape[0]]}
        row["seconds"] = round(time.perf_counter() - t0, 2)
        rows.append(row)
        print(f"  {k:>2}/{len(plan)} {code:<10} QP {qp:>3}: cycle 3 reproduced, {row['status']}, round trip "
              f"{row['round_trip']}, stream {row['stream_bytes']} B, HEIF {row.get('heif_bytes')} B ({row['seconds']} s)")
    repacked = []
    for code, qp in repack:
        coded = encode(paths, by_code[code]["array"], qp, out_dir)
        again = hashlib.sha256(package(paths, coded, out_dir).file).hexdigest()
        repacked.append({"image_code": code, "qp": qp, "first_sha256": files.get((code, qp)), "second_sha256": again,
                         "equal": again == files.get((code, qp))})
    n_rt = sum(r["round_trip"] for r in rows)
    summary = {"n_pairs": len(rows), "n_images": len({r["image_code"] for r in rows}),
               "n_reproduces_cycle_3": sum(r["reproduces_cycle_3"] for r in rows), "n_round_trip": n_rt,
               "n_repack": len(repacked), "n_repack_equal": sum(r["equal"] for r in repacked),
               "verdict_v2e": "pass" if rows and n_rt == len(rows) else "fail",
               "refutation_fired": n_rt != len(rows)}
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        **state, "declared_before_numbers": DECLARED, "level": LEVEL, "frame_rate": FRAME_RATE, "timeout_s": TIMEOUT_S,
        "build_manifest": paths["manifest"], "cycle_3_run_file_sha256": sha256_file(paths["cache_root"] / "cycles" / "3" / "v2a_vtm.json"),
        "gpac_writes": {k: first_provenance[k] for k in ("gpac_version", "ftyp_brands", "pixi_bits", "free_text",
                                                         "mp4box_sha256", "gpac_sha256", "libgpac_sha256")}
        if first_provenance else None,
        "commands_template": {k: [a if not a.endswith((".266", ".heic", ".vvc", ".yuv", ".266:primary")) else Path(a).name
                                  for a in first_provenance[k]] for k in ("pack_command", "extract_command", "decode_command")}
        if first_provenance else None,
        "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__,
                        "numpy": np.__version__, "platform": platform.platform()},
        "images": [{"code": i["code"], "class": i["class"], "shape": list(i["array"].shape),
                    "pixels_sha256": hashlib.sha256(i["array"].tobytes()).hexdigest()} for i in images],
        "unit_of_analysis": "pair = image x QP; n is counted by images",
        "rows": rows, "repack": repacked, "summary": summary,
    }
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        json.dump(record, f, ensure_ascii=False, indent=1, allow_nan=False)
    build_book(run_path, book_path)
    s = summary
    print(f"\nV2-e: cycle 3 reproduced on {s['n_reproduces_cycle_3']} of {s['n_pairs']} pairs; round trip on "
          f"{s['n_round_trip']} of {s['n_pairs']} -> {s['verdict_v2e']}; repack equal {s['n_repack_equal']} of "
          f"{s['n_repack']}; refutation fired: {s['refutation_fired']}")
    print(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}")
    return run_path, book_path


TABLE = [("№ пары", "pair_index"), ("Изображение (код)", "image_code"), ("Класс", "image_class"), ("QP", "qp"),
         ("H, пикс", "h"), ("W, пикс", "w"), ("Цикл 3 воспроизведён", "reproduces_cycle_3"), ("Статус", "status"),
         ("Причина", "reason"), ("Туда-обратно = recon", "round_trip"), ("Поток, байт", "stream_bytes"),
         ("HEIF, байт", "heif_bytes"), ("mdat, байт", "mdat_bytes"), ("free, байт", "free_bytes"),
         ("meta, байт", "meta_bytes"), ("HEIF − поток, байт", "overhead_bytes"), ("ispe = (W, H)", "ispe_ok"),
         ("sha256 потока", "bitstream_sha256"), ("sha256 HEIF", "heif_sha256"), ("Level", "level"),
         ("Статус Level", "level_status"), ("FrameRate", "frame_rate"), ("Статус FrameRate", "frame_rate_status"),
         ("VTM, с", "vtm_seconds"), ("Всего, с", "seconds")]
YES_NO = {True: "да", False: "нет"}


def book_cells(record, run_sha256):
    cells, s, g = {}, record["summary"], record["gpac_writes"] or {}
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
             ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
             *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
             ("Level / FrameRate", f"{record['level']['value']} / {record['frame_rate']['value']} (approved, B05)"),
             ("GPAC", g.get("gpac_version")), ("Бренды ftyp", ", ".join(g.get("ftyp_brands") or [])),
             ("pixi, бит по каналам", ", ".join(str(b) for b in g.get("pixi_bits") or [])),
             ("Текст бокса free", g.get("free_text")),
             ("Пар, шт", s["n_pairs"]), ("Изображений, шт", s["n_images"]),
             ("Цикл 3 воспроизведён, шт", s["n_reproduces_cycle_3"]), ("Туда-обратно совпало, шт", s["n_round_trip"]),
             ("Повторная упаковка: пар, шт", s["n_repack"]), ("из них файл тот же, шт", s["n_repack_equal"]),
             ("Вердикт V2-e", s["verdict_v2e"]), ("Опровержение сработало", YES_NO[s["refutation_fired"]]),
             ("Точность чисел книги", "16 значащих цифр (openpyxl пишет %.16g); каждая клетка сверена с файлом прогона"),
             ("Листы «Рисунки» и «Страницы»", "не строятся: страницы не применяются (объявлено до чисел)"),
             ("На пальцах", "Каждый поток VVC из сверки цикла 3 получен заново, уложен программой GPAC в файл HEIF, "
                            "извлечён обратно и раскодирован; картинка сравнивалась с исходной раскодированной.")]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v
    for j, (head, key) in enumerate(TABLE, start=1):
        cells["Таблица", 1, j], cells["Таблица", 2, j] = head, f"rows[].{key}"
    for i, row in enumerate(record["rows"]):
        for j, (_, key) in enumerate(TABLE, start=1):
            v = row.get(key)
            cells["Таблица", 3 + i, j] = YES_NO[v] if isinstance(v, bool) else v
    cells["График", 1, 1] = "Служебные байты HEIF сверх потока по парам, байт. Ключ: rows[].overhead_bytes"
    cells["График", 2, 1], cells["График", 2, 2] = "№ пары", "HEIF − поток, байт"
    for i, row in enumerate(record["rows"]):
        cells["График", 3 + i, 1], cells["График", 3 + i, 2] = row["pair_index"], row.get("overhead_bytes")
    return cells


def build_book(run_path, book_path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Таблица", "График"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    wb["Описание"].column_dimensions["A"].width = 36
    wb["Описание"].column_dimensions["B"].width = 120
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for c in wb["Таблица"][1]:
        c.font = Font(bold=True)
    g = wb["График"]
    n_rows = max(r for s, r, _ in cells if s == "График")
    if n_rows >= 3:
        chart = BarChart()
        chart.title, chart.y_axis.title, chart.x_axis.title = "HEIF − поток, байт", "байт", "№ пары"
        chart.x_axis.delete = chart.y_axis.delete = False
        chart.add_data(Reference(g, min_col=2, max_col=2, min_row=2, max_row=n_rows), titles_from_data=True)
        chart.set_categories(Reference(g, min_col=1, min_row=3, max_row=n_rows))
        chart.width, chart.height = 26, 12
        g.add_chart(chart, "E2")
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g, as openpyxl writes them); no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        images = images_of_the_cycle(paths["images_root"])
        plan = [(i["code"], qp) for i in images if i["class"] != "synthetic" for qp in QPS]
        plan += [(f"syn{h}x{w}", 32) for h, w in ODD_SHAPES]
        run(paths, images, plan, REPACK, out_dir, state)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
