"""Lab tests of run_v3a.py (study P001, cycle 5): test images, lab_paths.json and the V3-a reference needed.

- a run on one pair writes its run file and a book equal to it cell by cell, and refuses by name
  the second time; the cell check refuses a changed cell;
- a reference row made from other pixels is refused before any comparison.
"""

import importlib.util
import json
from pathlib import Path

import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def v3a():
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    spec = importlib.util.spec_from_file_location("run_v3a", STUDY / "run_v3a.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def setup(v3a):
    try:
        paths = v3a.load_paths()
    except v3a.RunRefusal as refusal:
        pytest.skip(f"V3-a is not ready: {refusal}")
    return paths, v3a.load_test_images(paths["images_root"], v3a.TEST_IMAGES, (512, 512))


def test_run_on_one_pair(v3a, setup, tmp_path):
    paths, images = setup
    state = {"core_commit": "test", "core_tree": "test"}
    run_path, book_path = v3a.run(paths, images[:1], (50,), tmp_path, state)
    record = json.loads(run_path.read_text(encoding="utf-8"))
    (row,) = record["rows"]
    assert (row["image_code"], row["quality"], row["status"], row["match"]) == ("baboon", 50, "valid", True)
    assert row["sha256"] == row["reference_sha256"] and record["summary"]["verdict_v3a"] == "pass"
    with pytest.raises(v3a.RunRefusal, match=v3a.RUN_FILE):
        v3a.run(paths, images[:1], (50,), tmp_path, state)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Таблица"].cell(row=3, column=6)  # bytes of our file
    assert cell.value == row["bytes"]
    cell.value += 1
    changed = tmp_path / "changed.xlsx"
    book.save(changed)
    with pytest.raises(v3a.RunRefusal, match=r"Таблица!R3C6 .* differs from the run file"):
        v3a.check_book(run_path, changed)


def test_reference_from_other_pixels_is_refused(v3a, setup):
    paths, images = setup
    reference = json.loads(paths["v3a_reference"].read_text(encoding="utf-8"))
    cycle1 = json.loads((paths["cache_root"] / "cycles" / "1" / "v1_metrics.json").read_text(encoding="utf-8"))
    v3a.check_reference(reference, images, cycle1["images"])  # the intact reference passes
    reference["rows"][0] = dict(reference["rows"][0], pixels_sha256="0" * 64)
    with pytest.raises(v3a.RunRefusal, match="pixels sha256 differs from cycle 1"):
        v3a.check_reference(reference, images, cycle1["images"])
