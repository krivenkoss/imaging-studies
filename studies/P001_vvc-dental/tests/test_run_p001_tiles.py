"""Lab tests of run_p001_tiles.py (study P001, cycle 8) on a synthetic copy in tmp_path (no medical image).

- a run writes its run file, one page per image and a book equal to the run file cell by cell; every
  page claim matches its run-file entry; the second run refuses by name;
- the book check refuses a changed cell and accepts an owner's verdict; a rebuild carries the verdict,
  and refuses by name when a page with a verdict is gone.
"""

import hashlib
import importlib.util
import json
from pathlib import Path

import numpy as np
import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def tiles_run():
    spec = importlib.util.spec_from_file_location("run_p001_tiles", STUDY / "run_p001_tiles.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def synthetic_copy(root):
    from PIL import Image

    rng = np.random.default_rng(3)
    (root / "P001").mkdir(parents=True)
    (root / "registry").mkdir()
    entries = []
    for code, (h, w) in (("D01", (1100, 1600)), ("D02", (600, 700))):
        values = rng.integers(0, 256, (h, w)).astype(np.uint8)
        values[:300, :] = 0  # a dark band: background in the first row of tiles
        Image.fromarray(values, "L").save(root / "P001" / f"{code}.png")
        entries.append({"code": code, "H": h, "W": w, "full_tiles": (h // 512) * (w // 512),
                        "png_sha256": hashlib.sha256((root / "P001" / f"{code}.png").read_bytes()).hexdigest()})
    (root / "registry" / "P001.json").write_text(json.dumps({"contract": "B10 0.2.0", "core_commit": "test",
                                                             "images": entries}), encoding="utf-8")
    return root


def test_run_book_pages_and_verdicts(tiles_run, tmp_path):
    data = synthetic_copy(tmp_path / "data")
    out = tmp_path / "cycle8"
    state = {"core_commit": "test", "core_tree": "test"}
    lines = []
    run_path, book_path = tiles_run.run(data, out, state, log=lines.append)
    record = json.loads(run_path.read_text(encoding="utf-8"))
    assert [(e["code"], e["n_tiles"]) for e in record["per_image"]] == [("D01", 6), ("D02", 1)]
    assert record["set"]["n_tiles"] == 7 and record["set"]["best_achievable_images"] == 2
    for e, page in zip(record["per_image"], record["pages"]):
        tiles_run.tiles.check_statement(page["claim"], e)
        assert (out / page["file"]).is_file()
    first_row = [r for r in record["rows"] if r["code"] == "D01" and r["r"] == 0]
    assert all(r["share_black"] > 0.5 for r in first_row)  # the dark band is seen as background
    with pytest.raises(tiles_run.RunRefusal, match="already exists"):
        tiles_run.run(data, out, state, log=lines.append)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Фрагменты"].cell(row=3, column=10)  # share_any of the first tile
    assert cell.value == record["rows"][0]["share_any"]
    cell.value = 0.5 if cell.value != 0.5 else 0.6
    book.save(tmp_path / "changed.xlsx")
    with pytest.raises(tiles_run.RunRefusal, match=r"Фрагменты!R3C10 .* differs from the run file"):
        tiles_run.check_book(run_path, tmp_path / "changed.xlsx")

    book = load_workbook(book_path)
    book["Страницы"].cell(row=3, column=tiles_run.VERDICT_COLUMN, value="ок")
    book.save(book_path)
    tiles_run.check_book(run_path, book_path)  # the owner's verdict is allowed
    assert tiles_run.rebuild_book(run_path, book_path) == {record["pages"][0]["file"]: "ок"}
    assert tiles_run.read_verdicts(book_path) == {record["pages"][0]["file"]: "ок"}
    (out / record["pages"][0]["file"]).unlink()
    with pytest.raises(tiles_run.RunRefusal, match="pages with a verdict are gone from the disk"):
        tiles_run.rebuild_book(run_path, book_path)
