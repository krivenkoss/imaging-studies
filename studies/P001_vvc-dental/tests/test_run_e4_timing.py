"""Lab tests of run_e4_timing.py (study P001, cycle 10).

- the mosaic: row r shifted left by r, crop from the top-left corner - on synthetic constant tiles;
- the estimate formula recomputed by hand in the test (medians, VTM whole images by linear interpolation of
  the QP points, nearest outside, metrics 3 per file, division by k) on rows whose mean differs from the
  median;
- a run on one small configuration (VTM 2 QP, two values of each other codec, a mosaic 64 x 96 of the same
  images; lab_paths.json with the coders and the B12 environment needed): the pool equals the sequential run,
  the estimate equals the hand formula on the rows of the run file, the book equals the run file and refuses a
  changed cell, the second run refuses by name, no working folder is left under cache_root.
- the other side of the check: a pool with another file at one QP is a stop recorded in the run file, no
  estimate and nothing coded after it (VTM replaced by a stub).
"""

import importlib.util
import json
from pathlib import Path

import numpy as np
import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def e4():
    spec = importlib.util.spec_from_file_location("run_e4_timing", STUDY / "run_e4_timing.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_mosaic_rows_shifted_and_cropped(e4):
    tiles = [np.full((512, 512), k, np.uint8) for k in range(6)]
    m = e4.mosaic(tiles, (1504, 2773))
    assert m.shape == (1504, 2773)
    for r in range(3):
        for c in range(6):
            y, x = r * 512, c * 512
            if y < 1504 and x < 2773:
                assert m[y, x] == (c + r) % 6, (r, c)
    assert m[1503, 2772] == (5 + 2) % 6  # the last pixel lies in row 2, column 5


def middle(values):
    """Median by hand: the middle value, or the mean of the two middle values."""
    v = sorted(values)
    return v[len(v) // 2] if len(v) % 2 else (v[len(v) // 2 - 1] + v[len(v) // 2]) / 2


def hand_hours(rows, k):
    """The declared formula, written again without the module."""
    seq = [r for r in rows if r["mode"] == 1]
    n = {"vtm": 76, "jpeg": 100, "heif": 101, "avif": 101}
    total = 0.0
    for codec, n_values in n.items():
        total += 154 * n_values * middle(r["seconds"] for r in seq if (r["codec"], r["image"]) == (codec, "f"))
        if codec == "vtm":
            pts = sorted((r["param"], r["seconds"]) for r in seq if (r["codec"], r["image"]) == ("vtm", "m"))
            whole = 0.0
            for qp in range(-12, 64):
                if qp <= pts[0][0]:
                    whole += pts[0][1]
                elif qp >= pts[-1][0]:
                    whole += pts[-1][1]
                else:
                    (x1, y1), (x2, y2) = next(((a, b) for a, b in zip(pts, pts[1:]) if a[0] <= qp <= b[0]))
                    whole += y1 + (y2 - y1) * (qp - x1) / (x2 - x1)
            total += 16 * whole
        else:
            total += 16 * n_values * middle(r["seconds"] for r in seq if (r["codec"], r["image"]) == (codec, "m"))
    for size, units in (("f", 154), ("m", 16)):
        per_file = sum(middle(r["seconds"] for r in seq if (r["codec"], r["image"], r["param"]) == ("metric", size, name))
                       for name in ("psnr", "psnr_hvs_m", "haarpsi"))
        total += units * 378 * per_file
    return total / 3600, total / k / 3600


def test_estimate_by_hand_with_median_not_mean(e4):
    rows = []
    for codec in ("vtm", "jpeg", "heif", "avif"):
        for size in ("f", "m"):
            for param, seconds in ((0, 1.0), (22, 2.0), (42, 9.0)):  # mean 4, median 2
                rows.append({"codec": codec, "image": size, "param": param, "mode": 1, "seconds": seconds * (2 if size == "m" else 1)})
    rows.append({"codec": "vtm", "image": "f", "param": 63, "mode": 4, "seconds": 100.0})  # pool rows are not used
    for size in ("f", "m"):
        for name, base in (("psnr", 0.01), ("psnr_hvs_m", 0.5), ("haarpsi", 0.2)):
            rows += [{"codec": "metric", "image": size, "param": name, "mode": 1, "seconds": base * x} for x in (1, 2, 7)]
    got = e4.estimate(rows, 3.5, 2.0)
    h1, h4 = hand_hours(rows, 3.5)
    assert got["hours_1_process"] == pytest.approx(h1, rel=1e-12) and got["hours_4_processes"] == pytest.approx(h4, rel=1e-12)
    assert got["terms"]["vtm"]["t_f_median_s"] == 2.0 and 0 < got["vtm_share"] < 1


def test_pool_differing_from_sequential_is_a_recorded_stop(e4, tmp_path, monkeypatch):
    def fake_set(paths, size, image, qps, work, workers):  # the pool gives another file at the last QP
        rows = [{"param": q, "status": "valid", "seconds": 1.0, "file_bytes": 100, "bitstream_sha256": "s",
                 "file_sha256": f"{q}-{workers if q == qps[-1] else 1}", "image": size, "codec": "vtm", "mode": workers,
                 "set_wall_seconds": 2.0, "cpu_busy_share": None, "cr_file": 1.0} for q in qps]
        return rows, 2.0
    monkeypatch.setattr(e4, "vtm_set", fake_set)
    plan = {"vtm": {"f": (32, 42), "m": (32, 42)}, "metric_repeats": 1, "mosaic_shape": (8, 8)}
    images = {"f": np.zeros((8, 8), np.uint8), "m": np.zeros((8, 8), np.uint8)}
    run_path, _ = e4.run({"cache_root": tmp_path}, images, plan, tmp_path / "out", {"core_commit": "t", "core_tree": "t"})
    record = json.loads(run_path.read_text(encoding="utf-8"))
    assert record["pool_equals_sequential"] is False and record["estimate"] is None
    assert record["stopped"].startswith("VTM (f): pool differs from the sequential run in 1 of 2 QP: ['QP 42']")
    assert {r["codec"] for r in record["rows"]} == {"vtm"}  # nothing after the stop


@pytest.fixture(scope="module")
def setup(e4):
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    try:
        paths = e4.load_paths()
    except e4.RunRefusal as refusal:
        pytest.skip(f"the codecs of E4 timing are not configured: {refusal}")
    tiles = [i["array"] for i in e4.load_test_images(paths["images_root"], e4.MOSAIC_ORDER, (512, 512))]
    return paths, {"f": tiles[4], "m": e4.mosaic(tiles, (64, 96))}


def test_run_on_one_small_configuration(e4, setup, tmp_path):
    paths, images = setup
    plan = {"vtm": {"f": (32, 42), "m": (32, 42)}, "jpeg": {"f": (50, 90), "m": (50, 90)},
            "heif": {"f": (50, 90), "m": (50, 90)}, "avif": {"f": (50, 90), "m": (50, 90)},
            "metric_repeats": 3, "mosaic_shape": (64, 96)}
    state = {"core_commit": "test", "core_tree": "test"}
    before = set(paths["cache_root"].iterdir())
    run_path, book_path = e4.run(paths, images, plan, tmp_path, state)
    assert set(paths["cache_root"].iterdir()) == before  # the working folder of the run is removed
    record = json.loads(run_path.read_text(encoding="utf-8"))
    assert record["pool_equals_sequential"] and record["stopped"] is None
    assert len(record["rows"]) == 8 + 12 + 18
    k = record["estimate"]["k_pool_f"]
    walls = {(w["image"], w["mode"]): w["seconds"] for w in record["set_walls"]}
    assert k == walls["f", 1] / walls["f", 4]
    h1, h4 = hand_hours(record["rows"], k)
    assert record["estimate"]["hours_1_process"] == pytest.approx(h1, rel=1e-12)
    assert record["estimate"]["hours_4_processes"] == pytest.approx(h4, rel=1e-12)
    with pytest.raises(e4.RunRefusal, match=e4.RUN_FILE):
        e4.run(paths, images, plan, tmp_path, state)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Строки"].cell(row=3, column=6)  # seconds of the first row
    assert cell.value == pytest.approx(record["rows"][0]["seconds"])
    cell.value = cell.value + 1
    book.save(tmp_path / "changed.xlsx")
    with pytest.raises(e4.RunRefusal, match=r"Строки!R3C6 .* differs from the run file"):
        e4.check_book(run_path, tmp_path / "changed.xlsx")
