"""Lab tests of run_v1.py (study P001, cycle 1): Octave, the authors' files and local images needed.

- the Octave side reproduces the stored values of tests/vectors/metrics_reference.json;
- a run on one pair writes its run file and a book equal to it cell by cell, and refuses
  by name the second time;
- the cell check of the book refuses a book that does not match the run file;
- the dental selection gives four distinct images with codes, never names.
"""

import base64
import importlib.util
import json
from pathlib import Path

import numpy as np
import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def v1():
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    spec = importlib.util.spec_from_file_location("run_v1", STUDY / "run_v1.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def paths(v1):
    return v1.load_paths()


def test_octave_reproduces_stored_vectors(v1, paths, tmp_path):
    vectors = json.loads((ROOT / "tests" / "vectors" / "metrics_reference.json").read_text(encoding="utf-8"))
    cases = [c for c in vectors["cases"] if c["psnrhvsm_m"] is not None]
    pairs = [tuple(np.frombuffer(base64.b64decode(c[k]), np.uint8).reshape(c["shape"]) for k in ("ref", "dist"))
             for c in cases]
    values, _ = v1.octave_metrics(paths["octave_cli"], paths["metric_sources"], pairs, tmp_path)
    assert [v[0] for v in values] == [c["psnrhvsm_m"] for c in cases]
    assert [v[1] for v in values] == [c["haarpsi_m"] for c in cases]


def test_run_on_one_pair(v1, paths, tmp_path):
    images = v1.test_set(paths["images_root"])[:1]
    state = {"core_commit": "test", "core_tree": "test"}
    run_path, book_path, leaks = v1.run(paths, images, tmp_path, state, [], distortions=v1.DISTORTIONS[:1])
    record = json.loads(run_path.read_text(encoding="utf-8"))
    assert record["path_control"]["passed"]
    assert len(record["pairs"]) == 1 and len(record["rows"]) == 4
    assert {(r["metric"], r["reference"]) for r in record["rows"]} == {(m, r) for m, r, *_ in v1.CHECKS}
    assert all(r["status"] == "valid" and r["verdict"] == "pass" for r in record["rows"])
    assert record["distortions"][0]["seed"] == 1 and record["environment"]["octave"] == v1.OCTAVE_VERSION
    assert book_path.is_file() and leaks == 0
    with pytest.raises(v1.RunRefusal, match=v1.RUN_FILE):
        v1.run(paths, images, tmp_path, state, [], distortions=v1.DISTORTIONS[:1])


def test_book_check_refuses_a_changed_value(v1, paths, tmp_path):
    images = v1.test_set(paths["images_root"])[:1]
    run_path, book_path, _ = v1.run(paths, images, tmp_path, {"core_commit": "t", "core_tree": "t"}, [],
                                    distortions=v1.DISTORTIONS[2:3])
    from openpyxl import load_workbook

    v1.check_book(run_path, book_path)  # the intact book passes
    book = load_workbook(book_path)
    cell = book["Таблица"].cell(row=3, column=8)  # PSNR-HVS-M ours of the pair
    assert isinstance(cell.value, float)
    cell.value += 1e-9
    changed = tmp_path / "changed.xlsx"
    book.save(changed)
    with pytest.raises(v1.RunRefusal, match=r"Таблица!R3C8 .* differs from the run file"):
        v1.check_book(run_path, changed)


def test_dental_set_gives_four_distinct_coded_images(v1, paths):
    images, codes = v1.dental_set(paths["images_root"])
    assert [i["code"] for i in images] == ["v1d01", "v1d02", "v1d03", "v1d04"]
    assert len({c["file_sha256"] for c in codes}) == 4
    for img in images:
        assert img["array"].shape == (512, 512) and img["array"].dtype == np.uint8
        meta = json.dumps({k: v for k, v in img.items() if k != "array"}, ensure_ascii=False)
        assert all(c["file_name"] not in meta and Path(c["file_name"]).stem not in meta for c in codes)
