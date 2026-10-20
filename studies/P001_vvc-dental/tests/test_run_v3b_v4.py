"""Lab tests of run_v3b_v4.py (study P001, cycle 7): test images, lab_paths.json and the V3-b/V4 reference needed.

- the OOP rule, the crossing rule of the steep zone and the noise rule against hand-computed truths;
- a reference copy with one changed number is refused by its sha256; a bad row is refused by section
  and index;
- a run on one small configuration writes its run file and a book equal to it cell by cell, and
  refuses by name the second time; the cell check refuses a changed cell.
"""

import importlib.util
import json
from pathlib import Path

import numpy as np
import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def v3b():
    spec = importlib.util.spec_from_file_location("run_v3b_v4", STUDY / "run_v3b_v4.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def setup(v3b):
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    try:
        paths = v3b.load_paths()
    except v3b.RunRefusal as refusal:
        pytest.skip(f"V3-b/V4 is not ready: {refusal}")
    return paths, v3b.load_test_images(paths["images_root"], v3b.TEST_IMAGES, (512, 512))


def test_oop_largest_psnr_smaller_qf_on_tie(v3b):
    assert v3b.oop([(10, 30.0), (20, 31.0), (30, 31.0), (40, 29.0)]) == 20
    assert v3b.oop([(40, 29.0), (30, 31.5), (20, 31.0)]) == 30
    assert v3b.oop([]) is None


def test_crossing_linear_in_cr_nearest_to_reference(v3b):
    assert v3b.crossing([0, 10, 20], [100.0, 50.0, 0.0], 75.0, 3) == (5.0, 1)
    assert v3b.crossing([0, 10, 20], [100.0, 50.0, 0.0], 50.0, 3) == (10.0, 1)  # exactly on a grid point
    assert v3b.crossing([0, 10, 20], [100.0, 50.0, 0.0], 200.0, 3) == (None, 0)
    assert v3b.crossing([0, 10, 20], [10.0, 30.0, 10.0], 20.0, 14) == (15.0, 2)
    assert v3b.crossing([0, 10, 20], [10.0, 30.0, 10.0], 20.0, 10) == (5.0, 2)  # tie of distances: the smaller


def test_noise_rule_rounds_clips_and_follows_the_seed(v3b):
    flat = np.full((64, 64), 128, np.uint8)
    z = np.random.default_rng(3).standard_normal((64, 64))
    assert np.array_equal(v3b.noisy_image(flat, 25, 3), np.floor(128 + 5 * z + 0.5).astype(np.uint8))
    black, white = np.zeros((256, 256), np.uint8), np.full((256, 256), 255, np.uint8)
    assert 0.45 < (v3b.noisy_image(black, 100, 1) == 0).mean() < 0.55
    assert 0.45 < (v3b.noisy_image(white, 100, 1) == 255).mean() < 0.55
    assert np.array_equal(v3b.noisy_image(flat, 100, 2), v3b.noisy_image(flat, 100, 2))
    assert not np.array_equal(v3b.noisy_image(flat, 100, 2), v3b.noisy_image(flat, 100, 4))


def test_reference_changed_or_broken_is_refused(v3b, setup, tmp_path):
    paths, _ = setup
    reference = v3b.load_reference(paths["v3b_v4_reference"])  # the intact file passes
    text = paths["v3b_v4_reference"].read_text(encoding="utf-8")
    assert '"cr": 158.29' in text
    copy = tmp_path / "v3b_v4_reference.json"
    copy.write_text(text.replace('"cr": 158.29', '"cr": 158.3', 1), encoding="utf-8")
    with pytest.raises(v3b.RunRefusal, match=r"v3b_v4_reference\.json: sha256 \w+ differs from the declared"):
        v3b.load_reference(copy)
    broken = json.loads(json.dumps(reference))
    broken["as25_table1_heif_cr"][3]["qf"] = 40
    with pytest.raises(v3b.RunRefusal, match=r"reference row as25_table1_heif_cr\[3\]: unexpected or broken key"):
        v3b.check_reference(broken)
    broken = json.loads(json.dumps(reference))
    broken["rs24_oop_by_psnr"][1]["cr"] = -26.5
    with pytest.raises(v3b.RunRefusal, match=r"reference row rs24_oop_by_psnr\[1\]: value out of the declared kind"):
        v3b.check_reference(broken)
    broken = json.loads(json.dumps(reference))
    del broken["as25_heif_qf_oop_by_mse"][0]
    with pytest.raises(v3b.RunRefusal, match=r"as25_heif_qf_oop_by_mse lacks rows"):
        v3b.check_reference(broken)


def test_run_on_one_small_configuration(v3b, setup, tmp_path):
    paths, images = setup
    reference = v3b.load_reference(paths["v3b_v4_reference"])
    frisco = [img for img in images if img["code"] == "frisco"]
    plan = {"sigma2": (25,), "seeds": (1,), "qualities": {"jpeg": (44, 48, 52), "heif": (29, 33, 37), "avif": (47, 51, 55)}}
    state = {"core_commit": "test", "core_tree": "test"}
    run_path, book_path = v3b.run(paths, reference, frisco, plan, tmp_path, state)
    record = json.loads(run_path.read_text(encoding="utf-8"))
    rows = record["rows"]
    assert len(rows) == 9 and all(r["status"] == "valid" for r in rows)
    for r in rows:
        assert r["cr_file"] == 512 * 512 / r["file_bytes"]
        assert (r["mdat_payload_bytes"] is None) == (r["codec"] == "jpeg")
        if r["codec"] != "jpeg":
            assert 0 < r["mdat_payload_bytes"] < r["file_bytes"] and r["cr_mdat"] == 512 * 512 / r["mdat_payload_bytes"]
    for c in record["curves"]:
        mine = [r for r in rows if r["codec"] == c["codec"]]
        assert c["psnr_max_db"] == max(r["psnr_tc_db"] for r in mine)
    v3b_rows = [c for c in record["comparison"] if c["check"] == "V3-b" and c["sigma2"] == 25 and c["image"] == "frisco"]
    assert len(v3b_rows) == 3 and all(c["n_seeds"] == 1 and c["in_tolerance"] in (True, False) for c in v3b_rows)
    assert record["verdicts"]["V3-b"]["rows_without_data"] == 3  # the other three JPEG rows of the reference
    with pytest.raises(v3b.RunRefusal, match=v3b.RUN_FILE):
        v3b.run(paths, reference, frisco, plan, tmp_path, state)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Строки"].cell(row=3, column=9)  # file bytes of the first row
    assert cell.value == rows[0]["file_bytes"]
    cell.value += 1
    changed = tmp_path / "changed.xlsx"
    book.save(changed)
    with pytest.raises(v3b.RunRefusal, match=r"Строки!R3C9 .* differs from the run file"):
        v3b.check_book(run_path, changed)
