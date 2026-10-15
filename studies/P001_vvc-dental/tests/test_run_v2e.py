"""Lab tests of run_v2e.py (study P001, cycle 5): VTM build, GPAC, lab_paths.json and the cycle 3 run file needed.

- a run on one small pair reproduces cycle 3, packs, checks the round trip, writes its run file and a
  book equal to it cell by cell, and refuses by name the second time; a changed cell is refused;
- the images of the cycle are those of cycle 3 (pixels sha256).
"""

import importlib.util
import json
from pathlib import Path

import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def v2e():
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    spec = importlib.util.spec_from_file_location("run_v2e", STUDY / "run_v2e.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def setup(v2e):
    try:
        paths = v2e.load_paths()
    except v2e.RunRefusal as refusal:
        pytest.skip(f"VTM or GPAC is not ready: {refusal}")
    return paths, v2e.images_of_the_cycle(paths["images_root"])


def test_images_are_those_of_cycle_3(v2e, setup):
    paths, images = setup
    run3 = json.loads((paths["cache_root"] / "cycles" / "3" / "v2a_vtm.json").read_text(encoding="utf-8"))
    v2e.check_against_cycle_3(images, run3)
    assert [i["code"] for i in images] == [i["code"] for i in run3["images"]]


def test_run_on_one_pair(v2e, setup, tmp_path):
    paths, images = setup
    state = {"core_commit": "test", "core_tree": "test"}
    run_path, book_path = v2e.run(paths, images, [("syn21x13", 32)], [("syn21x13", 32)], tmp_path, state)
    record = json.loads(run_path.read_text(encoding="utf-8"))
    (row,) = record["rows"]
    assert (row["reproduces_cycle_3"], row["status"], row["round_trip"], row["ispe_ok"]) == (True, "valid", True, True)
    assert row["heif_bytes"] == row["stream_bytes"] + row["overhead_bytes"] and row["mdat_bytes"] > 0
    assert row["level_status"] == row["frame_rate_status"] == "approved"
    assert record["repack"][0]["equal"] and record["summary"]["verdict_v2e"] == "pass"
    assert record["gpac_writes"]["gpac_version"] == "26.08-DEV-rev171-gf9f54240-master"
    with pytest.raises(v2e.RunRefusal, match=v2e.RUN_FILE):
        v2e.run(paths, images, [("syn21x13", 32)], [], tmp_path, state)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Таблица"].cell(row=3, column=12)  # HEIF bytes
    assert cell.value == row["heif_bytes"]
    cell.value += 1
    changed = tmp_path / "changed.xlsx"
    book.save(changed)
    with pytest.raises(v2e.RunRefusal, match=r"Таблица!R3C12 .* differs from the run file"):
        v2e.check_book(run_path, changed)
