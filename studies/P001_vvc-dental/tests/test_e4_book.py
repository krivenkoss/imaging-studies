"""Lab tests of e4_book.py (study P001, cycle 11) on a synthetic run record - no coder needed.

The record holds two fragment units and the mosaic, four codecs, valid, lossless and error rows. The book is
built and read back; the tests check the sheets, 13 and 9 scatter charts with logarithmic CR axes, drop-down
lists, series references inside the tables; the formulas of «Просмотр» and «Обзор» are evaluated by the
test's own parser of their two templates for two units (and two levels, every metric) and compared with the
record; «Обзор-данные» is recomputed by the test's own linear percentile, level by level; lossless points give
no chart value and are counted. The journal: resume, torn line, a failed task runs again (its last entry
counts), another header refused.
"""

import importlib.util
import json
import re
from pathlib import Path

import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def book_module():
    spec = importlib.util.spec_from_file_location("e4_book", STUDY / "e4_book.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


PARAMS = {"vtm": [-4, 22, 40], "jpeg": [10, 50, 100], "heif": [0, 50, 100], "avif": [0, 50, 100]}
UNITS = {"baboon": "f", "frisco": "f", "mosaic": "m"}


def synthetic_rows():
    rows = []
    for u, (unit, level) in enumerate(UNITS.items()):
        for c, codec in enumerate(PARAMS):
            for param in PARAMS[codec]:
                base = 20 + 3 * u + c + param / 10
                row = {"key": f"{unit}|{codec}|{param}", "unit": unit, "level": level, "codec": codec, "param": param,
                       "status": "valid", "file_bytes": 1000 + param, "cr_file": 50.0 / (1 + param / 50 + u),
                       "bpp": 0.1 * (1 + param), "psnr_db": base, "psnr_hvs_m_db": base + 5, "haarpsi": 0.5 + param / 1000,
                       "seconds": 0.1, "seconds_metrics": 0.01, "file_sha256": f"{u}{c}{param}", "detail": None}
                if codec == "avif" and param == 100:
                    row |= {"status": "lossless", "psnr_db": "+inf", "psnr_hvs_m_db": "+inf", "haarpsi": 1.0}
                if unit == "frisco" and codec == "heif" and param == 50:
                    row |= {"status": "error", "detail": "RuntimeError: test", "file_bytes": None, "cr_file": None,
                            "bpp": None, "psnr_db": None, "psnr_hvs_m_db": None, "haarpsi": None}
                rows.append(row)
    return rows


@pytest.fixture(scope="module")
def built(book_module, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("book")
    rows = synthetic_rows()
    record = {"run": "e4_dry.json", "cycle": 11, "command": "python test", "question": "test", "rows": rows,
              "units": UNITS, "overview": book_module.overview(rows), "summary_lines": [["Строк, шт", len(rows)]]}
    run_path, book_path = tmp / "e4_dry.json", tmp / "e4_dry.xlsx"
    run_path.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
    describe = lambda rec: [("Где что", "тест")]

    def check(rp, bp):
        book_module.check_book(rp, bp, describe, AssertionError)

    book_module.build_book(run_path, book_path, describe, check)
    from openpyxl import load_workbook
    return record, run_path, book_path, load_workbook(book_path)


def test_sheets_charts_axes_and_lists(book_module, built):
    _, _, _, wb = built
    assert wb.sheetnames == list(book_module.SHEETS) and wb.calculation.fullCalcOnLoad
    view, over = wb["Просмотр"], wb["Обзор"]
    assert [type(c).__name__ for c in view._charts] == ["ScatterChart"] * 13
    assert [type(c).__name__ for c in over._charts] == ["ScatterChart"] * 9
    log_x = lambda c: c.x_axis.scaling.logBase == 10
    log_y = lambda c: c.y_axis.scaling.logBase == 10
    assert [log_x(c) for c in view._charts] == [False] * 4 + [True] * 3 + [False] * 4 + [False] * 2
    assert [log_y(c) for c in view._charts] == [False] * 4 + [False] * 3 + [True] * 4 + [False] * 2
    assert [log_x(c) for c in over._charts] == [False] * 4 + [True] + [False] * 4
    assert [log_y(c) for c in over._charts] == [False] * 5 + [True] * 4
    assert all(len(c._charts) == 2 for c in view._charts[:4])  # HaarPSI on a second axis
    validations = {(ws.title, str(d.sqref)): d.formula1 for ws in (view, over) for d in ws.data_validations.dataValidation}
    assert validations == {("Просмотр", "B1"): "'Списки'!$A$2:$A$4", ("Обзор", "B1"): "'Списки'!$B$2:$B$3",
                           ("Обзор", "B2"): "'Списки'!$C$2:$C$4"}
    lists = wb["Списки"]
    assert [lists.cell(r, 1).value for r in (2, 3, 4)] == list(UNITS)
    assert [lists.cell(r, 3).value for r in (2, 3, 4)] == ["PSNR", "PSNR-HVS-M", "HaarPSI"]


def test_series_references_inside_the_tables(book_module, built):
    record, _, _, wb = built
    for sheet, layout in (("Просмотр", book_module.view_layout(record["rows"])),
                          ("Обзор", book_module.over_layout(record["rows"]))):
        allowed_rows = set(range(book_module.TABLE_ROW, book_module.TABLE_ROW + 3)) | {4, 5}
        for chart in wb[sheet]._charts:
            for part in [chart, *getattr(chart, "_charts", [])]:
                for s in part.series:
                    for ref in (s.xVal.numRef.f, s.yVal.numRef.f):
                        m = re.fullmatch(r"'(.+)'!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", ref)
                        assert m and m[1] == sheet and m[2] == m[4], ref
                        rows = set(range(int(m[3]), int(m[5]) + 1))
                        assert rows <= allowed_rows, ref


VIEW_RE = re.compile(r"""=IFERROR\(IF\(INDEX\('Строки'!\$([A-Z]+)\$3:\$[A-Z]+\$(\d+),MATCH\(\$B\$1&"\|(\w+)\|"&\$([A-Z]+)(\d+),'Строки'!\$A\$3:\$A\$\d+,0\)\)="valid",INDEX\('Строки'!\$([A-Z]+)\$3:\$[A-Z]+\$\d+,MATCH\(\$B\$1&"\|\w+\|"&\$[A-Z]+\d+,'Строки'!\$A\$3:\$A\$\d+,0\)\),NA\(\)\),NA\(\)\)""")
OVER_RE = re.compile(r"""=IFERROR\(INDEX\('Обзор-данные'!\$([A-Z]+)\$3:\$([A-Z]+)\$(\d+),MATCH\(\$B\$1&"\|(\w+)\|"&\$([A-Z]+)(\d+),'Обзор-данные'!\$A\$3:\$A\$\d+,0\),MATCH\((.+),'Обзор-данные'!\$[A-Z]+\$2:\$[A-Z]+\$2,0\)\),NA\(\)\)""")
COUNT_RE = re.compile(r"""="(.+): "&COUNTIFS\('Строки'!\$B\$3:\$B\$\d+,\$B\$1,'Строки'!\$D\$3:\$D\$\d+,"(\w+)",'Строки'!\$F\$3:\$F\$\d+,"lossless"\)""")
NA = object()


def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def evaluate_view(ws, rows_sheet, formula, unit):
    """The test's own evaluation of the «Просмотр» template: lookup of unit|codec|param in «Строки»."""
    m = VIEW_RE.fullmatch(formula)
    assert m, formula
    status_col, codec, param_col, param_row, value_col = col_index(m[1]), m[3], m[4], int(m[5]), col_index(m[6])
    key = f"{unit}|{codec}|{ws[f'{param_col}{param_row}'].value}"
    for r in range(3, rows_sheet.max_row + 1):
        if rows_sheet.cell(r, 1).value == key:
            return rows_sheet.cell(r, value_col).value if rows_sheet.cell(r, status_col).value == "valid" else NA
    return NA


def test_view_formulas_evaluated_for_two_units(book_module, built):
    record, _, _, wb = built
    view, rows_sheet = wb["Просмотр"], wb["Строки"]
    by_key = {r["key"]: r for r in record["rows"]}
    heads = {"PSNR, дБ": "psnr_db", "PSNR-HVS-M, дБ": "psnr_hvs_m_db", "HaarPSI, безразм.": "haarpsi", "CR, безразм.": "cr_file"}
    checked = 0
    for unit in ("baboon", "frisco"):
        for codec, (c0, params) in book_module.view_layout(record["rows"]).items():
            for i, param in enumerate(params):
                r = book_module.TABLE_ROW + i
                for j in range(1, 5):
                    field = heads[view.cell(book_module.TABLE_ROW - 1, c0 + j).value]
                    got = evaluate_view(view, rows_sheet, view.cell(r, c0 + j).value, unit)
                    row = by_key.get(f"{unit}|{codec}|{param}")
                    want = row[field] if row and row["status"] == "valid" else NA
                    want = float(f"{want:.16g}") if isinstance(want, float) else want  # as the cell stores a float
                    assert got is want or got == want, (unit, codec, param, field)
                    checked += 1
    assert checked == 2 * 12 * 4
    counts = {}
    for k in range(4):
        m = COUNT_RE.fullmatch(view.cell(3, 2 + k).value)
        counts[m[2]] = sum(1 for r in record["rows"] if r["unit"] == "frisco" and r["codec"] == m[2] and r["status"] == "lossless")
    assert counts == {"vtm": 0, "jpeg": 0, "heif": 0, "avif": 1}  # the lossless point is counted, not plotted
    assert evaluate_view(view, rows_sheet, view.cell(book_module.TABLE_ROW + 2, 1 + 3 * 6 + 1).value, "frisco") is NA


def percentile(values, q):
    """Linear percentile, written again: position (n - 1) * q / 100 between neighbours."""
    v = sorted(values)
    pos = (len(v) - 1) * q / 100
    lo = int(pos)
    hi = min(lo + 1, len(v) - 1)
    return v[lo] + (v[hi] - v[lo]) * (pos - lo)


def own_overview(rows, level, codec, param):
    valid = [r for r in rows if (r["level"], r["codec"], r["param"], r["status"]) == (level, codec, param, "valid")]
    out = {"n": len(valid)}
    for field, name in (("psnr_db", "PSNR"), ("psnr_hvs_m_db", "PSNR-HVS-M"), ("haarpsi", "HaarPSI"), ("cr_file", "CR")):
        for label, q in (("медиана", 50), ("q25", 25), ("q75", 75)):
            out[f"{name}|{label}"] = percentile([r[field] for r in valid], q) if valid else None
    return out


def test_overview_data_recomputed_level_by_level(book_module, built):
    record, _, _, wb = built
    data = wb["Обзор-данные"]
    header = {data.cell(2, c).value: c for c in range(1, data.max_column + 1)}
    for r in range(3, data.max_row + 1):
        level_name, codec, param = data.cell(r, 1).value.split("|")
        level = {v: k for k, v in book_module.LEVELS.items()}[level_name]
        own = own_overview(record["rows"], level, codec, int(param))
        for key, value in own.items():
            got = data.cell(r, header[key]).value
            assert (got is None and value is None) or got == pytest.approx(value, abs=1e-12), (level, codec, param, key)


def evaluate_over(sheet_over, data, formula, level, metric):
    m = OVER_RE.fullmatch(formula)
    assert m, formula
    first, codec, param_col, param_row, column_key = col_index(m[1]), m[4], m[5], int(m[6]), m[7]
    key = f"{level}|{codec}|{sheet_over[f'{param_col}{param_row}'].value}"
    column_key = re.fullmatch(r'(?:\$B\$2&"(.+)"|"(.+)")', column_key)
    column_name = metric + column_key[1] if column_key[1] else column_key[2]
    header = {data.cell(2, c).value: c for c in range(first, data.max_column + 1)}
    for r in range(3, data.max_row + 1):
        if data.cell(r, 1).value == key:
            value = data.cell(r, header[column_name]).value
            return NA if value is None else value
    return NA


def test_overview_formulas_evaluated_for_both_levels_and_every_metric(book_module, built):
    record, _, _, wb = built
    over, data = wb["Обзор"], wb["Обзор-данные"]
    heads = ["медиана", "q25", "q75", "n единиц, шт", "CR медиана", "CR q25", "CR q75"]
    for level_code, level in book_module.LEVELS.items():
        for metric in ("PSNR", "PSNR-HVS-M", "HaarPSI"):
            for codec, (c0, params) in book_module.over_layout(record["rows"]).items():
                assert [over.cell(book_module.TABLE_ROW - 1, c0 + j).value for j in range(1, 8)] == heads
                for i, param in enumerate(params):
                    r = book_module.TABLE_ROW + i
                    own = own_overview(record["rows"], level_code, codec, param)
                    if not any(x["level"] == level_code and x["codec"] == codec and x["param"] == param
                               for x in record["rows"]):
                        continue
                    want = [own[f"{metric}|медиана"], own[f"{metric}|q25"], own[f"{metric}|q75"], own["n"],
                            own["CR|медиана"], own["CR|q25"], own["CR|q75"]]
                    got = [evaluate_over(over, data, over.cell(r, c0 + j).value, level, metric) for j in range(1, 8)]
                    for g, w in zip(got, want):
                        assert (g is NA and w is None) or g == pytest.approx(w, abs=1e-12), (level, metric, codec, param)


def test_book_check_refuses_a_changed_cell(book_module, built, tmp_path):
    _, run_path, book_path, _ = built
    from openpyxl import load_workbook
    wb = load_workbook(book_path)
    wb["Обзор-данные"].cell(3, 7).value = 999.0
    wb.save(tmp_path / "changed.xlsx")
    with pytest.raises(AssertionError, match=r"Обзор-данные!R3C7"):
        book_module.check_book(run_path, tmp_path / "changed.xlsx", lambda rec: [("Где что", "тест")], AssertionError)


def test_journal_resume_torn_line_and_another_header(book_module, tmp_path):
    path = tmp_path / "j.jsonl"
    header = {"declared": {"a": 1}, "plan": {"vtm": [1, 2]}, "core_commit": "x"}
    j = book_module.Journal(path, header)
    j.append("t1", [{"k": 1}])
    j.append("t2", [{"k": 2}])
    with open(path, "a", encoding="utf-8") as f:
        f.write('{"task": "t3", "rows": [')  # a hard stop in the middle of a line
    j2 = book_module.Journal(path, header)
    assert (j2.resumed, j2.torn_lines, sorted(j2.done)) == (2, 1, ["t1", "t2"])
    j2.append("t3", [{"k": 3}])
    assert sorted(book_module.Journal(path, header).done) == ["t1", "t2", "t3"]
    j3 = book_module.Journal(path, header)
    j3.append("t4", [{"k": 4, "status": "error"}])  # a failed call, or one cut by a stop
    j4 = book_module.Journal(path, header)
    assert "t4" not in j4.done and j4.last["t4"][0]["status"] == "error" and j4.to_retry == 1
    j4.append("t4", [{"k": 4, "status": "valid"}])
    j5 = book_module.Journal(path, header)
    assert j5.done["t4"][0]["status"] == "valid" and j5.to_retry == 0 and j5.last["t4"] == j5.done["t4"]
    with pytest.raises(book_module.JournalRefusal, match=r"differs in \['plan'\] \{'plan': \['vtm'\]\}"):
        book_module.Journal(path, header | {"plan": {"vtm": [1, 2, 3]}})
