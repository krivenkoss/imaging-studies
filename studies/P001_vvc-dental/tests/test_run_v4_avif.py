"""Lab tests of run_v4_avif.py (study P001, cycle 9): Frisco, lab_paths.json with the B12 environment and the run
files of cycles 7 and 8 needed.

- a run on one small configuration (Frisco, sigma^2 = 100, seed 1, QF 34, 35, 36, 80): the threads step
  passes (2 = 4 = 8), the regression gives 4 of 4 rows of cycle 8 variant (d), the V4 verdict is "incomplete"
  (sigma^2 25 and 196 have no data), the book equals the run file and refuses a changed cell, the second run
  refuses by name;
- the regression stops on a changed row of cycle 8 and names it; the verdict counts both sides.
"""

import importlib.util
import json
from pathlib import Path

import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def v4():
    spec = importlib.util.spec_from_file_location("run_v4_avif", STUDY / "run_v4_avif.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def setup(v4):
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    try:
        paths = v4.load_paths()
    except v4.RunRefusal as refusal:
        pytest.skip(f"V4 AVIF on B12 is not ready: {refusal}")
    for cycle, name in ((7, "v3b_v4.json"), (8, "avif_diag.json")):
        if not (paths["cache_root"] / "cycles" / str(cycle) / name).is_file():
            pytest.skip(f"the run file of cycle {cycle} is absent on this machine")
    return paths, v4.load_test_images(paths["images_root"], ["frisco.bmp"], (512, 512))


def test_run_on_one_small_configuration(v4, setup, tmp_path):
    paths, images = setup
    plan = {"sigma2": (100,), "seeds": (1,), "qualities": (34, 35, 36, 80)}
    state = {"core_commit": "test", "core_tree": "test"}
    run_path, book_path = v4.run(paths, images, plan, tmp_path, state)
    record = json.loads(run_path.read_text(encoding="utf-8"))
    assert record["steps"]["threads"]["verdict"] == "pass" and record["steps"]["threads"]["same_2_4_8"]
    assert record["steps"]["regression"]["rows_equal"] == 4 and record["steps"]["regression"]["verdict"] == "pass"
    assert record["stopped"] is None and len(record["rows"]) == 4
    assert all(r["status"] == "valid" and r["av1c_monochrome"] == 1 for r in record["rows"])
    assert record["verdict_v4_avif"]["verdict"] == "incomplete" and record["verdict_v4_avif"]["missing_sigma2"] == [25, 196]
    with pytest.raises(v4.RunRefusal, match=v4.RUN_FILE):
        v4.run(paths, images, plan, tmp_path, state)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Строки"].cell(row=3, column=8)  # file bytes of the first row
    assert cell.value == record["rows"][0]["file_bytes"]
    cell.value += 1
    book.save(tmp_path / "changed.xlsx")
    with pytest.raises(v4.RunRefusal, match=r"Строки!R3C8 .* differs from the run file"):
        v4.check_book(run_path, tmp_path / "changed.xlsx")


def test_regression_stops_and_names_the_row(v4):
    cycle8 = {"rows": [{"variant": "d", "sigma2": 25, "seed": 1, "qf": q, "file_sha256": f"{q:064x}", "psnr_tc_db": 30.0 + q}
                       for q in (10, 20)] + [{"variant": "a", "sigma2": 25, "seed": 1, "qf": 10, "file_sha256": "x",
                                              "psnr_tc_db": 0.0}]}
    ours = [{"sigma2": 25, "seed": 1, "qf": q, "file_sha256": f"{q:064x}", "psnr_tc_db": 30.0 + q} for q in (10, 20)]
    assert v4.regression(ours, cycle8)["verdict"] == "pass"
    ours[1] = ours[1] | {"psnr_tc_db": 50.0000001}
    got = v4.regression(ours, cycle8)
    assert (got["verdict"], got["rows_equal"], got["differ_first"]) == ("stop", 1, ["s2 25 seed 1 QF 20"])


def test_verdict_both_sides(v4):
    def entry(s2, inside, within):
        return {"criterion": "V4 AVIF", "sigma2": s2, "in_tolerance": inside and within}
    assert v4.verdict_of([entry(s, True, True) for s in (25, 100, 196)], (25, 100, 196))["verdict"] == "pass"
    failed = v4.verdict_of([entry(25, True, True), entry(100, True, False), entry(196, True, True)], (25, 100, 196))
    assert failed["verdict"] == "fail" and failed["failed"] == ["frisco s2 100"]
