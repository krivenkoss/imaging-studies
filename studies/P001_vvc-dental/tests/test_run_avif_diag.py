"""Lab tests of run_avif_diag.py (study P001, cycle 8): Frisco, lab_paths.json and the files of cycle 7 needed.

- the OOP rule (largest PSNR_tc, the smaller QF on a tie) against hand-computed cases;
- a pinned file with another sha256 is refused by name;
- a run on one small configuration (sigma^2 = 25, seed 1, four qualities; variants (a), (b) and, when the
  libheif venv is configured, (d)): variant (a) equals the rows of cycle 7 file by file (the regression),
  every file is monochrome, the run refuses by name the second time, and the book check refuses a
  changed cell. Coding runs in this process (workers = 1).
"""

import importlib.util
import json
from pathlib import Path

import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def diag():
    spec = importlib.util.spec_from_file_location("run_avif_diag", STUDY / "run_avif_diag.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def setup(diag):
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    paths = diag.load_paths()
    cycle7 = paths["cache_root"] / "cycles" / "7" / "v3b_v4.json"
    if not cycle7.is_file():
        pytest.skip("the run file of cycle 7 is absent on this machine")
    (image,) = diag.load_test_images(paths["images_root"], [diag.IMAGE], (512, 512))
    return paths, image


def test_oop_largest_psnr_smaller_qf_on_tie(diag):
    assert diag.oop([(10, 30.0), (20, 31.0), (30, 31.0), (40, 29.0)]) == 20
    assert diag.oop([(5, 1.0)]) == 5 and diag.oop([]) is None


def test_pinned_file_with_other_sha256_refused(diag, setup, tmp_path):
    paths, _ = setup
    copy = tmp_path / "v3b_v4_reference.json"
    copy.write_bytes(paths["v3b_v4_reference"].read_bytes() + b"\n")
    with pytest.raises(diag.RunRefusal, match=r"V3-b/V4 reference v3b_v4_reference\.json: sha256 \w+ differs from the pinned"):
        diag.load_pinned(copy, diag.REFERENCE_SHA256, "V3-b/V4 reference")
    assert diag.load_pinned(paths["v3b_v4_reference"], diag.REFERENCE_SHA256, "V3-b/V4 reference")["rs24_oop_by_psnr"]


def test_run_on_one_small_configuration(diag, setup, tmp_path):
    paths, image = setup
    variants = ("a", "b", "d") if paths["avif_libheif_python"] else ("a", "b")
    plan = {"sigma2": (25,), "seeds": (1,), "qualities": (49, 50, 51, 100), "variants": variants}
    state = {"core_commit": "test", "core_tree": "test"}
    run_path, book_path = diag.run(paths, image, plan, tmp_path, state, workers=1)
    record = json.loads(run_path.read_text(encoding="utf-8"))
    assert record["unavailable"] == {}
    assert len(record["rows"]) == 4 * len(variants) and all(r["status"] == "valid" for r in record["rows"])
    assert all(r["av1c_monochrome"] == 1 for r in record["rows"])
    assert record["regression_a"] == {"rows_compared": 4, "rows_equal": 4, "differ_first": [], "verdict": "pass"}
    assert {s["variant"] for s in record["summary"]} == set(variants)
    if "d" in variants:
        assert record["variant_d_environment"]["pillow_heif"] == "0.22.0"
    with pytest.raises(diag.RunRefusal, match=diag.RUN_FILE):
        diag.run(paths, image, plan, tmp_path, state, workers=1)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Строки"].cell(row=3, column=8)  # file bytes of the first row
    assert cell.value == record["rows"][0]["file_bytes"]
    cell.value += 1
    changed = tmp_path / "changed.xlsx"
    book.save(changed)
    with pytest.raises(diag.RunRefusal, match=r"Строки!R3C8 .* differs from the run file"):
        diag.check_book(run_path, changed)
