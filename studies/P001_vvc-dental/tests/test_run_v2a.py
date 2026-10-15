"""Lab tests of run_v2a.py (study P001, cycle 3): VTM build, lab_paths.json and test images needed.

- a run on one small pair writes its run file and a book equal to it cell by cell, and refuses
  by name the second time;
- the cell check of the book refuses a book that does not match the run file;
- the images of the cycle are the declared ones (no dental pixels).
"""

import importlib.util
import json
from pathlib import Path

import pytest

pytestmark = pytest.mark.lab

STUDY = Path(__file__).resolve().parents[1]
ROOT = STUDY.parents[1]


@pytest.fixture(scope="module")
def v2a():
    if not (ROOT / "lab_paths.json").is_file():
        pytest.skip("lab_paths.json is absent: the laboratory is not configured on this machine")
    spec = importlib.util.spec_from_file_location("run_v2a", STUDY / "run_v2a.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def paths(v2a):
    try:
        return v2a.load_paths()
    except v2a.RunRefusal as refusal:
        pytest.skip(f"VTM is not ready: {refusal}")


@pytest.fixture(scope="module")
def images(v2a, paths):
    return v2a.images_of_the_cycle(paths["images_root"])


def test_images_of_the_cycle(v2a, images):
    assert [i["code"] for i in images] == ["e0", "baboon", "barbara", "goldhill", "lenna", "frisco", "diego",
                                           "syn8x8", "syn21x13", "syn517x389"]
    assert [i["array"].shape for i in images[-3:]] == [(8, 8), (21, 13), (517, 389)]
    assert {i["class"] for i in images} == {"control", "test", "synthetic"}
    assert (images[-1]["array"][:21, :13] == images[-2]["array"]).all()  # crops of one recipe frame


def test_run_on_one_pair_and_refusal_by_name(v2a, paths, images, tmp_path):
    state = {"core_commit": "test", "core_tree": "test"}
    run_path, book_path = v2a.run(paths, images, [("syn21x13", 32)], [("syn21x13", 32)], tmp_path, state)
    record = json.loads(run_path.read_text(encoding="utf-8"))
    assert record["path_control"]["v2a_match"] and record["path_control"]["bits"] > 0
    (row,) = record["rows"]
    assert (row["status"], row["v2a_match"], row["shape_ok"], row["h"], row["w"]) == ("valid", True, True, 21, 13)
    assert row["level_status"] == row["frame_rate_status"] == "provisional"
    assert record["summary"]["verdict_v2a"] == record["summary"]["verdict_determinism"] == "pass"
    assert record["build_manifest"]["tag"] == "VTM-23.0" and book_path.is_file()
    assert record["timeout_s"]["value"] == record["wrapper_provenance"]["timeout_s"] == v2a.TIMEOUT_S["value"]

    # V2-a reproduced: the same bitstream and recon as the run file of cycle 3 (read, never rewritten)
    cycle3 = json.loads((paths["cache_root"] / "cycles" / "3" / v2a.RUN_FILE).read_text(encoding="utf-8"))
    recorded = {(r["image_code"], r["qp"]): r for r in cycle3["rows"]}
    for now in (record["path_control"], row):
        then = recorded[now["image_code"], now["qp"]]
        assert (now["bitstream_sha256"], now["recon_sha256"]) == (then["bitstream_sha256"], then["recon_sha256"])

    with pytest.raises(v2a.RunRefusal, match=v2a.RUN_FILE):
        v2a.run(paths, images, [("syn21x13", 32)], [], tmp_path, state)

    from openpyxl import load_workbook
    book = load_workbook(book_path)
    cell = book["Таблица"].cell(row=3, column=9)  # bits of the pair
    assert cell.value == row["bits"]
    cell.value += 8
    changed = tmp_path / "changed.xlsx"
    book.save(changed)
    with pytest.raises(v2a.RunRefusal, match=r"Таблица!R3C9 .* differs from the run file"):
        v2a.check_book(run_path, changed)
