"""P001, B11: how many 512 x 512 tiles of the 16 images keep share_any (background and frame) at or below 10, 25, 50 %? Run: python studies/P001_vvc-dental/run_p001_tiles.py

Study P001, cycle 9, stage E3: contract ../REFERENCE.md#b11 0.2.0 through imstudies.lab.p001_tiles. Reads
the copy of B10 (data_root), lays the grid centred in every image (origin "center", passed explicitly;
cycle 8 used "top_left"), measures four shares per full tile, counts tiles per image and over the set,
writes <cache_root>/cycles/9/p001_tiles.json (refuses by name if it exists), the book p001_tiles.xlsx and
the pages for the owner's eyes in p001_tiles_pages/ beside it (local, outside Git). Prints codes and
counts only. The counts are a measurement for the owner's choice, not an exclusion rule.
`--rebuild-book` rebuilds the book from the run file and carries the owner's verdicts by page file.
"""

import argparse
import datetime
import hashlib
import json
import platform
import sys
from pathlib import Path

import numpy as np

import imstudies
from imstudies.lab import p001_tiles as tiles
from imstudies.lab.runs import RunRefusal, git_state, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
COMMAND = "python studies/P001_vvc-dental/run_p001_tiles.py"
QUESTION = ("P001, B11: сколько фрагментов 512 × 512 у каждого снимка и в наборе имеют долю фона и рамки share_any "
            "не больше 10, 25 и 50 %?")
CYCLE, RUN_FILE, BOOK_FILE, PAGES_DIR = 9, "p001_tiles.json", "p001_tiles.xlsx", "p001_tiles_pages"
ORIGIN = "center"  # B11 0.2.0, owner decision 2026-10-20, 19:21 UTC
DECLARED = {
    "data": "Копия B10 0.2.0 (data/P001, реестр); роль check у всех снимков; пороги на них не подбираются.",
    "measurements": (f"share_white — доля пикселей ≥ {tiles.WHITE}; share_black — ≤ {tiles.BLACK}; share_flat — доля "
                     f"блоков {tiles.BLOCK} × {tiles.BLOCK} сетки фрагмента со стандартным отклонением < {tiles.FLAT_STD} "
                     f"(numpy std, float64, ddof 0); share_any — доля пикселей хотя бы в одной маске, пиксель плоского "
                     f"блока — целиком. Пороги — параметры измерения (B11), ПРЕДПОЛОЖЕНИЕ: рамка насыщена, фон однороден."),
    "selection": "В книге — число фрагментов с share_any ≤ 10, 25, 50 % по снимку и по набору; это не правило "
                 "исключения — его выбирает владелец по счёту.",
    "grid": "Сетка 512 × 512 по центру снимка (B11 0.2.0, origin = center): x0 = ⌊(W − 512·n_c)/2⌋, "
            "y0 = ⌊(H − 512·n_r)/2⌋, лишний пиксель нечётного остатка — справа и снизу; без перекрытия, только полные "
            "фрагменты; остатки по обе стороны не используются; число фрагментов снимка обязано равняться full_tiles "
            "реестра B10.",
    "pages": "Страницы для глаз: одна на снимок (все снимки) — уменьшенная картинка (среднее блоков 4 × 4) с сеткой и "
             "share_any в каждом фрагменте; локально, вне Git, ассистенту не передаются; вердикт — лист «Страницы».",
    "claim": "Утверждение страницы собирается из строки снимка файла прогона: число фрагментов, счёт при 10 / 25 / 50 % и "
             "наибольшая доля; сверяется со строкой перед записью.",
}


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {}
    for key in ("data_root", "cache_root"):
        if not raw.get(key):
            raise RunRefusal(f"lab_paths.json has no key '{key}'")
        paths[key] = (ROOT / raw[key]).resolve()
    return paths


def title_of(entry):
    return (f"{entry['code']} — ортопантомограмма, {entry['n_tiles']} полных фрагментов 512 × 512; страница — по правилу "
            f"«одна на снимок» (B11)")


def run(data_root, out_dir, state, log=print):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path, pages_dir = out_dir / RUN_FILE, out_dir / BOOK_FILE, out_dir / PAGES_DIR
    for path in (run_path, pages_dir):
        if path.exists():
            raise RunRefusal(f"{path} already exists: a run of cycle {CYCLE} is never overwritten")
    started = datetime.datetime.now(datetime.timezone.utc)
    log(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        log(f"  [{key}] {text}")
    registry, images = tiles.load_copy(data_root)
    rows = tiles.measure(images, origin=ORIGIN)
    per = tiles.per_image(registry, rows, origin=ORIGIN)
    totals = tiles.set_totals(per)
    pages_dir.mkdir()
    pages = []
    for img, entry in zip(images, per):
        claim = tiles.statement(entry)
        tiles.check_statement(claim, entry)
        name = f"{entry['code']}.png"
        tiles.draw_page(img["values"], [r for r in rows if r["code"] == entry["code"]], title_of(entry), claim,
                        pages_dir / name)
        pages.append({"code": entry["code"], "file": f"{PAGES_DIR}/{name}", "title": title_of(entry), "claim": claim,
                      "explanation": " ".join(tiles.EXPLANATION)})
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        **state, "contract": tiles.CONTRACT, "declared_before_numbers": DECLARED,
        "parameters": {"origin": ORIGIN, "tile": tiles.TILE, "block": tiles.BLOCK, "white": tiles.WHITE, "black": tiles.BLACK,
                       "flat_std": tiles.FLAT_STD, "selection_shares": list(tiles.SELECTION), "page_scale": tiles.PAGE_SCALE},
        "registry": {"file": "data/registry/P001.json", "sha256": sha256_file(Path(data_root) / "registry" / "P001.json"),
                     "contract": registry.get("contract"), "core_commit": registry.get("core_commit")},
        "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__, "numpy": np.__version__,
                        "platform": platform.platform()},
        "unit_of_analysis": "row = tile DNN_rRR_cCC; counts per image and over the set",
        "rows": rows, "per_image": per, "set": totals, "pages": pages,
    }
    text = json.dumps(record, ensure_ascii=False, indent=1, allow_nan=False)
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        f.write(text)
    build_book(run_path, book_path)
    for e in per:
        log(f"  {e['code']}: tiles {e['n_tiles']}, share_any <= 10/25/50 %: {e['n_any_le_10']} / {e['n_any_le_25']} / "
            f"{e['n_any_le_50']}")
    log(f"Set: {totals['n_images']} images, {totals['n_tiles']} tiles; share_any <= 10/25/50 %: {totals['n_any_le_10']} / "
        f"{totals['n_any_le_25']} / {totals['n_any_le_50']} tiles, images with at least one such tile: "
        f"{totals['images_with_tiles_n_any_le_10']} / {totals['images_with_tiles_n_any_le_25']} / "
        f"{totals['images_with_tiles_n_any_le_50']}; best achievable: {totals['best_achievable_tiles']} tiles of "
        f"{totals['best_achievable_images']} images")
    log(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}\nPages: {pages_dir} (local)")
    return run_path, book_path


PER = [("Код", "code"), ("H, пикс.", "H"), ("W, пикс.", "W"), ("Положение сетки", "origin"), ("x0, пикс.", "x0"),
       ("y0, пикс.", "y0"), ("Фрагментов всего, шт", "n_tiles"),
       ("share_any ≤ 10 %, шт", "n_any_le_10"), ("share_any ≤ 25 %, шт", "n_any_le_25"),
       ("share_any ≤ 50 %, шт", "n_any_le_50"), ("Наибольшая share_any, %", "max_share_any")]
ROWS = [("Фрагмент", "tile"), ("Код", "code"), ("Строка сетки", "r"), ("Столбец сетки", "c"), ("x, пикс.", "x"),
        ("y, пикс.", "y"), ("share_white, %", "share_white"), ("share_black, %", "share_black"),
        ("share_flat, %", "share_flat"), ("share_any, %", "share_any")]
SET = [("Снимков, шт", "n_images"), ("Фрагментов всего, шт", "n_tiles"), ("share_any ≤ 10 %, шт", "n_any_le_10"),
       ("share_any ≤ 25 %, шт", "n_any_le_25"), ("share_any ≤ 50 %, шт", "n_any_le_50"),
       ("Снимков хотя бы с одним фрагментом ≤ 10 %, шт", "images_with_tiles_n_any_le_10"),
       ("Снимков хотя бы с одним фрагментом ≤ 25 %, шт", "images_with_tiles_n_any_le_25"),
       ("Снимков хотя бы с одним фрагментом ≤ 50 %, шт", "images_with_tiles_n_any_le_50"),
       ("Наилучший достижимый итог: фрагментов, шт", "best_achievable_tiles"),
       ("Наилучший достижимый итог: снимков, шт", "best_achievable_images")]
PAGES = [("Код", "code"), ("Страница", "file"), ("Заголовок", "title"), ("Утверждение", "claim"),
         ("На пальцах", "explanation")]
VERDICT_COLUMN = len(PAGES) + 1
PERCENT = {"max_share_any", "share_white", "share_black", "share_flat", "share_any"}


def book_cells(record, run_sha256):
    cells, formats = {}, {}
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
             ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
             ("Контракт", record["contract"]),
             *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
             ("Реестр B10", f"{record['registry']['file']}, sha256 {record['registry']['sha256']}, "
                            f"{record['registry']['contract']}"),
             ("Ключи листов", "под шапкой — имя поля файла прогона: «Таблица» — per_image[] и set, «Фрагменты» — rows[], "
                              "«Страницы» — pages[]; колонка «Вердикт» — владельца, в файле прогона её нет"),
             ("Слова вердикта", "ок · не тот класс · ложная тревога · пропуск (LAB-PAGES, п. 6)"),
             ("На пальцах", "Каждый снимок разбит на квадраты 512 × 512; в каждом квадрате посчитано, какая его часть "
                            "похожа на белую рамку, чёрный фон или ровный участок; книга считает, сколько квадратов "
                            "останется при разных допусках такой части.")]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v

    def table(sheet, columns, records, first_row=1):
        for j, (head, key) in enumerate(columns, start=1):
            cells[sheet, first_row, j], cells[sheet, first_row + 1, j] = head, key
        for i, rec in enumerate(records):
            for j, (_, key) in enumerate(columns, start=1):
                cells[sheet, first_row + 2 + i, j] = rec.get(key)
                if key in PERCENT:
                    formats[sheet, first_row + 2 + i, j] = "0.0%"
        return first_row + 2 + len(records)

    end = table("Таблица", PER, record["per_image"])
    cells["Таблица", end + 1, 1] = "Набор (set)"
    table("Таблица", SET, [record["set"]], first_row=end + 2)
    table("Фрагменты", ROWS, record["rows"])
    table("Страницы", PAGES, record["pages"])
    cells["Страницы", 1, VERDICT_COLUMN], cells["Страницы", 2, VERDICT_COLUMN] = "Вердикт", "владелец"
    return cells, formats


def build_book(run_path, book_path, verdicts=None):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    record = json.loads(run_bytes.decode("utf-8"))
    cells, formats = book_cells(record, hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Таблица", "Фрагменты", "Страницы"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    for (sheet, r, c), fmt in formats.items():
        wb[sheet].cell(row=r, column=c).number_format = fmt
    pages = wb["Страницы"]
    for i, page in enumerate(record["pages"], start=3):
        pages.cell(row=i, column=2).hyperlink = page["file"]
        if verdicts and page["file"] in verdicts:
            pages.cell(row=i, column=VERDICT_COLUMN, value=verdicts[page["file"]])
    wb["Описание"].column_dimensions["A"].width = 34
    wb["Описание"].column_dimensions["B"].width = 140
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for name in ("Таблица", "Фрагменты", "Страницы"):
        for c in wb[name][1]:
            c.font = Font(bold=True)
        wb[name].freeze_panes = "A3"
    n = len(record["per_image"])
    if n:
        chart = BarChart()
        chart.title = "Фрагментов на снимок: всего и при share_any ≤ 10, 25, 50 %"
        chart.y_axis.title, chart.x_axis.title = "фрагментов, шт", "код снимка"
        chart.x_axis.delete = chart.y_axis.delete = False
        sheet = wb["Таблица"]
        chart.add_data(Reference(sheet, min_col=7, max_col=10, min_row=1, max_row=2 + n), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=2 + n))
        chart.width, chart.height = 28, 12
        sheet.add_chart(chart, "M2")
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g); only the verdict column may hold other values."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())[0]
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        extra = {e for e in extra - set(cells) if not (sheet == "Страницы" and e[2] == VERDICT_COLUMN and e[1] >= 3)}
        if extra:
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra)[:3]}")


def read_verdicts(book_path):
    """Owner's verdicts of an existing book by page file."""
    from openpyxl import load_workbook

    sheet = load_workbook(book_path)["Страницы"]
    return {sheet.cell(row=r, column=2).value: sheet.cell(row=r, column=VERDICT_COLUMN).value
            for r in range(3, sheet.max_row + 1) if sheet.cell(row=r, column=VERDICT_COLUMN).value not in (None, "")}


def rebuild_book(run_path, book_path):
    """Rebuild from the run file, carrying verdicts; a page with a verdict that is gone is refused by name."""
    verdicts = read_verdicts(book_path) if Path(book_path).is_file() else {}
    pages = {p["file"] for p in json.loads(Path(run_path).read_text(encoding="utf-8"))["pages"]}
    gone = sorted(set(verdicts) - pages)
    if gone:
        raise RunRefusal(f"pages with a verdict are gone from the run file: {gone}")
    missing = sorted(f for f in verdicts if not (Path(run_path).parent / f).is_file())
    if missing:
        raise RunRefusal(f"pages with a verdict are gone from the disk: {missing}")
    build_book(run_path, book_path, verdicts)
    return verdicts


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book, carrying the owner's verdicts")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            carried = rebuild_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell, {len(carried)} verdicts carried: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        run(paths["data_root"], out_dir, state)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
