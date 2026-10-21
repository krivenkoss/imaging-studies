"""E4 timing: how many hours will the main run of P001 take, measured on test images (not on P001)? Run: python studies/P001_vvc-dental/run_e4_timing.py

Study P001, cycle 10. Times the codec wrappers and metrics on two test images: (f) Frisco 512 x 512, the size
of a fragment, and (m) a mosaic 1504 x 2773, the size of a whole radiograph, made of the six 512 x 512 test
images of Images/BMP. VTM (B05) with HEIF packing (B09) - (f) QP -12..63, (m) QP 0, 22, 32, 42 - once in one
process and once by a pool of 4 coder processes (files of the pool must equal the sequential ones by sha256,
otherwise the run stops); JPEG (B06), HEIF (B07), AVIF (B12, batch calls) sequentially; PSNR, PSNR-HVS-M,
HaarPSI (B02-B04) on image / JPEG QF 50, three repeats. The estimate of the main run follows the formula
declared before the numbers. No image of P001 is read. Writes <cache_root>/cycles/10/e4_timing.json (refuses by
name if it exists) and the book e4_timing.xlsx beside it; `--rebuild-book` rebuilds the book from the run file.
"""

import argparse
import concurrent.futures
import datetime
import hashlib
import json
import os
import platform
import statistics
import sys
import tempfile
import time
from pathlib import Path

import numpy as np

import imstudies
from imstudies.core import metrics
from imstudies.lab.coders import avif_libheif, heif, jpeg, vtm, vvc_heif
from imstudies.lab.images import load_test_images
from imstudies.lab.runs import RunRefusal, git_state, json_number, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
COMMAND = "python studies/P001_vvc-dental/run_e4_timing.py"
QUESTION = ("E4: сколько часов займёт основной прогон P001 (154 фрагмента и 16 целых снимков × 4 кодека × все значения "
            "параметра × 3 метрики) — оценка по времени вызовов на тестовых изображениях?")
CYCLE, RUN_FILE, BOOK_FILE = 10, "e4_timing.json", "e4_timing.xlsx"
MOSAIC_ORDER = ["BABOON.BMP", "BARBARA.BMP", "goldhill.bmp", "lenna.bmp", "frisco.bmp", "DIEGO.BMP"]
MOSAIC_SHAPE = (1504, 2773)
LEVEL, FRAME_RATE = "5", 1  # B05, owner 2026-10-14 18:13 UTC
TIMEOUT_S = {"f": 600, "m": 3600}
POOL = 4
N_FRAGMENTS, N_WHOLE = 154, 16
N_VALUES = {"vtm": 76, "jpeg": 100, "heif": 101, "avif": 101}  # full ranges of B05, B06, B07, B12
FULL_PLAN = {
    "vtm": {"f": tuple(range(vtm.QP_MIN, vtm.QP_MAX + 1)), "m": (0, 22, 32, 42)},
    "jpeg": {"f": tuple(range(1, 101)), "m": (1, *range(10, 101, 10))},
    "heif": {"f": tuple(range(0, 101)), "m": tuple(range(0, 101, 10))},
    "avif": {"f": tuple(range(0, 101)), "m": tuple(range(0, 101, 10))},
    "metric_repeats": 3, "mosaic_shape": MOSAIC_SHAPE,
}
DECLARED = {
    "images": ("Тестовые, в P001 не входят: (ф) Frisco 512 × 512 — размер фрагмента; (м) мозаика 1504 × 2773 — 3 строки × 6 "
               "столбцов изображений 512 × 512 из Images/BMP в порядке baboon, barbara, goldhill, lenna, frisco, diego; "
               "строка r сдвинута циклически на r влево (в столбце c строки r — изображение (c + r) mod 6), обрезка от "
               "левого верхнего угла; sha256 пикселей обоих — в файле прогона. Снимки P001 и data/P001 не читаются."),
    "vtm": ("VTM (B05, Level 5, FrameRate 1) с упаковкой в HEIF (B09): (ф) QP −12…63 (76 значений), (м) QP 0, 22, 32, 42; "
            "timeout_s = 600 (ф) и 3600 (м). Каждый набор — дважды: последовательно (1 процесс) и пулом из 4 (4 потока "
            "Python, каждый запускает свой EncoderApp, DecoderApp, MP4Box, gpac — одновременно 4 процесса кодера; "
            "у каждого вызова свой временный каталог обёртки). Проверка: sha256 потока и файла HEIF пула = "
            "последовательному прогону, иначе стоп."),
    "still": ("JPEG (B06): (ф) QF 1…100, (м) QF 1, 10, …, 100; HEIF (B07): (ф) 0…100, (м) 0, 10, …, 100; AVIF (B12): те же, "
              "одним пакетным вызовом на набор — секунды значения = стена пакетного вызова / число значений; "
              "последовательно."),
    "metrics": "PSNR (B02), PSNR-HVS-M (B03), HaarPSI (B04) на паре «изображение — JPEG QF 50», (ф) и (м), по 3 повтора.",
    "row": ("Строка — изображение, кодек, параметр, режим (1 или 4), секунды вызова (стена; у VTM — B05 + B09, и отдельно), "
            "размер файла; у пула — ещё стена всего набора. PSNR и CR тестовых изображений печатаются, результатом не являются."),
    "estimate": ("Время = Σ по кодекам (154 · n_знач · t_ф + 16 · n_знач · t_м) / k, где t — медиана секунд вызова "
                 "последовательного режима, n_знач — полный диапазон (VTM 76, JPEG 100, HEIF и AVIF 101); для VTM (м) "
                 "слагаемое 16 · Σ по QP −12…63 t_м(QP), t_м — линейная интерполяция 4 точек по QP, вне [0; 42] — "
                 "ближайшая; k — ускорение пула 4 = стена 1 процесса / стена пула на наборе VTM (ф) (k на наборе (м) "
                 "печатается рядом); метрики — 3 на каждый файл: (154 · Σn · Σt_метр,ф + 16 · Σn · Σt_метр,м), Σn = 378, "
                 "t_метр — медиана 3 повторов. Часы для 1 процесса (k = 1) и для 4 (k измеренное); доля VTM. "
                 "ПРЕДПОЛОЖЕНИЕ: время VTM на тестовых изображениях близко к времени на снимках P001 — содержимое "
                 "влияет на время кодера (риск называется рядом с числом)."),
    "pages": "Страницы для глаз не применяются — замер времени.",
}


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {}
    for key in ("images_root", "cache_root", "avif_libheif_python"):
        if not raw.get(key) or raw[key].startswith("<"):
            raise RunRefusal(f"lab_paths.json has no key '{key}' or it is the example placeholder")
        paths[key] = (ROOT / raw[key]).resolve()
    for codec, keys in (("vtm", ("encoder", "decoder", "cfg")), ("gpac", ("mp4box", "gpac"))):
        for key in keys:
            value = (raw.get("coders", {}).get(codec) or {}).get(key)
            if not value:
                raise RunRefusal(f"lab_paths.json has no key coders.{codec}.{key}")
            paths[key] = (ROOT / value).resolve()
            if not paths[key].is_file():
                raise RunRefusal(f"{codec} {key} not found: {paths[key]}")
    return paths


def mosaic(tiles, shape):
    """3 rows x 6 columns of the 512 x 512 tiles, row r shifted left by r, cropped from the top-left corner."""
    rows = [np.concatenate([tiles[(c + r) % len(tiles)] for c in range(len(tiles))], axis=1) for r in range(3)]
    return np.ascontiguousarray(np.concatenate(rows, axis=0)[:shape[0], :shape[1]])


def vtm_call(task):
    """B05 then B09 for one QP; seconds of each and of both, sha256 of stream and HEIF file."""
    paths, image, qp, timeout_s, work = task
    t0 = time.perf_counter()
    coded = vtm.encode_decode(image, qp, max_pixels=image.size, level=LEVEL, frame_rate=FRAME_RATE, timeout_s=timeout_s,
                              encoder=paths["encoder"], decoder=paths["decoder"], cfg=paths["cfg"], work_root=work)
    t1 = time.perf_counter()
    packed = vvc_heif.pack(coded.bitstream, coded.recon, timeout_s=timeout_s, mp4box=paths["mp4box"],
                           gpac=paths["gpac"], decoder=paths["decoder"], work_root=work)
    t2 = time.perf_counter()
    return {"param": qp, "seconds": t2 - t0, "seconds_b05": t1 - t0, "seconds_b09": t2 - t1,
            "file_bytes": len(packed.file), "file_sha256": hashlib.sha256(packed.file).hexdigest(),
            "bitstream_sha256": hashlib.sha256(coded.bitstream).hexdigest(),
            "psnr_db": metrics.psnr(image, coded.recon, max_pixels=image.size).value}


def cpu_busy():
    """System CPU times now (psutil, if installed) - for the mean load over a set."""
    try:
        import psutil
    except ImportError:
        return None
    return psutil.cpu_times()


def busy_share(before, after):
    if before is None or after is None:
        return None
    total = sum(after) - sum(before)
    return None if total <= 0 else 1 - (after.idle - before.idle) / total


def vtm_set(paths, size, image, qps, work, workers):
    """One VTM set: sequential (workers = 1) or a pool of `workers` threads, each running its own coder processes."""
    tasks = [(paths, image, qp, TIMEOUT_S[size], work) for qp in qps]
    before, t0 = cpu_busy(), time.perf_counter()
    if workers == 1:
        results = [vtm_call(t) for t in tasks]
    else:
        with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as pool:
            results = list(pool.map(vtm_call, tasks))
    wall = time.perf_counter() - t0
    load = busy_share(before, cpu_busy())
    h, w = image.shape
    return [r | {"image": size, "codec": "vtm", "mode": workers, "set_wall_seconds": wall, "cpu_busy_share": load,
                 "cr_file": h * w / r["file_bytes"]} for r in results], wall


def still_set(paths, size, image, codec, params, work):
    h, w = image.shape
    rows = []
    if codec == "avif":
        t0 = time.perf_counter()
        coded = avif_libheif.encode_decode(image, list(params), max_pixels=image.size, timeout_s=TIMEOUT_S[size],
                                           python=paths["avif_libheif_python"], work_root=work)
        each = (time.perf_counter() - t0) / len(params)
        pairs = [(q, c, each) for q, c in zip(params, coded)]
    else:
        module = {"jpeg": jpeg, "heif": heif}[codec]
        pairs = []
        for q in params:
            t0 = time.perf_counter()
            c = module.encode_decode(image, q, max_pixels=image.size)
            pairs.append((q, c, time.perf_counter() - t0))
    for q, c, seconds in pairs:
        rows.append({"image": size, "codec": codec, "param": q, "mode": 1, "seconds": seconds,
                     "file_bytes": len(c.file), "file_sha256": hashlib.sha256(c.file).hexdigest(),
                     "cr_file": h * w / len(c.file), "psnr_db": metrics.psnr(image, c.recon, max_pixels=image.size).value})
    return rows


def metric_rows(size, image, repeats):
    recon = jpeg.encode_decode(image, 50, max_pixels=image.size).recon
    rows = []
    for name, fn in (("psnr", metrics.psnr), ("psnr_hvs_m", metrics.psnr_hvs_m), ("haarpsi", metrics.haarpsi)):
        for k in range(1, repeats + 1):
            t0 = time.perf_counter()
            value = fn(image, recon, max_pixels=image.size).value
            rows.append({"image": size, "codec": "metric", "param": name, "repeat": k, "mode": 1,
                         "seconds": time.perf_counter() - t0, "value": value})
    return rows


def interp(points, x):
    """Linear interpolation of (x, y) points sorted by x; outside the points the nearest end value."""
    if x <= points[0][0]:
        return points[0][1]
    if x >= points[-1][0]:
        return points[-1][1]
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        if x1 <= x <= x2:
            return y1 + (y2 - y1) * (x - x1) / (x2 - x1)
    raise ValueError(x)


def estimate(rows, k_f, k_m):
    """The declared formula: hours for 1 and for 4 processes, terms per codec and the share of VTM."""
    seq = [r for r in rows if r["mode"] == 1]

    def t(codec, size):
        values = [r["seconds"] for r in seq if (r["codec"], r["image"]) == (codec, size)]
        return statistics.median(values) if values else None

    terms = {}
    for codec, n in N_VALUES.items():
        frag = N_FRAGMENTS * n * t(codec, "f")
        if codec == "vtm":
            points = sorted({r["param"]: r["seconds"] for r in seq if (r["codec"], r["image"]) == ("vtm", "m")}.items())
            whole = N_WHOLE * sum(interp(points, qp) for qp in range(vtm.QP_MIN, vtm.QP_MAX + 1))
        else:
            whole = N_WHOLE * n * t(codec, "m")
        terms[codec] = {"t_f_median_s": t(codec, "f"), "t_m_median_s": t(codec, "m"), "n_values": n,
                        "fragments_s": frag, "whole_s": whole, "total_s": frag + whole}
    n_files = sum(N_VALUES.values())
    t_metric = {size: sum(statistics.median([r["seconds"] for r in seq if (r["codec"], r["image"], r["param"]) ==
                                             ("metric", size, name)]) for name in ("psnr", "psnr_hvs_m", "haarpsi"))
                for size in ("f", "m")}
    metric_s = N_FRAGMENTS * n_files * t_metric["f"] + N_WHOLE * n_files * t_metric["m"]
    total_1 = sum(v["total_s"] for v in terms.values()) + metric_s
    return {"terms": terms, "metrics": {"t_f_sum_of_medians_s": t_metric["f"], "t_m_sum_of_medians_s": t_metric["m"],
                                        "files": n_files, "total_s": metric_s},
            "k_pool_f": k_f, "k_pool_m": k_m, "total_1_process_s": total_1, "total_4_processes_s": total_1 / k_f,
            "hours_1_process": total_1 / 3600, "hours_4_processes": total_1 / k_f / 3600,
            "vtm_share": terms["vtm"]["total_s"] / total_1}


def machine():
    out = {"processor": platform.processor(), "logical_cpus": os.cpu_count(), "platform": platform.platform()}
    try:
        import psutil
        out |= {"physical_cpus": psutil.cpu_count(logical=False),
                "memory_total_gib": round(psutil.virtual_memory().total / 2 ** 30, 1)}
    except ImportError:
        out |= {"physical_cpus": None, "memory_total_gib": None}
    return out


def _jsonable(obj):
    if isinstance(obj, float):
        return json_number(obj)
    if isinstance(obj, dict):
        return {k: _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    return obj


def run(paths, images, plan, out_dir, state):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path = out_dir / RUN_FILE, out_dir / BOOK_FILE
    if run_path.exists():
        raise RunRefusal(f"{run_path} already exists: a run file of cycle {CYCLE} is never overwritten")
    with tempfile.TemporaryDirectory(dir=paths["cache_root"], prefix="work_e4_timing_") as work:  # removed at the end
        return _run(paths, images, plan, run_path, book_path, state, Path(work))


def _run(paths, images, plan, run_path, book_path, state, work):
    started, t0 = datetime.datetime.now(datetime.timezone.utc), time.perf_counter()
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")
    rows, walls, stopped = [], {}, None
    for size in ("f", "m"):
        image = images[size]
        seq, walls[size, 1] = vtm_set(paths, size, image, plan["vtm"][size], work, 1)
        pool, walls[size, POOL] = vtm_set(paths, size, image, plan["vtm"][size], work, POOL)
        rows += seq + pool
        same = all(a["file_sha256"] == b["file_sha256"] and a["bitstream_sha256"] == b["bitstream_sha256"]
                   for a, b in zip(seq, pool))
        print(f"  VTM ({size}): {len(seq)} QP, sequential {walls[size, 1]:.1f} s, pool of {POOL} {walls[size, POOL]:.1f} s, "
              f"pool = sequential by sha256: {same}")
        if not same:
            stopped = f"VTM ({size}): files of the pool differ from the sequential run"
            break
    if not stopped:
        for size in ("f", "m"):
            for codec in ("jpeg", "heif", "avif"):
                got = still_set(paths, size, images[size], codec, plan[codec][size], work)
                rows += got
                print(f"  {codec} ({size}): {len(got)} values, median {statistics.median(r['seconds'] for r in got):.4f} s")
            rows += metric_rows(size, images[size], plan["metric_repeats"])
    for k, r in enumerate(rows, start=1):
        r["row_index"] = k
    k_f = walls["f", 1] / walls["f", POOL] if ("f", POOL) in walls else None
    k_m = walls["m", 1] / walls["m", POOL] if ("m", POOL) in walls else None
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "wall_seconds": round(time.perf_counter() - t0, 1), **state, "declared_before_numbers": DECLARED,
        "plan": {k: ({s: [min(v), max(v), len(v)] for s, v in p.items()} if isinstance(p, dict) else p)
                 for k, p in plan.items()},
        "machine": machine(), "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__,
                                              "numpy": np.__version__},
        "images": {s: {"shape": list(images[s].shape), "pixels_sha256": hashlib.sha256(images[s].tobytes()).hexdigest()}
                   for s in images},
        "set_walls": [{"image": s, "mode": m, "seconds": v} for (s, m), v in sorted(walls.items())],
        "pool_equals_sequential": stopped is None, "stopped": stopped, "rows": rows,
        "estimate": None if stopped else estimate(rows, k_f, k_m),
    }
    text = json.dumps(_jsonable(record), ensure_ascii=False, indent=1, allow_nan=False)
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        f.write(text)
    build_book(run_path, book_path)
    e = record["estimate"]
    if e:
        print(f"Pool speed-up k: (f) {k_f:.2f}, (m) {k_m:.2f}; estimate: {e['hours_1_process']:.1f} h in 1 process, "
              f"{e['hours_4_processes']:.1f} h in 4; VTM share {e['vtm_share']:.1%}")
    print(f"Stopped: {stopped}" if stopped else "Pool = sequential by sha256: yes")
    print(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}")
    return run_path, book_path


SUMMARY_KEYS = ("codec", "image", "mode", "n", "median_s", "min_s", "max_s", "set_wall_s")


def summary_rows(record):
    groups = {}
    for r in record["rows"]:
        key = (r["codec"], r["image"], r["mode"]) if r["codec"] != "metric" else (f"metric {r['param']}", r["image"], 1)
        groups.setdefault(key, []).append(r)
    walls = {(w["image"], w["mode"]): w["seconds"] for w in record["set_walls"]}
    out = []
    for (codec, image, mode), rs in sorted(groups.items()):
        s = [r["seconds"] for r in rs]
        out.append({"codec": codec, "image": image, "mode": mode, "n": len(rs), "median_s": statistics.median(s),
                    "min_s": min(s), "max_s": max(s), "set_wall_s": walls.get((image, mode)) if codec == "vtm" else None})
    return out


ROWS = [("№", "row_index"), ("Изображение (f/m)", "image"), ("Кодек или метрика", "codec"), ("Параметр", "param"),
        ("Режим, процессов", "mode"), ("Секунды вызова, с", "seconds"), ("B05, с", "seconds_b05"), ("B09, с", "seconds_b09"),
        ("Стена набора, с", "set_wall_seconds"), ("Файл, байт", "file_bytes"), ("CR, безразм.", "cr_file"),
        ("PSNR, дБ (не результат)", "psnr_db"), ("Повтор", "repeat"), ("sha256 файла", "file_sha256")]
SUMMARY = [("Кодек", "codec"), ("Размер (f/m)", "image"), ("Режим, процессов", "mode"), ("Вызовов, шт", "n"),
           ("Медиана, с", "median_s"), ("Наименьшее, с", "min_s"), ("Наибольшее, с", "max_s"),
           ("Стена набора, с", "set_wall_s")]


def book_cells(record, run_sha256):
    cells = {}
    e = record["estimate"]
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
             ("Закончен, UTC", record["finished_utc"]), ("Время прогона, с", record["wall_seconds"]),
             ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
             *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
             ("Машина", json.dumps(record["machine"], ensure_ascii=False)),
             ("Пул = последовательно по sha256", "да" if record["pool_equals_sequential"] else "нет"),
             ("Остановлен", record["stopped"] or "нет"),
             ("Ключи листов", "под шапкой — имя поля: «Оценка» — estimate, «Сводка» — свод rows[] по кодеку, размеру и "
                              "режиму (медиана, наименьшее, наибольшее), «Строки» — rows[]"),
             ("Листы «Рисунки» и «Страницы»", "не строятся: страницы не применяются (объявлено до чисел)"),
             ("На пальцах", "Каждый кодек и каждая метрика запущены на тестовом снимке размера фрагмента и на снимке "
                            "размера целой панорамы; по их времени формула прикидывает, сколько часов займёт основной "
                            "прогон в один и в четыре процесса.")]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v
    if e:
        est = [("Ускорение пула k (ф), раз", "k_pool_f", e["k_pool_f"]), ("Ускорение пула (м), раз", "k_pool_m", e["k_pool_m"])]
        for codec, term in e["terms"].items():
            est += [(f"{codec}: t_ф медиана, с", f"terms.{codec}.t_f_median_s", term["t_f_median_s"]),
                    (f"{codec}: t_м медиана, с", f"terms.{codec}.t_m_median_s", term["t_m_median_s"]),
                    (f"{codec}: n_знач", f"terms.{codec}.n_values", term["n_values"]),
                    (f"{codec}: фрагменты, с", f"terms.{codec}.fragments_s", term["fragments_s"]),
                    (f"{codec}: целые снимки, с", f"terms.{codec}.whole_s", term["whole_s"])]
        est += [("Метрики: Σ медиан (ф), с", "metrics.t_f_sum_of_medians_s", e["metrics"]["t_f_sum_of_medians_s"]),
                ("Метрики: Σ медиан (м), с", "metrics.t_m_sum_of_medians_s", e["metrics"]["t_m_sum_of_medians_s"]),
                ("Метрики: всего, с", "metrics.total_s", e["metrics"]["total_s"]),
                ("Оценка, 1 процесс, ч", "hours_1_process", e["hours_1_process"]),
                ("Оценка, 4 процесса, ч", "hours_4_processes", e["hours_4_processes"]),
                ("Доля VTM, безразм.", "vtm_share", e["vtm_share"])]
        cells["Оценка", 1, 1], cells["Оценка", 1, 2], cells["Оценка", 1, 3] = "Величина", "Ключ estimate", "Значение"
        for i, (label, key, value) in enumerate(est, start=2):
            cells["Оценка", i, 1], cells["Оценка", i, 2], cells["Оценка", i, 3] = label, key, value
    for sheet, columns, records in (("Сводка", SUMMARY, summary_rows(record)), ("Строки", ROWS, record["rows"])):
        for j, (head, key) in enumerate(columns, start=1):
            cells[sheet, 1, j], cells[sheet, 2, j] = head, key
        for i, rec in enumerate(records):
            for j, (_, key) in enumerate(columns, start=1):
                cells[sheet, 3 + i, j] = rec.get(key)
    return cells


def build_book(run_path, book_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Оценка", "Сводка", "Строки"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    wb["Описание"].column_dimensions["A"].width = 34
    wb["Описание"].column_dimensions["B"].width = 140
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for name in ("Оценка", "Сводка", "Строки"):
        for c in wb[name][1]:
            c.font = Font(bold=True)
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g, as openpyxl writes them); no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        tiles = [i["array"] for i in load_test_images(paths["images_root"], MOSAIC_ORDER, (512, 512))]
        images = {"f": tiles[MOSAIC_ORDER.index("frisco.bmp")], "m": mosaic(tiles, FULL_PLAN["mosaic_shape"])}
        run(paths, images, FULL_PLAN, out_dir, state)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
