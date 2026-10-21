"""E4 book of study P001 and the resume journal: shared by the dry run (cycle 11) and the main run E4.

Contract: STUDY.md section 7 ("Книга E4") and ../../REFERENCE.md#lab-experiment (chain, point 3; sheets
«Просмотр» and «Обзор»). The book holds the rows of the run file («Строки»), overview data per level, codec
and parameter («Обзор-данные»: median and quartiles over the units of the level, n), lists for the drop-down
choices («Списки»), and two sheets of native Excel charts driven by INDEX/MATCH formulas: «Просмотр» (one unit
chosen from a list: metric = f(parameter), metric = f(CR), CR = f(parameter), metric against metric) and «Обзор»
(level and metric chosen from lists: median and quartiles). No dynamic arrays - the book opens in Excel 2016;
full recalculation on load. Every static cell is checked against the run file; formula cells are checked as
text. Lossless rows (MSE = 0) keep their values in «Строки» but give no chart point (NA) and are counted under
the charts. Shared lab code lives here until a second study needs it (then src/, by its own contract).
"""

import hashlib
import json
import os
import threading
from pathlib import Path

import numpy as np

CODECS = ("vtm", "jpeg", "heif", "avif")
CODEC_NAMES = {"vtm": "VVC (VTM)", "jpeg": "JPEG", "heif": "HEIF", "avif": "AVIF"}
PARAM_NAMES = {"vtm": "QP", "jpeg": "качество (quality)", "heif": "качество (quality)", "avif": "качество (quality)"}
LEVELS = {"f": "фрагменты", "m": "снимки"}
METRICS = (("psnr_db", "PSNR", "PSNR, дБ"), ("psnr_hvs_m_db", "PSNR-HVS-M", "PSNR-HVS-M, дБ"),
           ("haarpsi", "HaarPSI", "HaarPSI, безразм."))
QUANTITIES = (*((key, name) for key, name, _ in METRICS), ("cr_file", "CR"))  # overview columns
STATS = (("median", "медиана"), ("q25", "q25"), ("q75", "q75"))
THRESHOLDS_DB = (40, 41, 42)
ROWS = (("Ключ «единица|кодек|параметр»", "key"), ("Единица", "unit"), ("Уровень", "level"), ("Кодек", "codec"),
        ("Параметр", "param"), ("Статус", "status"), ("Файл, байт", "file_bytes"), ("CR, безразм.", "cr_file"),
        ("bpp, бит/пикс.", "bpp"), ("PSNR, дБ", "psnr_db"), ("PSNR-HVS-M, дБ", "psnr_hvs_m_db"),
        ("HaarPSI, безразм.", "haarpsi"), ("Кодер, с", "seconds"), ("Метрики, с", "seconds_metrics"),
        ("sha256 файла", "file_sha256"), ("Причина сбоя", "detail"))
ROW_COL = {key: chr(ord("A") + i) for i, (_, key) in enumerate(ROWS)}


# ---------------------------------------------------------------- resume journal (LAB-EXPERIMENT, chain, 3)

class JournalRefusal(Exception):
    """The journal cannot be continued; the message says why in words."""


def _canonical(obj):
    return json.dumps(obj, ensure_ascii=False, sort_keys=True)


class Journal:
    """Append-only journal: the first line is the header (declared, plan, commit, versions, images); then one line
    per finished task {"task": id, "rows": [...]}. A run started again with an equal header skips the finished
    tasks; another header is refused in words naming what differs. A torn last line (a hard stop while writing)
    is ignored and counted."""

    def __init__(self, path, header):
        self.path, self.header, self.lock = Path(path), header, threading.Lock()
        self.done, self.entries, self.torn_lines = {}, [], 0
        if self.path.is_file():
            lines = self.path.read_text(encoding="utf-8").split("\n")
            old = json.loads(lines[0])
            if _canonical(old) != _canonical(header):
                differ = sorted(k for k in set(old) | set(header) if _canonical(old.get(k)) != _canonical(header.get(k)))
                inner = {k: sorted(s for s in set(old.get(k) or {}) | set(header.get(k) or {})
                                   if _canonical((old.get(k) or {}).get(s)) != _canonical((header.get(k) or {}).get(s)))
                         for k in differ if isinstance(old.get(k), dict) and isinstance(header.get(k), dict)}
                raise JournalRefusal(f"the journal {self.path.name} was started with another header: differs in "
                                     f"{differ} {inner} - finish or remove it by hand, or restore the same setup")
            for line in lines[1:]:
                if not line.strip():
                    continue
                try:
                    entry = json.loads(line)
                except json.JSONDecodeError:
                    self.torn_lines += 1
                    continue
                self.done[entry["task"]] = entry["rows"]
                self.entries.append(entry)
            if self.torn_lines:  # rewrite without the torn line, so the next appends start on a fresh line
                good = [lines[0]] + [ln for ln in lines[1:] if ln.strip() and _loads_ok(ln)]
                self.path.write_text("\n".join(good) + "\n", encoding="utf-8")
        else:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            with open(self.path, "x", encoding="utf-8", newline="\n") as f:
                f.write(_canonical(header) + "\n")
        self.resumed = len(self.done)

    def append(self, task, rows, **extra):
        """One finished task: its rows (and extra fields, e.g. the wall time of a finished pool)."""
        line = json.dumps({"task": task, "rows": rows, **extra}, ensure_ascii=False, allow_nan=False)
        with self.lock:
            with open(self.path, "a", encoding="utf-8", newline="\n") as f:
                f.write(line + "\n")
                f.flush()
                os.fsync(f.fileno())
            self.done[task] = rows
            self.entries.append({"task": task, "rows": rows, **extra})


def _loads_ok(line):
    try:
        json.loads(line)
        return True
    except json.JSONDecodeError:
        return False


# ---------------------------------------------------------------- overview data

def overview(rows):
    """Per (level, codec, param): n valid units, n lossless, and median, q25, q75 of each quantity over the valid
    units of that level only (numpy.percentile, linear)."""
    groups = {}
    for r in rows:
        groups.setdefault((r["level"], r["codec"], r["param"]), []).append(r)
    out = []
    for (level, codec, param), rs in sorted(groups.items(), key=lambda kv: (kv[0][0], CODECS.index(kv[0][1]), kv[0][2])):
        valid = [r for r in rs if r["status"] == "valid"]
        entry = {"key": f"{LEVELS[level]}|{codec}|{param}", "level": LEVELS[level], "codec": codec, "param": param,
                 "n": len(valid), "n_lossless": sum(r["status"] == "lossless" for r in rs)}
        for key, name in QUANTITIES:
            values = np.array([r[key] for r in valid], dtype=np.float64)
            for stat, label in STATS:
                q = {"median": 50, "q25": 25, "q75": 75}[stat]
                entry[f"{name}|{label}"] = float(np.percentile(values, q)) if values.size else None
        out.append(entry)
    return out


OVERVIEW_COLUMNS = (("Ключ «уровень|кодек|параметр»", "key"), ("Уровень", "level"), ("Кодек", "codec"),
                    ("Параметр", "param"), ("n единиц, шт", "n"), ("Без потерь, шт", "n_lossless"),
                    *((f"{name} {label}", f"{name}|{label}") for _, name in QUANTITIES for _, label in STATS))


# ---------------------------------------------------------------- the book

def _col(n):
    """1-based column number -> letters."""
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _q(sheet):
    return "'" + sheet + "'"


def params_of(rows, codec):
    return sorted({r["param"] for r in rows if r["codec"] == codec})


def view_formula(codec, param_cell, quantity, last):
    """«Просмотр»: the value of `quantity` of the unit in $B$1 for (codec, param), NA unless the row is valid."""
    rng = lambda key: f"{_q('Строки')}!${ROW_COL[key]}$3:${ROW_COL[key]}${last}"
    match = f'MATCH($B$1&"|{codec}|"&{param_cell},{rng("key")},0)'
    return f'=IFERROR(IF(INDEX({rng("status")},{match})="valid",INDEX({rng(quantity)},{match}),NA()),NA())'


def overview_formula(codec, param_cell, column_key, last, n_cols):
    """«Обзор»: a column of «Обзор-данные» for (level in $B$1, codec, param); column_key may name the metric of $B$2."""
    first, end = _col(5), _col(n_cols)
    table = f"{_q('Обзор-данные')}!${first}$3:${end}${last}"
    header = f"{_q('Обзор-данные')}!${first}$2:${end}$2"
    keys = f"{_q('Обзор-данные')}!$A$3:$A${last}"
    return (f'=IFERROR(INDEX({table},MATCH($B$1&"|{codec}|"&{param_cell},{keys},0),'
            f'MATCH({column_key},{header},0)),NA())')


def _table_cells(cells, sheet, columns, records):
    for j, (head, key) in enumerate(columns, start=1):
        cells[sheet, 1, j], cells[sheet, 2, j] = head, key
    for i, rec in enumerate(records):
        for j, (_, key) in enumerate(columns, start=1):
            cells[sheet, 3 + i, j] = rec.get(key)


VIEW_QUANTITIES = (("psnr_db", "PSNR, дБ"), ("psnr_hvs_m_db", "PSNR-HVS-M, дБ"), ("haarpsi", "HaarPSI, безразм."),
                   ("cr_file", "CR, безразм."))
VIEW_BLOCK = 6  # param + 4 quantities + gap
OVER_BLOCK = 9  # param, median, q25, q75, n, CR median, CR q25, CR q75, gap
TABLE_ROW = 8   # first data row of the tables on «Просмотр» and «Обзор» (row 7 - headers)


def view_layout(rows):
    """Columns of the codec blocks and the data rows of «Просмотр»: {codec: (first column, [params])}."""
    return {codec: (1 + k * VIEW_BLOCK, params_of(rows, codec)) for k, codec in enumerate(CODECS)}


def over_layout(rows):
    return {codec: (1 + k * OVER_BLOCK, params_of(rows, codec)) for k, codec in enumerate(CODECS)}


def book_cells(record, run_sha256, describe):
    """All cells of the book: static values and formula texts, {(sheet, row, column): value}."""
    cells = {}
    rows = record["rows"]
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), *describe(record)]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v
    for i, (k, v) in enumerate(record["summary_lines"], start=1):
        cells["Сводка", i, 1], cells["Сводка", i, 2] = k, v
    _table_cells(cells, "Строки", ROWS, rows)
    over = record["overview"]
    _table_cells(cells, "Обзор-данные", OVERVIEW_COLUMNS, over)
    units = list(record["units"])
    cells["Списки", 1, 1], cells["Списки", 1, 2], cells["Списки", 1, 3] = "Единица", "Уровень", "Метрика"
    for i, u in enumerate(units, start=2):
        cells["Списки", i, 1] = u
    for i, level in enumerate(LEVELS.values(), start=2):
        cells["Списки", i, 2] = level
    for i, (_, name, _) in enumerate(METRICS, start=2):
        cells["Списки", i, 3] = name
    last_row, last_over = 2 + len(rows), 2 + len(over)
    # «Просмотр»
    cells["Просмотр", 1, 1], cells["Просмотр", 1, 2] = "Единица (выбор из списка):", units[0]
    cells["Просмотр", 2, 1] = "Одна единица анализа; точки без потерь (MSE = 0) на графиках не показываются."
    cells["Просмотр", 3, 1] = "Без потерь, точек нет:"
    for k, codec in enumerate(CODECS):
        cells["Просмотр", 3, 2 + k] = (f'="{CODEC_NAMES[codec]}: "&COUNTIFS({_q("Строки")}!$B$3:$B${last_row},$B$1,'
                                       f'{_q("Строки")}!$D$3:$D${last_row},"{codec}",'
                                       f'{_q("Строки")}!$F$3:$F${last_row},"lossless")')
    for codec, (c0, params) in view_layout(rows).items():
        cells["Просмотр", TABLE_ROW - 2, c0] = CODEC_NAMES[codec]
        cells["Просмотр", TABLE_ROW - 1, c0] = PARAM_NAMES[codec]
        for j, (_, head) in enumerate(VIEW_QUANTITIES, start=1):
            cells["Просмотр", TABLE_ROW - 1, c0 + j] = head
        for i, p in enumerate(params):
            r = TABLE_ROW + i
            cells["Просмотр", r, c0] = p
            for j, (quantity, _) in enumerate(VIEW_QUANTITIES, start=1):
                cells["Просмотр", r, c0 + j] = view_formula(codec, f"${_col(c0)}{r}", quantity, last_row)
    # threshold lines of PSNR-HVS-M = f(CR): x from the smallest to the largest CR of the unit
    cr_ranges = [f"${_col(c0 + 4)}${TABLE_ROW}:${_col(c0 + 4)}${TABLE_ROW + len(params) - 1}"
                 for c0, params in view_layout(rows).values()]
    cells["Просмотр", 4, 1], cells["Просмотр", 5, 1] = "Пороги PSNR-HVS-M: CR от", "CR до"
    cells["Просмотр", 4, 2] = f"=MIN({','.join(cr_ranges)})"
    cells["Просмотр", 5, 2] = f"=MAX({','.join(cr_ranges)})"
    for k, t in enumerate(THRESHOLDS_DB):
        cells["Просмотр", 4, 3 + k], cells["Просмотр", 5, 3 + k] = t, t
    # «Обзор»
    cells["Обзор", 1, 1], cells["Обзор", 1, 2] = "Уровень (выбор из списка):", LEVELS["f"]
    cells["Обзор", 2, 1], cells["Обзор", 2, 2] = "Метрика (выбор из списка):", METRICS[1][1]
    cells["Обзор", 3, 1] = "Обзор, не результат: медиана и квартили по единицам выбранного уровня при каждом значении параметра; n — рядом."
    n_cols = len(OVERVIEW_COLUMNS)
    heads = ("медиана", "q25", "q75", "n единиц, шт", "CR медиана", "CR q25", "CR q75")
    keys = ('$B$2&"|медиана"', '$B$2&"|q25"', '$B$2&"|q75"', '"n"', '"CR|медиана"', '"CR|q25"', '"CR|q75"')
    for codec, (c0, params) in over_layout(rows).items():
        cells["Обзор", TABLE_ROW - 2, c0] = CODEC_NAMES[codec]
        cells["Обзор", TABLE_ROW - 1, c0] = PARAM_NAMES[codec]
        for j, head in enumerate(heads, start=1):
            cells["Обзор", TABLE_ROW - 1, c0 + j] = head
        for i, p in enumerate(params):
            r = TABLE_ROW + i
            cells["Обзор", r, c0] = p
            for j, key in enumerate(keys, start=1):
                cells["Обзор", r, c0 + j] = overview_formula(codec, f"${_col(c0)}{r}", key, last_over, n_cols)
    return cells


def _scatter(title, x_title, y_title, log_x=False, log_y=False):
    from openpyxl.chart import ScatterChart

    chart = ScatterChart()
    chart.title, chart.style = title, 13
    chart.x_axis.title, chart.y_axis.title = x_title, y_title
    chart.x_axis.delete = chart.y_axis.delete = False
    if log_x:
        chart.x_axis.scaling.logBase = 10
    if log_y:
        chart.y_axis.scaling.logBase = 10
    chart.width, chart.height = 16, 8
    return chart


def _series(ws, x_col, y_col, first, last, title):
    from openpyxl.chart import Reference, Series
    from openpyxl.chart.marker import Marker

    s = Series(Reference(ws, min_col=y_col, min_row=first, max_row=last),
               Reference(ws, min_col=x_col, min_row=first, max_row=last), title=title)
    s.marker = Marker(symbol="circle", size=3)
    s.smooth = False
    return s


def add_view_charts(ws, rows):
    """13 charts of «Просмотр», anchored right of the tables in a grid of two columns."""
    layout = view_layout(rows)
    anchor_col = 1 + len(CODECS) * VIEW_BLOCK + 1
    charts = []
    for codec, (c0, params) in layout.items():  # metric = f(parameter): PSNR, PSNR-HVS-M left, HaarPSI right
        first, last = TABLE_ROW, TABLE_ROW + len(params) - 1
        chart = _scatter(f"{CODEC_NAMES[codec]}: метрики = f({PARAM_NAMES[codec]})", PARAM_NAMES[codec], "PSNR, PSNR-HVS-M, дБ")
        chart.series.append(_series(ws, c0, c0 + 1, first, last, "PSNR"))
        chart.series.append(_series(ws, c0, c0 + 2, first, last, "PSNR-HVS-M"))
        right = _scatter("", PARAM_NAMES[codec], "HaarPSI, безразм.")
        right.series.append(_series(ws, c0, c0 + 3, first, last, "HaarPSI"))
        right.y_axis.axId, right.y_axis.crosses = 200, "max"
        chart += right
        charts.append(chart)
    for m, (key, name, unit_title) in enumerate(METRICS):  # metric = f(CR), four codecs
        chart = _scatter(f"{name} = f(CR), четыре кодека", "CR, безразм. (лог. шкала)", unit_title, log_x=True)
        for codec, (c0, params) in layout.items():
            chart.series.append(_series(ws, c0 + 4, c0 + 1 + m, TABLE_ROW, TABLE_ROW + len(params) - 1, CODEC_NAMES[codec]))
        if key == "psnr_hvs_m_db":
            for k, t in enumerate(THRESHOLDS_DB):
                chart.series.append(_series(ws, 2, 3 + k, 4, 5, f"порог {t} дБ"))
        charts.append(chart)
    for codec, (c0, params) in layout.items():  # CR = f(parameter)
        chart = _scatter(f"{CODEC_NAMES[codec]}: CR = f({PARAM_NAMES[codec]})", PARAM_NAMES[codec],
                         "CR, безразм. (лог. шкала)", log_y=True)
        chart.series.append(_series(ws, c0, c0 + 4, TABLE_ROW, TABLE_ROW + len(params) - 1, "CR"))
        charts.append(chart)
    for x_m, y_m in ((0, 1), (1, 2)):  # metric against metric
        chart = _scatter(f"{METRICS[y_m][1]} против {METRICS[x_m][1]}", METRICS[x_m][2], METRICS[y_m][2])
        for codec, (c0, params) in layout.items():
            chart.series.append(_series(ws, c0 + 1 + x_m, c0 + 1 + y_m, TABLE_ROW, TABLE_ROW + len(params) - 1,
                                        CODEC_NAMES[codec]))
        charts.append(chart)
    for k, chart in enumerate(charts):
        ws.add_chart(chart, f"{_col(anchor_col + (k % 2) * 11)}{2 + (k // 2) * 18}")
    return charts


def add_overview_charts(ws, rows):
    """9 charts of «Обзор»."""
    layout = over_layout(rows)
    anchor_col = 1 + len(CODECS) * OVER_BLOCK + 1
    charts = []
    for codec, (c0, params) in layout.items():  # selected metric = f(parameter): median and quartiles
        first, last = TABLE_ROW, TABLE_ROW + len(params) - 1
        chart = _scatter(f"{CODEC_NAMES[codec]}: выбранная метрика = f({PARAM_NAMES[codec]}), медиана и квартили",
                         PARAM_NAMES[codec], "выбранная метрика (дБ или безразм.)")
        for j, title in ((1, "медиана"), (2, "q25"), (3, "q75")):
            chart.series.append(_series(ws, c0, c0 + j, first, last, title))
        charts.append(chart)
    chart = _scatter("Медиана метрики против медианы CR, четыре кодека", "медиана CR, безразм. (лог. шкала)",
                     "медиана выбранной метрики", log_x=True)
    for codec, (c0, params) in layout.items():
        chart.series.append(_series(ws, c0 + 5, c0 + 1, TABLE_ROW, TABLE_ROW + len(params) - 1, CODEC_NAMES[codec]))
    charts.append(chart)
    for codec, (c0, params) in layout.items():  # CR = f(parameter): median and quartiles
        first, last = TABLE_ROW, TABLE_ROW + len(params) - 1
        chart = _scatter(f"{CODEC_NAMES[codec]}: CR = f({PARAM_NAMES[codec]}), медиана и квартили", PARAM_NAMES[codec],
                         "CR, безразм. (лог. шкала)", log_y=True)
        for j, title in ((5, "медиана"), (6, "q25"), (7, "q75")):
            chart.series.append(_series(ws, c0, c0 + j, first, last, title))
        charts.append(chart)
    for k, chart in enumerate(charts):
        ws.add_chart(chart, f"{_col(anchor_col + (k % 2) * 11)}{2 + (k // 2) * 18}")
    return charts


SHEETS = ("Описание", "Сводка", "Просмотр", "Обзор", "Строки", "Обзор-данные", "Списки")


def build_book(run_path, book_path, describe, check):
    """Build the book from the run file and check it (`check` - a function (run_path, book_path))."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    run_bytes = Path(run_path).read_bytes()
    record = json.loads(run_bytes.decode("utf-8"))
    cells = book_cells(record, hashlib.sha256(run_bytes).hexdigest(), describe)
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    n_units = len(record["units"])
    for sheet, cell, rng in (("Просмотр", "B1", f"$A$2:$A${1 + n_units}"), ("Обзор", "B1", f"$B$2:$B${1 + len(LEVELS)}"),
                             ("Обзор", "B2", f"$C$2:$C${1 + len(METRICS)}")):
        dv = DataValidation(type="list", formula1=f"{_q('Списки')}!{rng}", allow_blank=False)
        wb[sheet].add_data_validation(dv)
        dv.add(cell)
        wb[sheet][cell].font = Font(bold=True)
    add_view_charts(wb["Просмотр"], record["rows"])
    add_overview_charts(wb["Обзор"], record["rows"])
    for name in ("Описание", "Сводка"):
        wb[name].column_dimensions["A"].width = 44
        wb[name].column_dimensions["B"].width = 130
        for row in wb[name].iter_rows(min_col=2, max_col=2):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for name in ("Строки", "Обзор-данные"):
        for c in wb[name][1]:
            c.font = Font(bold=True)
        wb[name].freeze_panes = "B3"
    wb.calculation.fullCalcOnLoad = True
    wb.save(book_path)
    check(run_path, book_path)


def check_book(run_path, book_path, describe, refusal):
    """Every cell (values and formula texts) equals what the run file gives; no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest(), describe)
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise refusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise refusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")
