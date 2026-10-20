"""V3-b and V4: do our JPEG, HEIF and AVIF on noisy Frisco and Diego reproduce the group's RS24 and AS25 numbers within STUDY.md section 8? Run: python studies/P001_vvc-dental/run_v3b_v4.py

Study P001, cycle 7: checks V3-b (JPEG, B06) and V4 (HEIF, B07; AVIF, B08) of STUDY.md section 8.
Frisco and Diego with AWGN (sigma^2 = 25, 100, 196; seeds 1..5) are coded at every quality of each
wrapper; the best point (OOP), PSNR_max and CR are compared with the assistant's reference file
(path: key v3b_v4_reference of lab_paths.json; its sha256 is pinned below). Writes
<cache_root>/cycles/7/v3b_v4.json (refuses by name if it exists) and the book v3b_v4.xlsx beside it;
`--rebuild-book` rebuilds the book from the run file.
"""

import argparse
import datetime
import hashlib
import json
import math
import platform
import statistics
import sys
import time
from pathlib import Path

import numpy as np

import imstudies
from imstudies.core.metrics import Refusal, psnr
from imstudies.lab.coders import avif, heif, isobmff, jpeg
from imstudies.lab.images import load_test_images
from imstudies.lab.runs import RunRefusal, git_state, json_number, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
COMMAND = "python studies/P001_vvc-dental/run_v3b_v4.py"
QUESTION = ("V3-b и V4: воспроизводят ли наши JPEG (B06), HEIF (B07) и AVIF (B08) на зашумлённых Frisco и Diego "
            "числа группы RS24 (табл. 1, 3) и AS25 (табл. 1) в допусках STUDY.md §8?")
CYCLE, RUN_FILE, BOOK_FILE = 7, "v3b_v4.json", "v3b_v4.xlsx"
TEST_IMAGES = ["frisco.bmp", "DIEGO.BMP"]
SIGMA2 = (25, 100, 196)
SEEDS = (1, 2, 3, 4, 5)
CODECS = {"jpeg": jpeg, "heif": heif, "avif": avif}
FULL_PLAN = {"sigma2": SIGMA2, "seeds": SEEDS,
             "qualities": {c: tuple(range(m.QUALITY_RANGE[0], m.QUALITY_RANGE[1] + 1)) for c, m in CODECS.items()}}
# sha256 of tmp/assistant_handover/tools/v3b_v4_reference.json in the boundary of cycle 7 (boundary_7_in.sha256)
REFERENCE_SHA256 = "88f7661b38309874accdf8bbf9d5f720596d341e9594c070076eadeb198a85e0"
AS25_QF = (9, 19, 29, 39, 49, 59, 69, 79, 89, 99)
STEEP_QF = (19, 29, 39)
TOL = {"v3b_qf": 3, "v3b_psnr_db": 0.3, "v3b_cr": 0.10, "heif_cr": 0.15, "heif_shift": 4, "heif_qf": 4, "avif_qf": 6}
DECLARED = {
    "data": ("Frisco и Diego из Images/BMP, 512 × 512 (imstudies.lab.images.load_test_images — чтение цикла 1; "
             "пиксели сверяются с файлом прогона цикла 1). Роль — тестовые изображения сверки; в выборку P001 не входят."),
    "noise": ("AWGN, σ² = 25, 100, 196 (σ = 5, 10, 14), зёрна 1…5: numpy.random.default_rng(seed).normal(0, σ, (H, W)); "
              "сумма с истинным изображением в float64, numpy.rint, отсечение 0…255, uint8; при данном зерне один "
              "генератор для обоих изображений и всех σ². ПРЕДПОЛОЖЕНИЕ о методике группы (в статьях не названа): "
              "иной способ сдвинет числа."),
    "codecs": ("B06 JPEG QF 1…100, B07 HEIF QF 0…100, B08 AVIF QF 0…100 (max_threads = 1), шаг 1; версии библиотек "
               "проверяются до прогона."),
    "metrics": ("PSNR_tc = B02(истинное, восстановленное), PSNR_nc = B02(зашумлённое, восстановленное), дБ; "
                "max_pixels = H·W изображения (ключ core_params.io.max_image_pixels не назначен; как в прогонах "
                "циклов 3 и 5)."),
    "oop": ("OOP кривой (изображение, σ², зерно, кодек) — QF с наибольшим PSNR_tc среди строк valid, при равенстве — "
            "меньший QF; PSNR_max — PSNR_tc в OOP; CR в OOP — CR по файлу в OOP. Сводка — медиана, наименьшее и "
            "наибольшее по пяти зёрнам."),
    "cr": ("CR по файлу = 8·W·H / (8·байт файла) = W·H / байт файла — критерий; CR по mdat = W·H / байт полезной "
           "нагрузки бокса mdat (без заголовка бокса) — диагностика вклада контейнера (HEIF, AVIF)."),
    "criterion_v3b": ("V3-b, JPEG, RS24 табл. 1 и 3 (Frisco σ² = 25, 100, 196; Diego σ² = 196), по медиане пяти зёрен: "
                      "|QF_OOP − эталон| ≤ 3; |PSNR_max − эталон| ≤ 0,3 дБ; |CR_OOP / эталон − 1| ≤ 10 %. Строка эталона "
                      "в допуске, если в допуске все три величины."),
    "criterion_v4_heif": ("V4 HEIF: AS25 табл. 1 — при QF 9 и 49…99 |медиана CR по файлу при этом QF / эталон − 1| ≤ 15 %; "
                          "при QF 19, 29, 39 (крутая зона) — сдвиг QF по медианной кривой, |сдвиг| ≤ 4; RS24 — "
                          "|медиана QF_OOP − эталон| ≤ 4."),
    "criterion_v4_avif": ("V4 AVIF, качественно: на Frisco при σ² = 25, 100, 196 медиана QF_OOP строго внутри 0…100 и "
                          "|медиана QF_OOP − эталон| ≤ 6; число зёрен с OOP внутри — рядом. У группы AVIF шёл через "
                          "libheif, у нас — через libavif (B08)."),
    "interpolation": ("Медианная кривая m(q) — медиана по зёрнам CR по файлу HEIF при каждом QF сетки прогона. Для точки "
                      "эталона (q_ref, CR_ref) пересечения — QF сетки с m(q) = CR_ref и интервалы между соседними QF "
                      "сетки, на которых m − CR_ref меняет знак; внутри интервала q* = q1 + (CR_ref − m(q1))·(q2 − q1) / "
                      "(m(q2) − m(q1)) — линейно по CR. Из нескольких пересечений — ближайшее к q_ref (при равенстве — "
                      "меньшее q*); число пересечений печатается. Нет пересечения — сдвиг не определён, строка вне "
                      "допуска. Сдвиг = q* − q_ref. То же правило на кривой каждого зерна даёт наименьший и наибольший "
                      "сдвиг."),
    "spread": ("Разброс по зёрнам — наибольшее минус наименьшее (для CR — делённое на медиану). Разброс больше "
               "полуширины допуска строки называется в строке сверки; допуск пересматривает владелец до следующего "
               "цикла, не задним числом."),
    "no_criterion": ("Без критерия печатаются: CR по mdat; CR по файлу в крутой зоне и сдвиг QF вне её; QF_OOP HEIF по "
                     "MSE из текста AS25; PSNR_max и CR в OOP HEIF и AVIF; Diego AVIF (в RS24 — локальный максимум)."),
    "refutation": ("Выход за допуск — V3-b или V4 не пройдена в названной строке эталона; разбор, не подгонка: допуск, "
                   "метод шума и правило OOP после чисел не меняются."),
    "reference": (f"Эталон — файл ассистента v3b_v4_reference.json (RS24, AS25, сняты по PDF); sha256 сверяется с "
                  f"объявленным {REFERENCE_SHA256} до чтения строк; состав строк проверяется."),
    "pages": "Страницы для глаз не применяются — численная сверка с опубликованными таблицами (как V1 и V3-a).",
}


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {}
    for key in ("images_root", "cache_root", "v3b_v4_reference"):
        if not raw.get(key):
            raise RunRefusal(f"lab_paths.json has no key '{key}'")
        paths[key] = (ROOT / raw[key]).resolve()
    if not paths["v3b_v4_reference"].is_file():
        raise RunRefusal(f"V3-b/V4 reference not found: {paths['v3b_v4_reference']}")
    return paths


def load_reference(path):
    """The reference file: sha256 equal to the declared one, then rows checked by check_reference."""
    found = sha256_file(path)
    if found != REFERENCE_SHA256:
        raise RunRefusal(f"{Path(path).name}: sha256 {found} differs from the declared {REFERENCE_SHA256} "
                         f"(boundary of cycle 7): the reference is not the assistant's file")
    reference = json.loads(Path(path).read_text(encoding="utf-8"))
    check_reference(reference)
    return reference


def _positive(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool) and math.isfinite(x) and x > 0


def check_reference(reference):
    """Every expected row once, with values of the declared kind; a bad row is refused by section and index."""
    expect = {
        "as25_table1_heif_cr": {(i, s, q) for i in ("frisco", "diego") for s in SIGMA2 for q in AS25_QF},
        "rs24_oop_by_psnr": {(i, s, c) for i in ("frisco",) for s in SIGMA2 for c in CODECS}
                            | {("diego", 196, c) for c in CODECS},
        "as25_heif_qf_oop_by_mse": {("frisco", s) for s in SIGMA2} | {("diego", 196)},
    }
    for section, wanted in expect.items():
        rows = reference.get(section)
        if not isinstance(rows, list):
            raise RunRefusal(f"reference has no list '{section}'")
        seen = set()
        for i, row in enumerate(rows):
            where = f"reference row {section}[{i}]"
            if section == "as25_table1_heif_cr":
                key, ok = (row.get("image"), row.get("sigma2"), row.get("qf")), _positive(row.get("cr"))
            elif section == "rs24_oop_by_psnr":
                key = (row.get("image"), row.get("sigma2"), row.get("codec"))
                low, high = CODECS[key[2]].QUALITY_RANGE if key[2] in CODECS else (None, None)
                qf = row.get("qf_oop")
                ok = (isinstance(qf, int) and not isinstance(qf, bool) and low is not None and low <= qf <= high
                      and _positive(row.get("psnr_max_db")) and _positive(row.get("cr")))
            else:
                key, qf = (row.get("image"), row.get("sigma2")), row.get("qf_oop")
                ok = isinstance(qf, int) and not isinstance(qf, bool) and 0 <= qf <= 100
            if key not in wanted:
                raise RunRefusal(f"{where}: unexpected or broken key {key}")
            if key in seen:
                raise RunRefusal(f"{where}: duplicate key {key}")
            if not ok:
                raise RunRefusal(f"{where}: value out of the declared kind: {row}")
            seen.add(key)
        if seen != wanted:
            raise RunRefusal(f"reference section {section} lacks rows {sorted(wanted - seen)[:5]}")


def check_against_cycle_1(images, cycle1_images):
    cycle1 = {i["code"]: i["pixels_sha256"] for i in cycle1_images if i["class"] == "test"}
    differ = [img["code"] for img in images if cycle1.get(img["code"]) != hashlib.sha256(img["array"].tobytes()).hexdigest()]
    if differ:
        raise RunRefusal(f"the test images read now differ from cycle 1 (pixels sha256): {differ}")


def noisy_image(true, sigma2, seed):
    """AWGN of variance sigma2 with seed: float64 sum, rint, clip 0..255, uint8 (declared before the numbers)."""
    noise = np.random.default_rng(seed).normal(0, math.sqrt(sigma2), true.shape)
    return np.clip(np.rint(true.astype(np.float64) + noise), 0, 255).astype(np.uint8)


def mdat_payload_bytes(file):
    sizes = [len(payload) for path, kind, payload, _ in isobmff.boxes(file) if path == "mdat"]
    return sum(sizes) if sizes else None


def oop(points):
    """(qf, psnr_tc) pairs -> the qf of the largest psnr_tc, the smaller qf on a tie; None if empty."""
    return min(points, key=lambda p: (-p[1], p[0]))[0] if points else None


def crossing(qs, values, target, q_ref):
    """QF where the piecewise-linear curve (qs ascending) equals `target`, nearest to q_ref: (q*, n crossings)."""
    found = set()
    for i, (q, v) in enumerate(zip(qs, values)):
        if v == target:
            found.add(float(q))
        if i + 1 < len(qs):
            a, b = v - target, values[i + 1] - target
            if a * b < 0:
                found.add(q + (target - v) * (qs[i + 1] - q) / (values[i + 1] - v))
    if not found:
        return None, 0
    return min(found, key=lambda x: (abs(x - q_ref), x)), len(found)


def _jsonable(obj):
    """Non-finite floats as the strings of json_number ('+inf' for PSNR at zero MSE, B02), recursively."""
    if isinstance(obj, float):
        return json_number(obj)
    if isinstance(obj, dict):
        return {k: _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    return obj


def _stats(values):
    values = [v for v in values if v is not None]
    if not values:
        return None
    return {"median": statistics.median(values), "min": min(values), "max": max(values), "n": len(values)}


def code_one(true, noisy, codec, q, h, w):
    row = {"codec": codec, "qf": q}
    try:
        coded = CODECS[codec].encode_decode(noisy, q, max_pixels=noisy.size)
        tc = psnr(true, coded.recon, max_pixels=true.size).value
        nc = psnr(noisy, coded.recon, max_pixels=true.size).value
    except Refusal as refusal:
        return row | {"status": "no_result", "reason": refusal.reason}
    except Exception as exc:  # noqa: BLE001 - status `error` by SPEC-RESULT
        return row | {"status": "error", "reason": "internal_error", "detail": str(exc)[:600]}
    mdat = mdat_payload_bytes(coded.file) if codec != "jpeg" else None
    return row | {"status": "valid", "reason": None, "file_bytes": len(coded.file), "mdat_payload_bytes": mdat,
                  "cr_file": h * w / len(coded.file), "cr_mdat": h * w / mdat if mdat else None,
                  "psnr_tc_db": tc, "psnr_nc_db": nc, "file_sha256": hashlib.sha256(coded.file).hexdigest()}


def curves_of(rows):
    """Per (image, sigma2, seed, codec): the OOP and the values there."""
    groups = {}
    for r in rows:
        groups.setdefault((r["image"], r["sigma2"], r["seed"], r["codec"]), []).append(r)
    out = []
    for (image, s2, seed, codec), rs in groups.items():
        valid = [r for r in rs if r["status"] == "valid"]
        q = oop([(r["qf"], r["psnr_tc_db"]) for r in valid])
        at = next((r for r in valid if r["qf"] == q), {})
        low, high = CODECS[codec].QUALITY_RANGE
        out.append({"image": image, "sigma2": s2, "seed": seed, "codec": codec, "n_valid": len(valid),
                    "n_not_valid": len(rs) - len(valid), "qf_oop": q, "psnr_max_db": at.get("psnr_tc_db"),
                    "psnr_nc_at_oop_db": at.get("psnr_nc_db"), "cr_file_oop": at.get("cr_file"),
                    "cr_mdat_oop": at.get("cr_mdat"), "oop_inside_range": None if q is None else low < q < high})
    return out


def median_curves(rows):
    groups = {}
    for r in rows:
        if r["status"] == "valid":
            groups.setdefault((r["image"], r["sigma2"], r["codec"], r["qf"]), []).append(r)
    return [{"image": i, "sigma2": s, "codec": c, "qf": q, "n_seeds": len(rs),
             "cr_file_median": statistics.median(r["cr_file"] for r in rs),
             "psnr_tc_median_db": statistics.median(r["psnr_tc_db"] for r in rs)}
            for (i, s, c, q), rs in sorted(groups.items())]


def summary_of(curves):
    out = []
    for key in sorted({(c["image"], c["sigma2"], c["codec"]) for c in curves}):
        cs = [c for c in curves if (c["image"], c["sigma2"], c["codec"]) == key]
        entry = {"image": key[0], "sigma2": key[1], "codec": key[2], "n_seeds": len(cs),
                 "n_oop_inside_range": sum(bool(c["oop_inside_range"]) for c in cs)}
        for name in ("qf_oop", "psnr_max_db", "cr_file_oop", "cr_mdat_oop"):
            st = _stats([c[name] for c in cs]) or {}
            entry |= {f"{name}_{k}": st.get(k) for k in ("median", "min", "max")}
        out.append(entry)
    return out


def _compare_row(check, source, image, s2, codec, quantity, reference, st, kind, tol, criterion, note=None, **extra):
    """One comparison: deviation of the median from the reference, tolerance, spread over seeds."""
    row = {"check": check if criterion else "—", "source": source, "image": image, "sigma2": s2, "codec": codec,
           "quantity": quantity, "qf_ref": extra.get("qf_ref"), "reference_cr": extra.get("reference_cr"),
           "reference": reference, "ours_median": None, "ours_min": None, "ours_max": None, "n_seeds": 0,
           "deviation": None, "deviation_kind": kind, "tolerance": tol, "in_tolerance": None, "spread": None,
           "spread_exceeds_tolerance": None, "n_crossings": extra.get("n_crossings"), "note": note}
    if st is None:
        row["note"] = "; ".join(x for x in (note, "нет данных в этом прогоне") if x)
        return row
    med = st["median"]
    row |= {"ours_median": med, "ours_min": st["min"], "ours_max": st["max"], "n_seeds": st["n"]}
    if med is None:  # the value is undefined; the caller names why in the note
        return row
    if kind == "relative":
        row["deviation"] = med / reference - 1
        row["spread"] = (st["max"] - st["min"]) / med
    elif kind == "inside_range":
        row["spread"] = st["max"] - st["min"]
    else:
        row["deviation"] = med - reference
        row["spread"] = None if st["min"] is None else st["max"] - st["min"]
    if criterion:
        if kind == "inside_range":
            low, high = CODECS[codec].QUALITY_RANGE
            row["in_tolerance"] = low < med < high
        else:
            row["in_tolerance"] = abs(row["deviation"]) <= tol
        if row["spread"] is not None and tol is not None and kind != "inside_range":
            row["spread_exceeds_tolerance"] = row["spread"] > tol
    return row


def compare(reference, rows, curves, medians):
    out = []
    by_curve = {}
    for c in curves:
        by_curve.setdefault((c["image"], c["sigma2"], c["codec"]), []).append(c)

    def seed_stats(image, s2, codec, name):
        return _stats([c[name] for c in by_curve.get((image, s2, codec), [])])

    table = {"frisco": "RS24 табл. 1", "diego": "RS24 табл. 3"}
    for ref in reference["rs24_oop_by_psnr"]:
        i, s2, codec = ref["image"], ref["sigma2"], ref["codec"]
        src, note = table[i], ref.get("note")
        if codec == "jpeg":
            crit = [("V3-b", "qf_oop", ref["qf_oop"], "difference", TOL["v3b_qf"], True),
                    ("V3-b", "psnr_max_db", ref["psnr_max_db"], "difference", TOL["v3b_psnr_db"], True),
                    ("V3-b", "cr_file_oop", ref["cr"], "relative", TOL["v3b_cr"], True)]
        elif codec == "heif":
            crit = [("V4 HEIF", "qf_oop", ref["qf_oop"], "difference", TOL["heif_qf"], True),
                    ("V4 HEIF", "psnr_max_db", ref["psnr_max_db"], "difference", None, False),
                    ("V4 HEIF", "cr_file_oop", ref["cr"], "relative", None, False),
                    ("V4 HEIF", "cr_mdat_oop", ref["cr"], "relative", None, False)]
        else:
            frisco = i == "frisco"
            crit = [("V4 AVIF", "oop_inside_range", None, "inside_range", None, frisco),
                    ("V4 AVIF", "qf_oop", ref["qf_oop"], "difference", TOL["avif_qf"], frisco),
                    ("V4 AVIF", "psnr_max_db", ref["psnr_max_db"], "difference", None, False),
                    ("V4 AVIF", "cr_file_oop", ref["cr"], "relative", None, False),
                    ("V4 AVIF", "cr_mdat_oop", ref["cr"], "relative", None, False)]
        for check, quantity, value, kind, tol, criterion in crit:
            name = "qf_oop" if quantity == "oop_inside_range" else quantity
            st = seed_stats(i, s2, codec, name)
            extra_note = note
            if quantity == "oop_inside_range":
                inside = sum(bool(c["oop_inside_range"]) for c in by_curve.get((i, s2, codec), []))
                extra_note = "; ".join(x for x in (note, f"зёрен с OOP внутри: {inside}") if x)
            out.append(_compare_row(check, src, i, s2, codec, quantity, value, st, kind, tol, criterion, extra_note))

    med = {}
    for m in medians:
        med.setdefault((m["image"], m["sigma2"], m["codec"]), []).append(m)
    per_seed = {}
    for r in rows:
        if r["codec"] == "heif" and r["status"] == "valid":
            per_seed.setdefault((r["image"], r["sigma2"], r["seed"]), []).append(r)
    for ref in reference["as25_table1_heif_cr"]:
        i, s2, q, cr = ref["image"], ref["sigma2"], ref["qf"], ref["cr"]
        steep = q in STEEP_QF
        at_q = [r for rs in (per_seed.get((i, s2, seed), []) for seed in SEEDS) for r in rs if r["qf"] == q]
        out.append(_compare_row("V4 HEIF", "AS25 табл. 1", i, s2, "heif", "cr_file_at_qf", cr,
                                _stats([r["cr_file"] for r in at_q]), "relative", TOL["heif_cr"], not steep,
                                qf_ref=q, reference_cr=cr))
        out.append(_compare_row("V4 HEIF", "AS25 табл. 1", i, s2, "heif", "cr_mdat_at_qf", cr,
                                _stats([r["cr_mdat"] for r in at_q]), "relative", None, False,
                                qf_ref=q, reference_cr=cr))
        curve = sorted(med.get((i, s2, "heif"), []), key=lambda m: m["qf"])
        q_star, n_cross = crossing([m["qf"] for m in curve], [m["cr_file_median"] for m in curve], cr, q)
        seeds_q = []
        for seed in SEEDS:
            rs = sorted(per_seed.get((i, s2, seed), []), key=lambda r: r["qf"])
            if rs:
                seeds_q.append(crossing([r["qf"] for r in rs], [r["cr_file"] for r in rs], cr, q)[0])
        st = None
        if curve:
            defined = [x for x in seeds_q if x is not None]
            st = {"median": q_star, "min": min(defined) if defined else None, "max": max(defined) if defined else None,
                  "n": len(seeds_q)}
        note = None
        if curve and q_star is None:
            note = "нет пересечения: сдвиг не определён"
        if seeds_q and any(x is None for x in seeds_q):
            note = "; ".join(x for x in (note, f"у {sum(x is None for x in seeds_q)} зёрен нет пересечения") if x)
        row = _compare_row("V4 HEIF", "AS25 табл. 1", i, s2, "heif", "qf_at_reference_cr", q, st, "difference",
                           TOL["heif_shift"], steep, note, qf_ref=q, reference_cr=cr, n_crossings=n_cross if curve else None)
        if steep and curve and q_star is None:
            row["in_tolerance"] = False
        out.append(row)
    for ref in reference["as25_heif_qf_oop_by_mse"]:
        out.append(_compare_row("V4 HEIF", "AS25, текст", ref["image"], ref["sigma2"], "heif", "qf_oop_by_mse",
                                ref["qf_oop"], seed_stats(ref["image"], ref["sigma2"], "heif", "qf_oop"), "difference",
                                None, False, ref.get("note")))
    return out


def verdicts(comparison):
    out = {}
    for check in ("V3-b", "V4 HEIF", "V4 AVIF"):
        rows = [c for c in comparison if c["check"] == check]
        groups = {}
        for c in rows:
            key = (c["source"], c["image"], c["sigma2"], c["codec"], c["qf_ref"])
            groups.setdefault(key, []).append(c)
        failed = [f"{k[1]} σ²={k[2]} {k[3]}" + (f" QF {k[4]}" if k[4] is not None else "") + ": "
                  + ", ".join(c["quantity"] for c in g if c["in_tolerance"] is False)
                  for k, g in groups.items() if any(c["in_tolerance"] is False for c in g)]
        no_data = sum(any(c["in_tolerance"] is None for c in g) for g in groups.values())
        n_in = sum(all(c["in_tolerance"] is True for c in g) for g in groups.values())
        out[check] = {"reference_rows": len(groups), "rows_in_tolerance": n_in, "rows_out": len(failed),
                      "rows_without_data": no_data, "checks": len(rows),
                      "checks_in_tolerance": sum(c["in_tolerance"] is True for c in rows),
                      "spread_exceeds_tolerance": sum(bool(c["spread_exceeds_tolerance"]) for c in rows),
                      "failed": failed,
                      "verdict": "pass" if groups and not failed and not no_data else ("fail" if failed else "incomplete")}
    return out


def run(paths, reference, images, plan, out_dir, state):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path = out_dir / RUN_FILE, out_dir / BOOK_FILE
    if run_path.exists():
        raise RunRefusal(f"{run_path} already exists: a run file of cycle {CYCLE} is never overwritten")
    for name, module in CODECS.items():
        try:
            module.check_versions(name.upper(), module.EXPECTED_VERSIONS, module.versions())
        except RuntimeError as exc:
            raise RunRefusal(str(exc)) from exc
    started, t0 = datetime.datetime.now(datetime.timezone.utc), time.perf_counter()
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")
    rows, noise = [], []
    for img in images:
        true, (h, w) = img["array"], img["array"].shape
        for s2 in plan["sigma2"]:
            for seed in plan["seeds"]:
                noisy = noisy_image(true, s2, seed)
                noise.append({"image": img["code"], "sigma2": s2, "seed": seed,
                              "noisy_sha256": hashlib.sha256(noisy.tobytes()).hexdigest(),
                              "psnr_tn_db": psnr(true, noisy, max_pixels=true.size).value})
                for codec, qs in plan["qualities"].items():
                    t, start = time.perf_counter(), len(rows)
                    for q in qs:
                        rows.append({"row_index": len(rows) + 1, "image": img["code"], "sigma2": s2, "seed": seed}
                                    | code_one(true, noisy, codec, q, h, w))
                    group = rows[start:]
                    valid = [r for r in group if r["status"] == "valid"]
                    best = oop([(r["qf"], r["psnr_tc_db"]) for r in valid])
                    at = next((r for r in valid if r["qf"] == best), None)
                    print(f"  {img['code']:<6} s2 {s2:>3} seed {seed} {codec:<4}: {len(valid)}/{len(group)} valid, "
                          f"OOP QF {best}, PSNR_max {at['psnr_tc_db']:.3f} dB, CR {at['cr_file']:.3f} "
                          f"({time.perf_counter() - t:.1f} s)" if at else
                          f"  {img['code']:<6} s2 {s2:>3} seed {seed} {codec:<4}: no valid rows")
    curves = curves_of(rows)
    medians = median_curves(rows)
    comparison = compare(reference, rows, curves, medians)
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "wall_seconds": round(time.perf_counter() - t0, 1),
        **state, "declared_before_numbers": DECLARED, "tolerances": TOL,
        "plan": {"sigma2": list(plan["sigma2"]), "seeds": list(plan["seeds"]),
                 "qualities": {c: [min(qs), max(qs), len(qs)] for c, qs in plan["qualities"].items()}},
        "reference": {"file": paths["v3b_v4_reference"].name, "sha256": sha256_file(paths["v3b_v4_reference"]),
                      "declared_sha256": REFERENCE_SHA256, "what": reference.get("_what"),
                      "sources": reference.get("sources")},
        "wrappers": {"jpeg": {"versions": jpeg.versions(), "encode_args": jpeg.ENCODE_ARGS},
                     "heif": {"versions": heif.versions(), "enc_params": heif.ENC_PARAMS, "options": heif.OPTIONS},
                     "avif": {"versions": avif.versions(), "encode_args": avif.ENCODE_ARGS}},
        "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__,
                        "numpy": np.__version__, "platform": platform.platform()},
        "images": [{k: v for k, v in img.items() if k != "array"}
                   | {"pixels_sha256": hashlib.sha256(img["array"].tobytes()).hexdigest()} for img in images],
        "unit_of_analysis": "row = image x sigma2 x seed x codec x QF; summary per (image, sigma2, codec) over seeds",
        "noise": noise, "rows": rows, "curves": curves, "median_curves": medians, "summary": summary_of(curves),
        "as25_reference": reference["as25_table1_heif_cr"], "comparison": comparison,
        "verdicts": verdicts(comparison),
    }
    text = json.dumps(_jsonable(record), ensure_ascii=False, indent=1, allow_nan=False)  # before the file is opened
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        f.write(text)
    build_book(run_path, book_path)
    print_verdicts(record, run_path, book_path)
    return run_path, book_path


def print_verdicts(record, run_path, book_path):
    print(f"\nRows: {len(record['rows'])} ({sum(r['status'] == 'valid' for r in record['rows'])} valid); "
          f"wall time {record['wall_seconds']} s")
    for check, v in record["verdicts"].items():
        print(f"{check}: {v['rows_in_tolerance']} of {v['reference_rows']} reference rows in tolerance "
              f"({v['checks_in_tolerance']} of {v['checks']} checks; without data {v['rows_without_data']}; "
              f"seed spread over tolerance {v['spread_exceeds_tolerance']}) -> {v['verdict']}")
        for f in v["failed"]:
            print(f"    out of tolerance: {f}")
    print(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}")


QUANTITY = {"qf_oop": "QF_OOP, QF", "psnr_max_db": "PSNR_max, дБ", "cr_file_oop": "CR по файлу в OOP, безразм.",
            "cr_mdat_oop": "CR по mdat в OOP, безразм.", "oop_inside_range": "OOP внутри диапазона QF",
            "cr_file_at_qf": "CR по файлу при QF эталона, безразм.", "cr_mdat_at_qf": "CR по mdat при QF эталона, безразм.",
            "qf_at_reference_cr": "QF, где медианная кривая даёт CR эталона, QF",
            "qf_oop_by_mse": "QF_OOP по MSE (AS25), QF"}
KIND = {"difference": "разность медианы и эталона", "relative": "медиана / эталон − 1, %",
        "inside_range": "внутри 0…100"}
IN_TOL = {True: "да", False: "нет", None: "без критерия"}
COMPARE = [("Проверка", "check"), ("Источник", "source"), ("Изображение", "image"), ("σ²", "sigma2"),
           ("Кодек", "codec"), ("Величина", "quantity"), ("QF эталона", "qf_ref"), ("CR эталона, безразм.", "reference_cr"),
           ("Эталон", "reference"), ("Наше: медиана по зёрнам", "ours_median"), ("Наше: наименьшее", "ours_min"),
           ("Наше: наибольшее", "ours_max"), ("Зёрен, шт", "n_seeds"), ("Отклонение", "deviation"),
           ("Вид отклонения", "deviation_kind"), ("Допуск (полуширина)", "tolerance"), ("В допуске", "in_tolerance"),
           ("Разброс по зёрнам", "spread"), ("Разброс > допуска", "spread_exceeds_tolerance"),
           ("Пересечений, шт", "n_crossings"), ("Примечание", "note")]
SUMMARY = [("Изображение", "image"), ("σ²", "sigma2"), ("Кодек", "codec"), ("Зёрен, шт", "n_seeds"),
           ("OOP внутри диапазона, зёрен", "n_oop_inside_range")] + [
    (f"{label}: {stat}", f"{key}_{k}") for key, label in (("qf_oop", "QF_OOP, QF"), ("psnr_max_db", "PSNR_max, дБ"),
                                                          ("cr_file_oop", "CR по файлу в OOP"),
                                                          ("cr_mdat_oop", "CR по mdat в OOP"))
    for k, stat in (("median", "медиана"), ("min", "наименьшее"), ("max", "наибольшее"))]
CURVES = [("Изображение", "image"), ("σ²", "sigma2"), ("Зерно", "seed"), ("Кодек", "codec"), ("Valid, шт", "n_valid"),
          ("Не valid, шт", "n_not_valid"), ("QF_OOP, QF", "qf_oop"), ("PSNR_max (PSNR_tc в OOP), дБ", "psnr_max_db"),
          ("PSNR_nc в OOP, дБ", "psnr_nc_at_oop_db"), ("CR по файлу в OOP, безразм.", "cr_file_oop"),
          ("CR по mdat в OOP, безразм.", "cr_mdat_oop"), ("OOP внутри диапазона", "oop_inside_range")]
ROWS = [("№", "row_index"), ("Изображение", "image"), ("σ²", "sigma2"), ("Зерно", "seed"), ("Кодек", "codec"),
        ("QF", "qf"), ("Статус", "status"), ("Причина", "reason"), ("Файл, байт", "file_bytes"),
        ("mdat (нагрузка), байт", "mdat_payload_bytes"), ("CR по файлу, безразм.", "cr_file"),
        ("CR по mdat, безразм.", "cr_mdat"), ("PSNR_tc, дБ", "psnr_tc_db"), ("PSNR_nc, дБ", "psnr_nc_db"),
        ("sha256 файла", "file_sha256")]


def _shown(key, value):
    if key in ("in_tolerance",):
        return IN_TOL[value]
    if isinstance(value, bool):
        return "да" if value else "нет"
    if key == "quantity":
        return QUANTITY.get(value, value)
    if key == "deviation_kind":
        return KIND.get(value, value)
    return value


def _table(cells, formats, sheet, columns, records, percent=lambda rec, key: False):
    for j, (head, key) in enumerate(columns, start=1):
        cells[sheet, 1, j], cells[sheet, 2, j] = head, f"{key}"
    for i, rec in enumerate(records):
        for j, (_, key) in enumerate(columns, start=1):
            cells[sheet, 3 + i, j] = _shown(key, rec.get(key))
            if percent(rec, key):
                formats[sheet, 3 + i, j] = "0.0%"


def book_cells(record, run_sha256):
    cells, formats = {}, {}
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
             ("Закончен, UTC", record["finished_utc"]), ("Время прогона, с", record["wall_seconds"]),
             ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
             *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
             ("Эталон", f"{record['reference']['file']}, sha256 {record['reference']['sha256']}"),
             ("Версии JPEG", json.dumps(record["wrappers"]["jpeg"]["versions"], ensure_ascii=False)),
             ("Версии HEIF", json.dumps(record["wrappers"]["heif"]["versions"], ensure_ascii=False)),
             ("Версии AVIF", json.dumps(record["wrappers"]["avif"]["versions"], ensure_ascii=False)),
             ("Строк, шт", len(record["rows"])),
             ("Строк valid, шт", sum(r["status"] == "valid" for r in record["rows"]))]
    for check, v in record["verdicts"].items():
        lines.append((f"Вердикт {check}", f"{v['verdict']}: в допуске {v['rows_in_tolerance']} из {v['reference_rows']} "
                                           f"строк эталона ({v['checks_in_tolerance']} из {v['checks']} проверок); без "
                                           f"данных {v['rows_without_data']}; разброс по зёрнам больше допуска — "
                                           f"{v['spread_exceeds_tolerance']}; вне допуска: "
                                           f"{'; '.join(v['failed']) or 'нет'}"))
    lines += [("Ключи листов", "под шапкой каждого листа — имя поля строки файла прогона; «Сверка» — comparison[], "
                               "«Сводка» — summary[], «Зёрна» — curves[], «Строки» — rows[], «График» — median_curves[] "
                               "и as25_table1_heif_cr эталона"),
              ("Точность чисел книги", "как в файле прогона; каждая клетка сверена с файлом прогона"),
              ("Листы «Рисунки» и «Страницы»", "не строятся: страницы не применяются (объявлено до чисел)"),
              ("На пальцах", "Два тестовых аэроснимка зашумлены пятью разными случайными шумами трёх уровней и сжаты "
                             "тремя кодеками при каждом значении качества; для каждого случая найдена степень сжатия, "
                             "при которой сжатый снимок ближе всего к чистому, и это сравнивается с таблицами прежних "
                             "статей группы.")]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v
    _table(cells, formats, "Сверка", COMPARE, record["comparison"],
           lambda rec, key: rec["deviation_kind"] == "relative" and key in ("deviation", "tolerance", "spread"))
    _table(cells, formats, "Сводка", SUMMARY, record["summary"])
    _table(cells, formats, "Зёрна", CURVES, record["curves"])
    _table(cells, formats, "Строки", ROWS, record["rows"])
    conds = sorted({(m["image"], m["sigma2"]) for m in record["median_curves"]})
    combos = sorted({(m["image"], m["sigma2"], m["codec"]) for m in record["median_curves"]})
    cols = ([("heif_cr", i, s, "heif") for i, s in conds] + [("as25", i, s, "heif") for i, s in conds]
            + [("psnr", i, s, c) for i, s, c in combos])
    qs = sorted({m["qf"] for m in record["median_curves"]})
    med = {(m["image"], m["sigma2"], m["codec"], m["qf"]): m for m in record["median_curves"]}
    as25 = {(r["image"], r["sigma2"], r["qf"]): r["cr"] for r in record.get("as25_reference", [])}
    cells["График", 1, 1] = "median_curves[].qf"
    cells["График", 2, 1] = "QF"
    for j, (kind, i, s, c) in enumerate(cols, start=2):
        cells["График", 2, j] = {"heif_cr": f"HEIF CR медиана {i} σ²={s}", "as25": f"AS25 CR {i} σ²={s}",
                                 "psnr": f"PSNR_tc медиана, дБ {c} {i} σ²={s}"}[kind]
        cells["График", 1, j] = {"heif_cr": "median_curves[].cr_file_median", "as25": "as25_reference[].cr",
                                 "psnr": "median_curves[].psnr_tc_median_db"}[kind]
    for r, q in enumerate(qs, start=3):
        cells["График", r, 1] = q
        for j, (kind, i, s, c) in enumerate(cols, start=2):
            m = med.get((i, s, c, q))
            v = {"heif_cr": m and m["cr_file_median"], "as25": as25.get((i, s, q)),
                 "psnr": m and m["psnr_tc_median_db"]}[kind]
            if v is not None:
                cells["График", r, j] = v
    return cells, formats, len(conds), len(cols)


def build_book(run_path, book_path):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.marker import Marker
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    cells, formats, n_cond, n_cols = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Сверка", "Сводка", "Зёрна", "Строки", "График"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    for (sheet, r, c), fmt in formats.items():
        wb[sheet].cell(row=r, column=c).number_format = fmt
    wb["Описание"].column_dimensions["A"].width = 34
    wb["Описание"].column_dimensions["B"].width = 140
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for name in ("Сверка", "Сводка", "Зёрна", "Строки", "График"):
        for c in wb[name][2 if name == "График" else 1]:
            c.font = Font(bold=True)
        wb[name].freeze_panes = "A3"
    g = wb["График"]
    n_rows = max(r for s, r, _ in cells if s == "График")
    if n_rows >= 3 and n_cond:
        chart = LineChart()
        chart.title = "HEIF: CR по файлу (медиана по зёрнам) и AS25 по QF"
        chart.y_axis.title, chart.x_axis.title = "CR, безразм. (лог. шкала)", "QF"
        chart.y_axis.scaling.logBase = 10
        chart.x_axis.delete = chart.y_axis.delete = False
        chart.display_blanks = "gap"
        chart.add_data(Reference(g, min_col=2, max_col=1 + 2 * n_cond, min_row=2, max_row=n_rows), titles_from_data=True)
        chart.set_categories(Reference(g, min_col=1, min_row=3, max_row=n_rows))
        for series in chart.series[n_cond:]:
            series.marker = Marker(symbol="circle", size=7)
            series.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        chart.width, chart.height = 28, 14
        g.add_chart(chart, "B" + str(n_rows + 3))
        chart2 = LineChart()
        chart2.title = "PSNR_tc (медиана по зёрнам) по QF"
        chart2.y_axis.title, chart2.x_axis.title = "PSNR_tc, дБ", "QF"
        chart2.x_axis.delete = chart2.y_axis.delete = False
        chart2.display_blanks = "gap"
        chart2.add_data(Reference(g, min_col=2 + 2 * n_cond, max_col=n_cols + 1, min_row=2, max_row=n_rows),
                        titles_from_data=True)
        chart2.set_categories(Reference(g, min_col=1, min_row=3, max_row=n_rows))
        chart2.width, chart2.height = 28, 14
        g.add_chart(chart2, "B" + str(n_rows + 33))
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g, as openpyxl writes them); no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())[0]
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        reference = load_reference(paths["v3b_v4_reference"])
        images = load_test_images(paths["images_root"], TEST_IMAGES, (512, 512))
        cycle1 = json.loads((paths["cache_root"] / "cycles" / "1" / "v1_metrics.json").read_text(encoding="utf-8"))
        check_against_cycle_1(images, cycle1["images"])
        run(paths, reference, images, FULL_PLAN, out_dir, state)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
