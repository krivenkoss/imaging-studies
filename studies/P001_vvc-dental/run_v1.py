"""V1: do our PSNR, PSNR-HVS-M and HaarPSI agree with the authors' code on 40 pairs? Run: python studies/P001_vvc-dental/run_v1.py

Study P001, cycle 1: checks V1-a..V1-c of STUDY.md section 8 for blocks B02-B04 of
../REFERENCE.md. Paths come from lab_paths.json (images_root, cache_root, metric_sources,
octave_cli). Writes <cache_root>/cycles/1/v1_metrics.json (the run file; refuses by name
if it exists), v1_dental_codes.json (dental code -> source file; outside the repository)
and the book v1_metrics.xlsx beside them. `--rebuild-book` rebuilds the book from the
existing run file.
"""

import argparse
import datetime
import hashlib
import importlib.util
import io
import json
import math
import platform
import statistics
import subprocess
import sys
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path

import numpy as np
import scipy
from PIL import Image, features

import imstudies
from imstudies.core.metrics import Refusal, haarpsi, psnr, psnr_hvs_m
from imstudies.oracle import psnr as psnr_oracle

COMMAND = "python studies/P001_vvc-dental/run_v1.py"
QUESTION = ("V1: совпадают ли наши PSNR, PSNR-HVS-M и HaarPSI с авторскими реализациями "
            "на 40 парах (10 изображений x 4 искажения)?")
STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
CYCLE = 1
RUN_FILE = "v1_metrics.json"
CODES_FILE = "v1_dental_codes.json"
BOOK_FILE = "v1_metrics.xlsx"

# The authors' files as named in STUDY.md section 6 (sha256).
SOURCES_SHA256 = {
    "psnrhvsm/psnrhvsm.m": "4a373c54603f98adf1a4340b0bcdee8b0b8e1c7d4a2dce5e5699fb08147841b7",
    "haarpsi/HaarPSI.m": "ef1bed2872fd1f74842afdccfd1d6b0e3a93bc5479b14bf845460f08dec8ee4a",
    "haarpsi/haarPsi.py": "9b48464ad14126492ecc82725d49297718b0a6566b0e415dfbf85e94e6703d9e",
    "haarpsi/LICENSE": "113f167f37d16e380789e127632a3ca8d7e0dcdafcdef20fa282229cab87f1b7",
}
OCTAVE_VERSION, SIGNAL_VERSION = "11.3.0", "1.4.8"

TEST_IMAGES = ["BABOON.BMP", "BARBARA.BMP", "goldhill.bmp", "lenna.bmp", "frisco.bmp", "DIEGO.BMP"]
DENTAL_DIR, N_DENTAL, CROP = "Dental_panorama_original", 4, 512
DISTORTIONS = [
    {"id": "noise_s3_seed1", "kind": "gaussian_noise", "sigma": 3.0, "seed": 1},
    {"id": "noise_s8_seed2", "kind": "gaussian_noise", "sigma": 8.0, "seed": 2},
    {"id": "jpeg_q30", "kind": "jpeg_pillow", "quality": 30},
    {"id": "jpeg_q70", "kind": "jpeg_pillow", "quality": 70},
]
NOISE_RULE = ("numpy.random.default_rng(seed).normal(0, sigma, (H, W)) added to the image in float64, "
              "numpy.rint (half to even), clip to 0..255, uint8; a fresh generator per image")
JPEG_RULE = ("PIL.Image.fromarray(uint8 H x W) (mode L).save(BytesIO, format='JPEG', quality=q), other options "
             "Pillow defaults; decoded by PIL.Image.open to mode L uint8")

# (metric, reference, tolerance, tolerance unit, unit of the value); declared before the numbers.
CHECKS = [
    ("psnr", "oracle_decimal", 2.0, "ulp", "dB"),
    ("psnr_hvs_m", "psnrhvsm.m", 0.01, "dB", "dB"),
    ("haarpsi", "haarPsi.py", 1e-4, "dimensionless", "dimensionless"),
    ("haarpsi", "HaarPSI.m", 1e-4, "dimensionless", "dimensionless"),
]
DECLARED = {
    "roles": ("Все изображения цикла — отладочные для сверки реализаций; в роли данных P001 (tuning, check) "
              "не входят и в приёмке исследования не участвуют. Тестовые: BABOON, BARBARA, goldhill, lenna, "
              "frisco, DIEGO из Images/BMP (8 бит, 512 x 512). Стоматологические: из "
              "Images/Dental_panorama_original — первые четыре различных по SHA-256 файла BMP/PNG в порядке "
              "сортировки имён; яркость через палитру; центральная вырезка 512 x 512; коды v1d01…v1d04; "
              "таблица «код → файл» — только в cycles/1/ под cache_root. Итого 10 изображений."),
    "distortions": ("40 пар: на каждое изображение гауссов шум σ = 3 (seed 1) и σ = 8 (seed 2) с округлением "
                    "и отсечением к 0…255; JPEG средствами Pillow с качеством 30 и 70, декодирование в uint8."),
    "criterion": ("V1-a — SSE совпадает с аналитическим целым точно; PSNR побитово, где 65025·N/SSE = 10^k, "
                  "иначе ≤ 2 ulp от оракула повышенной точности. V1-b — |Δ PSNR-HVS-M| ≤ 0.01 дБ против "
                  "psnrhvsm.m под Octave на всех 40 парах. V1-c — |Δ HaarPSI| ≤ 1e-4 против haarPsi.py и против "
                  "HaarPSI.m под Octave на всех 40 парах. Сравниваемые поля: значение; +inf ↔ 100000 эталона "
                  "PSNR-HVS-M."),
    "refutation": ("Хотя бы одна пара сверх допуска V1-b или V1-c — вердикт «не пройдено», первая строка "
                   "отчёта, разбор без подгонки."),
    "expected": ("Расхождение PSNR-HVS-M порядка 1e-12 дБ, HaarPSI против haarPsi.py порядка 1e-14. "
                 "Порядок больше — дефект, не округление."),
    "path_control": ("До длинной части: синтетический кадр E0 и его копия с шумом σ = 3 (seed 1) — наше "
                     "значение и Octave для PSNR-HVS-M и HaarPSI в пределах допусков; расхождение — стоп."),
    "pages": "Правило выбора страниц для глаз не применяется — численная сверка; книга содержит все строки.",
    "run_part_of_v1a": ("В прогоне V1-a проверяется на тех же 40 парах против оракула decimal (≤ 2 ulp); "
                        "аналитические случаи V1-a — в тестах ядра (tests/test_metrics.py)."),
}
MAX_PIXELS_RULE = ("max_pixels = H*W of the pair itself: the limit is not under test in V1, and "
                   "config.json core_params.io.max_image_pixels is null (NOT ASSIGNED)")


class RunRefusal(Exception):
    """The run cannot go on; the message says why in words."""


# --- paths, sources, environment ---------------------------------------------------------------

def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {}
    for key in ("images_root", "cache_root", "metric_sources", "octave_cli"):
        if not raw.get(key):
            raise RunRefusal(f"lab_paths.json has no key '{key}'")
        paths[key] = (ROOT / raw[key]).resolve()
    for key in ("images_root", "metric_sources"):
        if not paths[key].is_dir():
            raise RunRefusal(f"lab_paths.json '{key}' is not a folder: {paths[key]}")
    if not paths["octave_cli"].is_file():
        raise RunRefusal(f"lab_paths.json 'octave_cli' is not a file: {paths['octave_cli']}")
    return paths


def sha256_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def check_sources(src):
    found = {}
    for rel, expected in SOURCES_SHA256.items():
        path = src / rel
        if not path.is_file():
            raise RunRefusal(f"reference file missing: {path}")
        found[rel] = sha256_file(path)
        if found[rel] != expected:
            raise RunRefusal(f"reference file {rel} has sha256 {found[rel]}, STUDY.md section 6 names {expected}")
    return found


def octave_versions(octave):
    code = "pkg load signal; p = pkg('list', 'signal'); printf('%s %s', version(), p{1}.version);"
    out = subprocess.run([str(octave), "--no-gui", "--quiet", "--norc", "--eval", code],
                         capture_output=True, text=True, timeout=120)
    parts = out.stdout.split()
    if out.returncode != 0 or len(parts) != 2:
        raise RunRefusal(f"Octave or its package 'signal' did not answer: {out.stderr.strip()[-300:]}")
    if parts != [OCTAVE_VERSION, SIGNAL_VERSION]:
        raise RunRefusal(f"Octave {parts[0]} with signal {parts[1]}; this run is pinned to "
                         f"Octave {OCTAVE_VERSION} with signal {SIGNAL_VERSION}")
    return {"octave": parts[0], "octave_signal": parts[1]}


def git_state():
    def git(*args):
        out = subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True)
        if out.returncode != 0:
            raise RunRefusal(f"git {' '.join(args)} failed: {out.stderr.strip()}")
        return out.stdout.strip()
    if git("status", "--porcelain", "--untracked-files=no"):
        raise RunRefusal("the working tree has uncommitted changes: the run goes from a committed tree")
    return {"core_commit": git("rev-parse", "HEAD"), "core_tree": git("rev-parse", "HEAD^{tree}")}


def environment():
    return {
        "python": platform.python_version(), "imstudies": imstudies.__version__, "numpy": np.__version__,
        "scipy": scipy.__version__, "pillow": Image.__version__,
        "libjpeg_turbo": features.version_feature("libjpeg_turbo"), "pillow_jpeg_api": features.version("jpg"),
        "platform": platform.platform(),
    }


def load_haarpsi_py(src):
    spec = importlib.util.spec_from_file_location("haarpsi_reference", src / "haarpsi" / "haarPsi.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# --- images and distortions -----------------------------------------------------------------------

def read_gray(path):
    """8-bit brightness of a palette or grey image: (uint8 array, how it was read)."""
    with Image.open(path) as im:
        mode = im.mode
        idx = np.asarray(im)
        palette = im.getpalette() if mode == "P" else None
    if mode == "L":
        if path.suffix.lower() == ".bmp":
            head = path.read_bytes()[:30]
            if head[:2] != b"BM" or int.from_bytes(head[28:30], "little") != 8:
                raise RunRefusal(f"{path.name}: not an 8-bit BMP")
        return idx.astype(np.uint8), "Pillow mode L (grey ramp palette or grey file): value = brightness"
    if mode == "P":
        pal = np.array(palette, dtype=np.int64).reshape(-1, 3)
        used = np.unique(idx)
        if used.max() >= len(pal) or not (pal[used, 0] == pal[used, 1]).all() or not (pal[used, 1] == pal[used, 2]).all():
            raise RunRefusal("a palette image is not grey on the indices it uses")
        return pal[idx, 0].astype(np.uint8), "Pillow mode P: brightness = grey palette entry of the index"
    raise RunRefusal(f"unsupported image mode {mode}")


def test_set(images_root):
    images = []
    for name in TEST_IMAGES:
        path = images_root / "BMP" / name
        if not path.is_file():
            raise RunRefusal(f"test image missing: {path}")
        arr, how = read_gray(path)
        if arr.shape != (CROP, CROP):
            raise RunRefusal(f"{name}: expected {CROP}x{CROP}, got {arr.shape}")
        images.append({"code": path.stem.lower(), "class": "test", "source": f"BMP/{name}",
                       "source_sha256": sha256_file(path), "full_shape": list(arr.shape), "crop_top_left": [0, 0],
                       "read": how, "array": arr})
    return images


def dental_set(images_root):
    """First N_DENTAL files distinct by file sha256, in Python sort order of names; centre crop."""
    folder = images_root / DENTAL_DIR
    names = sorted(p.name for p in folder.iterdir() if p.suffix.lower() in (".bmp", ".png"))
    images, codes, seen = [], [], set()
    for name in names:
        digest = sha256_file(folder / name)
        if digest in seen:
            continue
        seen.add(digest)
        full, how = read_gray(folder / name)
        h, w = full.shape
        if h < CROP or w < CROP:
            raise RunRefusal(f"dental image {len(images) + 1} is smaller than {CROP}x{CROP}")
        top, left = (h - CROP) // 2, (w - CROP) // 2
        code = f"v1d{len(images) + 1:02d}"
        images.append({"code": code, "class": "dental", "source": f"{DENTAL_DIR} (file name in {CODES_FILE})",
                       "source_sha256": None, "full_shape": [h, w], "crop_top_left": [top, left], "read": how,
                       "array": np.ascontiguousarray(full[top:top + CROP, left:left + CROP])})
        codes.append({"code": code, "file_name": name, "file_sha256": digest, "full_shape": [h, w],
                      "crop_top_left": [top, left]})
        if len(images) == N_DENTAL:
            return images, codes
    raise RunRefusal(f"only {len(images)} distinct dental files found, {N_DENTAL} needed")


def distort(img, d):
    if d["kind"] == "gaussian_noise":
        noise = np.random.default_rng(d["seed"]).normal(0, d["sigma"], img.shape)
        return np.rint(img + noise).clip(0, 255).astype(np.uint8)
    buf = io.BytesIO()
    source = Image.fromarray(img)
    if source.mode != "L":
        raise RunRefusal(f"uint8 array became Pillow mode {source.mode}, expected L")
    source.save(buf, format="JPEG", quality=d["quality"])
    with Image.open(io.BytesIO(buf.getvalue())) as im:
        if im.mode != "L":
            raise RunRefusal(f"JPEG decoded to mode {im.mode}, expected L")
        return np.asarray(im).astype(np.uint8)


def control_frame():
    """Synthetic frame E0 (PROMPT, cycle 1) and its copy with noise sigma 3, seed 1."""
    y, x = np.mgrid[0:512, 0:512]
    rng = np.random.default_rng(1)
    frame = (x / 2 + y / 4 + rng.normal(0, 8, (512, 512))).clip(0, 255).astype(np.uint8)
    return frame, distort(frame, DISTORTIONS[0])


# --- measurement --------------------------------------------------------------------------------------

def octave_metrics(octave, src, pairs, workdir):
    """psnrhvsm.m and HaarPSI.m under Octave for each (ref, dist): list of (psnrhvsm, haarpsi)."""
    with tempfile.TemporaryDirectory(dir=workdir, prefix="octave_io_") as tmp:
        tmp = Path(tmp)
        with open(tmp / "list.txt", "w", encoding="ascii") as f:
            for k, (ref, dist) in enumerate(pairs):
                np.ascontiguousarray(ref).tofile(tmp / f"r{k}.bin")
                np.ascontiguousarray(dist).tofile(tmp / f"d{k}.bin")
                f.write(f"{k} {ref.shape[0]} {ref.shape[1]}\n")
        out = subprocess.run([str(octave), "--no-gui", "--quiet", "--norc", str(STUDY / "octave_v1.m"),
                              src.as_posix(), tmp.as_posix()], capture_output=True, text=True, timeout=3600)
        if out.returncode != 0 or not (tmp / "out.txt").is_file():
            raise RunRefusal(f"Octave failed: {out.stderr.strip()[-500:]}")
        values = {}
        for line in (tmp / "out.txt").read_text(encoding="ascii").splitlines():
            k, p, q = line.split()
            values[int(k)] = (float(p), float(q))
    if sorted(values) != list(range(len(pairs))):
        raise RunRefusal("Octave returned a different number of pairs")
    stderr = [s for s in out.stderr.splitlines() if s.strip()]
    return [values[k] for k in range(len(pairs))], stderr


def ours(fn, ref, dist):
    """(status, reason, value) as SPEC-RESULT maps a core call."""
    try:
        return "valid", None, fn(ref, dist, max_pixels=ref.size).value
    except Refusal as refusal:
        return "no_result", refusal.reason, None
    except Exception:  # noqa: BLE001 - any other exception is status `error` by SPEC-RESULT
        return "error", "internal_error", None


def json_number(x):
    if x is None or math.isfinite(x):
        return x
    return "+inf" if x > 0 else ("-inf" if x < 0 else "nan")


def compare(metric, reference, status, reason, value, ref_value, ref_raw=None, exact=None):
    tol, tol_unit, unit = next((t, tu, u) for m, r, t, tu, u in CHECKS if (m, r) == (metric, reference))
    row = {"metric": metric, "reference": reference, "status": status, "reason": reason,
           "ours": json_number(value), "reference_value": json_number(ref_value),
           "reference_raw": json_number(ref_raw) if ref_raw is not None else None,
           "reference_exact": str(exact) if exact is not None else None, "unit": unit,
           "delta": None, "abs_delta": None, "delta_ulp": None,
           "tolerance": tol, "tolerance_unit": tol_unit, "abs_delta_over_tolerance": None, "within_tolerance": False}
    if status == "valid":
        if math.isinf(value) or math.isinf(ref_value):
            row["within_tolerance"] = value == ref_value
            if row["within_tolerance"]:
                row.update(delta=0.0, abs_delta=0.0, abs_delta_over_tolerance=0.0)
                if tol_unit == "ulp":
                    row["delta_ulp"] = 0.0
        else:
            delta = value - ref_value
            row.update(delta=delta, abs_delta=abs(delta))
            if tol_unit == "ulp":
                ulps = float(abs(Decimal(value) - exact) / Decimal(math.ulp(float(exact)))) if exact != 0 else (
                    0.0 if value == 0 else math.inf)
                row["delta_ulp"] = json_number(ulps)
                # bit-exact where the exact value is an integer (65025*N/SSE = 10**k)
                row["within_tolerance"] = value == float(exact) if exact == exact.to_integral_value() else ulps <= tol
                row["abs_delta_over_tolerance"] = json_number(ulps / tol)
            else:
                row["within_tolerance"] = abs(delta) <= tol
                row["abs_delta_over_tolerance"] = abs(delta) / tol
    row["verdict"] = "pass" if row["within_tolerance"] else "fail"
    return row


def measure_pair(ref, dist, octave_values, haarpsi_py):
    """The four rows of one pair (CHECKS order)."""
    rows = []
    exact = psnr_oracle.psnr(ref.tolist(), dist.tolist())
    status, reason, value = ours(psnr, ref, dist)
    rows.append(compare("psnr", "oracle_decimal", status, reason, value,
                        math.inf if exact.is_infinite() else float(exact), exact=exact))
    p_raw, h_m = octave_values
    status, reason, value = ours(psnr_hvs_m, ref, dist)
    rows.append(compare("psnr_hvs_m", "psnrhvsm.m", status, reason, value,
                        math.inf if p_raw == 100000 else p_raw, ref_raw=p_raw))
    status, reason, value = ours(haarpsi, ref, dist)
    rows.append(compare("haarpsi", "haarPsi.py", status, reason, value, float(haarpsi_py.haar_psi(ref, dist)[0])))
    rows.append(compare("haarpsi", "HaarPSI.m", status, reason, value, h_m))
    return rows


def summarize(rows):
    summary = []
    for metric, reference, tol, tol_unit, unit in CHECKS:
        sel = [r for r in rows if (r["metric"], r["reference"]) == (metric, reference)]
        key = "delta_ulp" if tol_unit == "ulp" else "abs_delta"
        absd = sorted(r[key] for r in sel if isinstance(r[key], float))
        signed = [r["delta"] for r in sel if isinstance(r["delta"], float)]
        q = statistics.quantiles(absd, n=4, method="inclusive") if len(absd) > 1 else [None] * 3
        summary.append({
            "metric": metric, "reference": reference, "n_pairs": len(sel),
            "n_images": len({r["image_code"] for r in sel}), "n_valid": sum(r["status"] == "valid" for r in sel),
            "n_within": sum(r["within_tolerance"] for r in sel), "measure": f"|delta|, {tol_unit}",
            "max": max(absd) if absd else None, "q1": q[0], "median": q[1], "q3": q[2],
            "signed_median_delta": statistics.median(signed) if signed else None,
            "max_abs_delta": max(abs(d) for d in signed) if signed else None, "delta_unit": unit,
            "tolerance": tol, "tolerance_unit": tol_unit,
            "verdict": "pass" if sel and all(r["within_tolerance"] for r in sel) else "fail",
        })
    return summary


def verdicts(summary):
    by = {(s["metric"], s["reference"]): s["verdict"] for s in summary}
    v1b = by["psnr_hvs_m", "psnrhvsm.m"]
    v1c = "pass" if by["haarpsi", "haarPsi.py"] == by["haarpsi", "HaarPSI.m"] == "pass" else "fail"
    return {"V1-a_run_pairs": by["psnr", "oracle_decimal"], "V1-b": v1b, "V1-c": v1c,
            "refutation_fired": not (v1b == v1c == "pass")}


def path_control(octave, src, haarpsi_py, workdir):
    frame, noisy = control_frame()
    octave_values, stderr = octave_metrics(octave, src, [(frame, noisy)], workdir)
    rows = measure_pair(frame, noisy, octave_values[0], haarpsi_py)
    for r in rows:
        r.update(image_code="E0_synthetic", image_class="control", distortion="noise_s3_seed1")
    passed = all(r["within_tolerance"] for r in rows)
    return {"frame": "y,x = np.mgrid[0:512,0:512]; rng = np.random.default_rng(1); "
                     "(x/2 + y/4 + rng.normal(0, 8, (512, 512))).clip(0, 255).astype(np.uint8)",
            "distortion": "noise_s3_seed1", "rows": rows, "passed": passed,
            "frame_sha256": hashlib.sha256(frame.tobytes()).hexdigest(),
            "octave_stderr_lines": len(stderr)}


# --- run file and book -----------------------------------------------------------------------------

def run(paths, images, out_dir, state, codes, distortions=DISTORTIONS):
    """Measure all pairs of `images` x `distortions`, write the run file, the code table and the book."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, codes_path, book_path = out_dir / RUN_FILE, out_dir / CODES_FILE, out_dir / BOOK_FILE
    for p in (run_path, codes_path):
        if p.exists():
            raise RunRefusal(f"{p} already exists: a run file of cycle {CYCLE} is never overwritten; "
                             f"a new run is a new cycle")
    started = datetime.datetime.now(datetime.timezone.utc)
    src, octave = paths["metric_sources"], paths["octave_cli"]
    sources = check_sources(src)
    env = {**environment(), **octave_versions(octave)}
    haarpsi_py = load_haarpsi_py(src)
    print_declared()

    control = path_control(octave, src, haarpsi_py, out_dir)
    for r in control["rows"]:
        print(f"  path control {r['metric']:<10} vs {r['reference']:<14} |delta| = {r['abs_delta']} "
              f"({r['tolerance']} {r['tolerance_unit']}) {r['verdict']}")
    if not control["passed"]:
        raise RunRefusal("path control (synthetic frame E0) is out of tolerance: stop before the 40 pairs")

    pairs = [(img, d, distort(img["array"], d)) for img in images for d in distortions]
    octave_values, stderr = octave_metrics(octave, src, [(img["array"], dist) for img, _, dist in pairs], out_dir)
    rows, pair_table = [], []
    for k, ((img, d, dist), ov) in enumerate(zip(pairs, octave_values, strict=True), start=1):
        pair_rows = measure_pair(img["array"], dist, ov, haarpsi_py)
        for r in pair_rows:
            r.update(pair_index=k, image_code=img["code"], image_class=img["class"], distortion=d["id"])
        rows.extend(pair_rows)
        pair_table.append({"pair_index": k, "image_code": img["code"], "image_class": img["class"],
                           "distortion": d["id"], "dist_pixels_sha256": hashlib.sha256(dist.tobytes()).hexdigest(),
                           "all_within": all(r["within_tolerance"] for r in pair_rows)})
    summary = summarize(rows)
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        **state, "declared_before_numbers": DECLARED, "environment": env,
        "metric_sources_sha256": sources,
        "octave_call": "octave-cli --no-gui --quiet --norc studies/P001_vvc-dental/octave_v1.m <metric_sources> <io>; "
                       "warning 'Octave:possible-matlab-short-circuit-operator' off",
        "octave_stderr_lines": len(stderr), "max_pixels_rule": MAX_PIXELS_RULE,
        "noise_rule": NOISE_RULE, "jpeg_rule": JPEG_RULE, "distortions": distortions,
        "images": [{k: v for k, v in img.items() if k != "array"}
                   | {"pixels_sha256": hashlib.sha256(img["array"].tobytes()).hexdigest()} for img in images],
        "unit_of_analysis": "pair = image x distortion; n is counted by images",
        "path_control": control, "pairs": pair_table, "rows": rows, "summary": summary,
        "verdicts": verdicts(summary),
    }
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        json.dump(record, f, ensure_ascii=False, indent=1, allow_nan=False)
    with open(codes_path, "x", encoding="utf-8", newline="\n") as f:
        json.dump({"_what": "Dental image code -> source file of run v1_metrics (cycle 1). Names identify "
                            "patients: this file stays in the cache, never in a repository.", "codes": codes},
                  f, ensure_ascii=False, indent=1)
    build_book(run_path, book_path)
    leaks = leak_count([run_path, book_path], codes)
    print_summary(json.loads(run_path.read_text(encoding="utf-8")), run_path, book_path, leaks)
    return run_path, book_path, leaks


def leak_count(files, codes):
    """Occurrences of dental source names (file name and stem) in the given files and inside xlsx parts."""
    needles = {c["file_name"] for c in codes} | {Path(c["file_name"]).stem for c in codes}
    found = 0
    for path in files:
        texts = []
        if path.suffix == ".xlsx":
            with zipfile.ZipFile(path) as z:
                texts = [z.read(n).decode("utf-8", "replace") for n in z.namelist()]
        else:
            texts = [path.read_text(encoding="utf-8")]
        found += sum(t.count(n) for t in texts for n in needles)
    return found


CLASS_RU = {"test": "тестовое", "dental": "стоматологическое"}
VERDICT_RU = {"pass": "в допуске", "fail": "сверх допуска"}
# wide table: (header, run-file key, getter)
TABLE = [
    ("№ пары", "pairs[].pair_index", lambda p, r: p["pair_index"]),
    ("Изображение (код)", "pairs[].image_code", lambda p, r: p["image_code"]),
    ("Класс", "pairs[].image_class", lambda p, r: CLASS_RU[p["image_class"]]),
    ("Искажение", "pairs[].distortion", lambda p, r: p["distortion"]),
    ("PSNR наш, дБ", "rows[psnr/oracle_decimal].ours", lambda p, r: r[0]["ours"]),
    ("PSNR оракул decimal, дБ", "rows[psnr/oracle_decimal].reference_value", lambda p, r: r[0]["reference_value"]),
    ("|Δ| PSNR, ulp", "rows[psnr/oracle_decimal].delta_ulp", lambda p, r: r[0]["delta_ulp"]),
    ("PSNR-HVS-M наш, дБ", "rows[psnr_hvs_m/psnrhvsm.m].ours", lambda p, r: r[1]["ours"]),
    ("PSNR-HVS-M psnrhvsm.m, дБ", "rows[psnr_hvs_m/psnrhvsm.m].reference_value", lambda p, r: r[1]["reference_value"]),
    ("Δ PSNR-HVS-M, дБ", "rows[psnr_hvs_m/psnrhvsm.m].delta", lambda p, r: r[1]["delta"]),
    ("HaarPSI наш, безразм.", "rows[haarpsi/haarPsi.py].ours", lambda p, r: r[2]["ours"]),
    ("HaarPSI haarPsi.py, безразм.", "rows[haarpsi/haarPsi.py].reference_value", lambda p, r: r[2]["reference_value"]),
    ("Δ к haarPsi.py, безразм.", "rows[haarpsi/haarPsi.py].delta", lambda p, r: r[2]["delta"]),
    ("HaarPSI HaarPSI.m, безразм.", "rows[haarpsi/HaarPSI.m].reference_value", lambda p, r: r[3]["reference_value"]),
    ("Δ к HaarPSI.m, безразм.", "rows[haarpsi/HaarPSI.m].delta", lambda p, r: r[3]["delta"]),
    ("Пара в допуске по всем сверкам", "pairs[].all_within", lambda p, r: "да" if p["all_within"] else "нет"),
]
SUMMARY_COLS = [
    ("Метрика", "metric"), ("Эталон", "reference"), ("Пар, шт", "n_pairs"), ("Изображений, шт", "n_images"),
    ("Пар со значением, шт", "n_valid"), ("Пар в допуске, шт", "n_within"), ("Мера", "measure"),
    ("Наибольшее", "max"), ("Q1", "q1"), ("Медиана", "median"), ("Q3", "q3"),
    ("Знаковая медиана Δ", "signed_median_delta"), ("Наибольшее |Δ|", "max_abs_delta"), ("Единица Δ", "delta_unit"),
    ("Допуск", "tolerance"), ("Единица допуска", "tolerance_unit"), ("Вердикт", "verdict"),
]
CHART_SERIES = ["PSNR / оракул", "PSNR-HVS-M / psnrhvsm.m", "HaarPSI / haarPsi.py", "HaarPSI / HaarPSI.m"]


def _rows_by_pair(record):
    by = {}
    for r in record["rows"]:
        by.setdefault(r["pair_index"], {})[(r["metric"], r["reference"])] = r
    order = [(m, ref) for m, ref, *_ in CHECKS]
    return {k: [v[c] for c in order] for k, v in by.items()}


def book_cells(record, run_sha256):
    """Every cell of the book as {(sheet, row, col): value}, built from the run file only."""
    cells = {}

    def put(sheet, row, col, value):
        cells[sheet, row, col] = value

    rows_by_pair = _rows_by_pair(record)
    d = "Описание"
    lines = [
        ("Команда", record["command"]), ("Вопрос", record["question"]),
        ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
        ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
        ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
        ("Роль данных", record["declared_before_numbers"]["roles"]),
        ("Искажения", record["declared_before_numbers"]["distortions"]),
        ("Критерий (до чисел)", record["declared_before_numbers"]["criterion"]),
        ("Опровержение (до чисел)", record["declared_before_numbers"]["refutation"]),
        ("Ожидаемо (до чисел)", record["declared_before_numbers"]["expected"]),
        ("Контроль пути (до чисел)", record["declared_before_numbers"]["path_control"]),
        ("Страницы для глаз", record["declared_before_numbers"]["pages"]),
        ("V1-a в прогоне", record["declared_before_numbers"]["run_part_of_v1a"]),
        ("Шум", record["noise_rule"]), ("JPEG", record["jpeg_rule"]), ("max_pixels", record["max_pixels_rule"]),
        ("Контроль пути пройден", "да" if record["path_control"]["passed"] else "нет"),
        ("V1-a (пары прогона)", VERDICT_RU[record["verdicts"]["V1-a_run_pairs"]]),
        ("V1-b", VERDICT_RU[record["verdicts"]["V1-b"]]), ("V1-c", VERDICT_RU[record["verdicts"]["V1-c"]]),
        ("Опровержение сработало", "да" if record["verdicts"]["refutation_fired"] else "нет"),
        ("Листы «Рисунки» и «Страницы»", "не строятся: правило страниц не применяется (объявлено до чисел)"),
        ("Точность чисел книги", "16 значащих цифр (openpyxl пишет %.16g); каждая клетка сверена с числом файла "
                                 "прогона, округлённым до 16 цифр; полная точность — в файле прогона"),
        ("На пальцах", "Мы посчитали три меры качества своим кодом и кодом авторов на одних и тех же парах "
                       "изображений и сравнили ответы. Таблица сводки ниже показывает, насколько ответы "
                       "расходятся и укладываются ли они в объявленный заранее допуск."),
    ]
    for i, (k, v) in enumerate(lines, start=1):
        put(d, i, 1, k)
        put(d, i, 2, v)
    top = len(lines) + 2
    put(d, top, 1, "Сводка по сверкам (ключ файла прогона: summary[])")
    for j, (head, key) in enumerate(SUMMARY_COLS, start=1):
        put(d, top + 1, j, head)
        put(d, top + 2, j, key)
    for i, s in enumerate(record["summary"]):
        for j, (_, key) in enumerate(SUMMARY_COLS, start=1):
            value = s[key]
            put(d, top + 3 + i, j, VERDICT_RU[value] if key == "verdict" else value)

    t = "Таблица"
    for j, (head, key, _) in enumerate(TABLE, start=1):
        put(t, 1, j, head)
        put(t, 2, j, key)
    for i, p in enumerate(record["pairs"]):
        r = rows_by_pair[p["pair_index"]]
        for j, (_, _, get) in enumerate(TABLE, start=1):
            put(t, 3 + i, j, get(p, r))

    g = "График"
    put(g, 1, 1, "|Δ| / допуск по парам, безразм.; логарифмическая ось. Пустая клетка — |Δ| = 0 "
                 "(побитовое совпадение): на логарифмической оси ноль не рисуется. "
                 "Ключ: rows[].abs_delta_over_tolerance")
    put(g, 2, 1, "№ пары")
    for j, name in enumerate(CHART_SERIES, start=2):
        put(g, 2, j, name)
    for i, p in enumerate(record["pairs"]):
        put(g, 3 + i, 1, p["pair_index"])
        for j, r in enumerate(rows_by_pair[p["pair_index"]], start=2):
            v = r["abs_delta_over_tolerance"]
            put(g, 3 + i, j, None if v == 0 else v)
    return cells


def build_book(run_path, book_path):
    """Write the book from the run file and check every cell against it; a mismatch refuses."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    record = json.loads(run_bytes.decode("utf-8"))
    cells = book_cells(record, hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Таблица", "График"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    desc = wb["Описание"]
    desc.column_dimensions["A"].width = 30
    desc.column_dimensions["B"].width = 120
    for row in desc.iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for sheet in ("Таблица",):
        for c in wb[sheet][1]:
            c.font = Font(bold=True)
    n = len(record["pairs"])
    chart = BarChart()
    chart.title = "|Δ| / допуск по 40 парам (log)"
    chart.y_axis.title = "|Δ| / допуск, безразм."
    chart.x_axis.title = "№ пары"
    chart.y_axis.scaling.logBase = 10
    chart.x_axis.delete = chart.y_axis.delete = False  # recent Excel hides axes otherwise
    chart.add_data(Reference(wb["График"], min_col=2, max_col=1 + len(CHART_SERIES), min_row=2, max_row=2 + n),
                   titles_from_data=True)
    chart.set_categories(Reference(wb["График"], min_col=1, min_row=3, max_row=2 + n))
    chart.width, chart.height = 32, 12
    wb["График"].add_chart(chart, "G2")
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell of the book equals the value built from the run file, and no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    check = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = check[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v  # openpyxl writes floats as %.16g
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in check.sheetnames:
        filled = {(sheet, c.row, c.column) for row in check[sheet].iter_rows() for c in row if c.value is not None}
        extra = filled - set(cells)
        if extra:
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra)[:3]}")


# --- console -------------------------------------------------------------------------------------------

def print_declared():
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")


def print_summary(record, run_path, book_path, leaks):
    print("\nSummary (n = pairs = image x distortion; independent units are images):")
    for s in record["summary"]:
        print(f"  {s['metric']:<10} vs {s['reference']:<14} n={s['n_pairs']:>2} images={s['n_images']:>2} "
              f"within={s['n_within']:>2}  max {s['measure']} = {s['max']!r}  median = {s['median']!r}  "
              f"signed median delta = {s['signed_median_delta']!r} {s['delta_unit']}  "
              f"tolerance {s['tolerance']} {s['tolerance_unit']}  -> {s['verdict']}")
    v = record["verdicts"]
    print(f"Verdicts: V1-a (run pairs) {v['V1-a_run_pairs']}, V1-b {v['V1-b']}, V1-c {v['V1-c']}; "
          f"refutation fired: {v['refutation_fired']}")
    print(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}\n"
          f"Leak check (dental source names in run file and book): {leaks}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        for name in (RUN_FILE, CODES_FILE):
            if (out_dir / name).exists():
                raise RunRefusal(f"{out_dir / name} already exists: a run file of cycle {CYCLE} is never "
                                 f"overwritten; a new run is a new cycle")
        state = git_state()
        dental, codes = dental_set(paths["images_root"])
        run(paths, test_set(paths["images_root"]) + dental, out_dir, state, codes)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
