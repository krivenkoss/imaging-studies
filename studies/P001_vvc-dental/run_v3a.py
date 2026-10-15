"""V3-a: are our JPEG files (Pillow, B06) bit-identical to IJG libjpeg 9e cjpeg -baseline on 42 pairs? Run: python studies/P001_vvc-dental/run_v3a.py

Study P001, cycle 5: check V3-a of STUDY.md section 8 for the JPEG wrapper (../REFERENCE.md#b06).
The reference is the assistant's file of IJG libjpeg 9e files (path: key v3a_reference of
lab_paths.json). Writes <cache_root>/cycles/5/v3a_jpeg.json (refuses by name if it exists) and the
book v3a_jpeg.xlsx beside it; `--rebuild-book` rebuilds the book from the run file.
"""

import argparse
import datetime
import hashlib
import json
import platform
import sys
from pathlib import Path

import numpy as np

import imstudies
from imstudies.core.metrics import Refusal
from imstudies.lab.coders import jpeg
from imstudies.lab.images import load_test_images
from imstudies.lab.runs import RunRefusal, git_state, sha256_file

STUDY = Path(__file__).resolve().parent
ROOT = STUDY.parents[1]
COMMAND = "python studies/P001_vvc-dental/run_v3a.py"
QUESTION = ("V3-a: совпадают ли побитово наши файлы JPEG (Pillow 12.3.0 → libjpeg-turbo, B06) с файлами "
            "IJG libjpeg 9e cjpeg -quality q -baseline на 42 парах «изображение × quality»?")
CYCLE, RUN_FILE, BOOK_FILE = 5, "v3a_jpeg.json", "v3a_jpeg.xlsx"
TEST_IMAGES = ["BABOON.BMP", "BARBARA.BMP", "goldhill.bmp", "lenna.bmp", "frisco.bmp", "DIEGO.BMP"]
QUALITIES = (1, 5, 30, 50, 75, 95, 100)
DECLARED = {
    "data": ("Шесть тестовых BMP цикла 1 (чтение imstudies.lab.images.load_test_images) × quality = 1, 5, 30, 50, "
             "75, 95, 100 — 42 файла B06. Изображения отладочные, в приёмке P001 не участвуют."),
    "criterion": ("sha256 каждого файла = jpeg_sha256 своей строки эталона ассистента v3a_ijg9e.json (IJG libjpeg 9e, "
                  "cjpeg -quality q -baseline, Linux); pixels_sha256 эталона = циклу 1 — сверяется до сравнения."),
    "refutation": "Любое несовпадение — V3-a не пройдена; стоп, разбор, не подгонка.",
    "sizes": "Размеры файлов в этом цикле — не результат исследования (тестовые изображения): печатаются для контроля.",
    "pages": "Страницы для глаз не применяются — побитовая сверка.",
}


def load_paths():
    path = ROOT / "lab_paths.json"
    if not path.is_file():
        raise RunRefusal(f"{path} not found: copy lab_paths.example.json to lab_paths.json and set the paths")
    raw = json.loads(path.read_text(encoding="utf-8"))
    paths = {}
    for key in ("images_root", "cache_root", "v3a_reference"):
        if not raw.get(key):
            raise RunRefusal(f"lab_paths.json has no key '{key}'")
        paths[key] = (ROOT / raw[key]).resolve()
    if not paths["v3a_reference"].is_file():
        raise RunRefusal(f"V3-a reference not found: {paths['v3a_reference']}")
    return paths


def check_reference(reference, images, cycle1_images):
    """The reference rows cover every (image, quality) and were made from the pixels of cycle 1 and of this run."""
    ours = {img["code"]: hashlib.sha256(img["array"].tobytes()).hexdigest() for img in images}
    cycle1 = {i["code"]: i["pixels_sha256"] for i in cycle1_images if i["class"] == "test"}
    differ = [code for code, sha in ours.items() if cycle1.get(code) != sha]
    if differ:
        raise RunRefusal(f"the test images read now differ from cycle 1 (pixels sha256): {differ}")
    for row in reference["rows"]:
        if row["pixels_sha256"] != cycle1.get(row["code"]):
            raise RunRefusal(f"reference row {row['code']} q {row['quality']}: pixels sha256 differs from cycle 1")
    keys = {(r["code"], r["quality"]) for r in reference["rows"]}
    missing = [(c, q) for c in ours for q in QUALITIES if (c, q) not in keys]
    if missing:
        raise RunRefusal(f"reference has no rows for {missing[:5]}")


def run(paths, images, qualities, out_dir, state):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    run_path, book_path = out_dir / RUN_FILE, out_dir / BOOK_FILE
    if run_path.exists():
        raise RunRefusal(f"{run_path} already exists: a run file of cycle {CYCLE} is never overwritten")
    try:
        jpeg.check_versions("JPEG", jpeg.EXPECTED_VERSIONS, jpeg.versions())
    except RuntimeError as exc:
        raise RunRefusal(str(exc)) from exc
    reference = json.loads(paths["v3a_reference"].read_text(encoding="utf-8"))
    cycle1 = json.loads((paths["cache_root"] / "cycles" / "1" / "v1_metrics.json").read_text(encoding="utf-8"))
    check_reference(reference, images, cycle1["images"])
    by_key = {(r["code"], r["quality"]): r for r in reference["rows"]}
    started = datetime.datetime.now(datetime.timezone.utc)
    print(f"{COMMAND}\n{QUESTION}\nDeclared before the numbers:")
    for key, text in DECLARED.items():
        print(f"  [{key}] {text}")
    rows = []
    for img in images:
        for q in qualities:
            ref = by_key[img["code"], q]
            row = {"pair_index": len(rows) + 1, "image_code": img["code"], "quality": q,
                   "reference_bytes": ref["bytes"], "reference_sha256": ref["jpeg_sha256"]}
            try:
                result = jpeg.encode_decode(img["array"], q, max_pixels=img["array"].size)
            except Refusal as refusal:
                row |= {"status": "no_result", "reason": refusal.reason, "bytes": None, "sha256": None, "match": False}
            except Exception as exc:  # noqa: BLE001 - status `error` by SPEC-RESULT
                row |= {"status": "error", "reason": "internal_error", "detail": str(exc)[:600], "bytes": None,
                        "sha256": None, "match": False}
            else:
                sha = hashlib.sha256(result.file).hexdigest()
                row |= {"status": "valid", "reason": None, "bytes": len(result.file), "sha256": sha,
                        "match": sha == ref["jpeg_sha256"],
                        "recon_shape_ok": result.recon.shape == img["array"].shape}
            rows.append(row)
            print(f"  {row['pair_index']:>2} {img['code']:<9} q {q:>3}: {row['status']} {row['bytes']} bytes, "
                  f"match {row['match']}")
    n_match = sum(r["match"] for r in rows)
    record = {
        "run": RUN_FILE, "study": "P001", "cycle": CYCLE, "question": QUESTION, "command": COMMAND,
        "started_utc": started.isoformat(timespec="seconds"),
        "finished_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        **state, "declared_before_numbers": DECLARED,
        "reference": {"file": paths["v3a_reference"].name, "sha256": sha256_file(paths["v3a_reference"]),
                      "what": reference.get("_what"), "ijg": reference.get("ijg")},
        "wrapper": {"versions": jpeg.versions(), "encode_args": jpeg.ENCODE_ARGS, "quality_range": jpeg.QUALITY_RANGE},
        "environment": {"python": platform.python_version(), "imstudies": imstudies.__version__,
                        "numpy": np.__version__, "platform": platform.platform()},
        "images": [{k: v for k, v in img.items() if k != "array"}
                   | {"pixels_sha256": hashlib.sha256(img["array"].tobytes()).hexdigest()} for img in images],
        "rows": rows,
        "summary": {"n_pairs": len(rows), "n_images": len(images), "n_match": n_match,
                    "verdict_v3a": "pass" if rows and n_match == len(rows) else "fail",
                    "refutation_fired": n_match != len(rows)},
    }
    with open(run_path, "x", encoding="utf-8", newline="\n") as f:
        json.dump(record, f, ensure_ascii=False, indent=1, allow_nan=False)
    build_book(run_path, book_path)
    s = record["summary"]
    print(f"\nV3-a: {s['n_match']} of {s['n_pairs']} files bit-identical to IJG libjpeg 9e ({s['n_images']} images) -> "
          f"{s['verdict_v3a']}; refutation fired: {s['refutation_fired']}")
    print(f"Run file: {run_path} (sha256 {sha256_file(run_path)})\nBook: {book_path}")
    return run_path, book_path


TABLE = [("№ пары", "pair_index"), ("Изображение (код)", "image_code"), ("quality", "quality"),
         ("Статус", "status"), ("Причина", "reason"), ("Наш файл, байт", "bytes"), ("Эталон, байт", "reference_bytes"),
         ("sha256 нашего файла", "sha256"), ("sha256 эталона", "reference_sha256"), ("Совпадение", "match")]
YES_NO = {True: "да", False: "нет"}


def book_cells(record, run_sha256):
    cells, s = {}, record["summary"]
    lines = [("Команда", record["command"]), ("Вопрос", record["question"]),
             ("Файл прогона", f"{record['run']} (cycles/{record['cycle']}/ под cache_root из lab_paths.json)"),
             ("sha256 файла прогона", run_sha256), ("Начат, UTC", record["started_utc"]),
             ("Коммит core", record["core_commit"]), ("Дерево core", record["core_tree"]),
             *((f"{k} (до чисел)", v) for k, v in record["declared_before_numbers"].items()),
             ("Эталон", f"{record['reference']['file']}, sha256 {record['reference']['sha256']}; {record['reference']['ijg']}"),
             ("Pillow / libjpeg-turbo", f"{record['wrapper']['versions']['pillow']} / "
                                        f"{record['wrapper']['versions']['libjpeg_turbo']}"),
             ("Пар, шт", s["n_pairs"]), ("Изображений, шт", s["n_images"]), ("Совпало, шт", s["n_match"]),
             ("Вердикт V3-a", s["verdict_v3a"]), ("Опровержение сработало", YES_NO[s["refutation_fired"]]),
             ("Точность чисел книги", "целые; каждая клетка сверена с файлом прогона"),
             ("Листы «Рисунки» и «Страницы»", "не строятся: страницы не применяются (объявлено до чисел)"),
             ("На пальцах", "Каждое тестовое изображение сжато в JPEG нашим кодом и эталонной программой IJG при семи "
                            "уровнях качества; сравнивались сами файлы байт в байт.")]
    for i, (k, v) in enumerate(lines, start=1):
        cells["Описание", i, 1], cells["Описание", i, 2] = k, v
    for j, (head, key) in enumerate(TABLE, start=1):
        cells["Таблица", 1, j], cells["Таблица", 2, j] = head, f"rows[].{key}"
    for i, row in enumerate(record["rows"]):
        for j, (_, key) in enumerate(TABLE, start=1):
            v = row.get(key)
            cells["Таблица", 3 + i, j] = YES_NO[v] if isinstance(v, bool) else v
    codes = list(dict.fromkeys(r["image_code"] for r in record["rows"]))
    qs = sorted({r["quality"] for r in record["rows"]})
    cells["График", 1, 1] = "Наш файл, байт, по quality для каждого изображения. Ключ: rows[].bytes"
    cells["График", 2, 1] = "quality"
    for j, code in enumerate(codes, start=2):
        cells["График", 2, j] = code
    for i, q in enumerate(qs):
        cells["График", 3 + i, 1] = q
        for j, code in enumerate(codes, start=2):
            cells["График", 3 + i, j] = next((r["bytes"] for r in record["rows"]
                                              if (r["image_code"], r["quality"]) == (code, q)), None)
    return cells


def build_book(run_path, book_path):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Описание", "Таблица", "График"):
        wb.create_sheet(name)
    for (sheet, r, c), v in cells.items():
        wb[sheet].cell(row=r, column=c, value=v)
    wb["Описание"].column_dimensions["A"].width = 32
    wb["Описание"].column_dimensions["B"].width = 120
    for row in wb["Описание"].iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    for c in wb["Таблица"][1]:
        c.font = Font(bold=True)
    g = wb["График"]
    n_rows = max(r for s, r, _ in cells if s == "График")
    n_cols = max(c for s, _, c in cells if s == "График")
    if n_rows >= 3 and n_cols >= 2:
        chart = LineChart()
        chart.title, chart.y_axis.title, chart.x_axis.title = "Файл JPEG, байт, по quality", "байт", "quality"
        chart.x_axis.delete = chart.y_axis.delete = False
        chart.add_data(Reference(g, min_col=2, max_col=n_cols, min_row=2, max_row=n_rows), titles_from_data=True)
        chart.set_categories(Reference(g, min_col=1, min_row=3, max_row=n_rows))
        chart.width, chart.height = 24, 12
        g.add_chart(chart, "J2")
    wb.save(book_path)
    check_book(run_path, book_path)


def check_book(run_path, book_path):
    """Every cell equals the run-file value (floats as %.16g, as openpyxl writes them); no other cell is filled."""
    from openpyxl import load_workbook

    run_bytes = Path(run_path).read_bytes()
    cells = book_cells(json.loads(run_bytes.decode("utf-8")), hashlib.sha256(run_bytes).hexdigest())
    book = load_workbook(book_path)
    for (sheet, r, c), v in cells.items():
        got = book[sheet].cell(row=r, column=c).value
        expected = float(f"{v:.16g}") if isinstance(v, float) else v
        if got != expected:
            raise RunRefusal(f"book cell {sheet}!R{r}C{c} = {got!r} differs from the run file value {v!r}")
    for sheet in book.sheetnames:
        extra = {(sheet, x.row, x.column) for row in book[sheet].iter_rows() for x in row if x.value is not None}
        if extra - set(cells):
            raise RunRefusal(f"book sheet {sheet} has cells not taken from the run file: {sorted(extra - set(cells))[:3]}")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--rebuild-book", action="store_true", help="rebuild the book from the existing run file")
    args = parser.parse_args(argv)
    try:
        paths = load_paths()
        out_dir = paths["cache_root"] / "cycles" / str(CYCLE)
        if args.rebuild_book:
            if not (out_dir / RUN_FILE).is_file():
                raise RunRefusal(f"no run file to rebuild the book from: {out_dir / RUN_FILE}")
            build_book(out_dir / RUN_FILE, out_dir / BOOK_FILE)
            print(f"Book rebuilt and checked cell by cell: {out_dir / BOOK_FILE}")
            return 0
        if (out_dir / RUN_FILE).exists():
            raise RunRefusal(f"{out_dir / RUN_FILE} already exists: a run file of cycle {CYCLE} is never overwritten")
        state = git_state(ROOT)
        images = load_test_images(paths["images_root"], TEST_IMAGES, (512, 512))
        run(paths, images, QUALITIES, out_dir, state)
    except RunRefusal as refusal:
        print(f"REFUSED: {refusal}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
